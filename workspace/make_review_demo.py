"""
Builds the Word review demo in review_demo/ (OFFICE_COLLAB_DESIGN.md, W9), and its answer key one folder up:

    review_demo/test_plan.docx      a three-page fire test plan, reviewed by two fictional people: A. Reviewer (7
                                    tracked changes, 4 comments) and B. Checker (3 changes, one of them a formatting
                                    change, and 2 comments). One thread has two replies, one thread is resolved, a
                                    table cell is changed, and a whole paragraph is inserted
    review_demo/test_plan_v2.docx   the same plan as a colleague sent it back: every change accepted, no comments, and
                                    four differences planted (a reworded sentence, a deleted paragraph, a new table row,
                                    a changed number)
    ../review_demo_answers.json     the census, the threads, the v1/v2 differences and the expected texts, outside
                                    workspace/ so no tool can show it to the LLM (compare_review.py scores against it)

    python workspace/make_review_demo.py

Everything here is synthetic: the repository is public, so no real proposal or report goes in. The tracked changes and
comments are written with the harness's own writer (harness/tools/docx_markup.py), so building the demo is also the
writer's first test; the expected texts in the key are worked out here with plain string replacements, not with the
reader they are used to check. The harness never changes these files.
"""

import json
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
PROJECT = HERE.parent
sys.path.insert(0, str(PROJECT))

from harness.tools import docx_markup as dm
from harness.tools.word_model import accepted_text, paragraphs_of

FOLDER = HERE / "review_demo"
A, B = "A. Reviewer", "B. Checker"
A_DATE, B_DATE, A_REPLY_DATE = "2026-10-02T10:00:00Z", "2026-10-03T14:30:00Z", "2026-10-04T09:15:00Z"

# The plan before anyone reviewed it: (style, text). "T" marks where Table 1 goes.
PLAN = [
    ("Title", "Fire Test Plan: Compartment Tests with an Automatic Sprinkler"),
    ("Normal", "Synthetic example for the glass-box teaching harness. Every name, number and place in it is made up."),
    ("Heading 1", "1 Purpose"),
    ("Normal", "This plan describes a compartment fire test series in a small room with one automatic sprinkler. The "
               "tests measure how quickly the hot gas layer forms, when the sprinkler activates, and how much the "
               "spray cools the room afterwards. The results will be used to check the sprinkler model."),
    ("Normal", "The plan is written for the test crew, the instrumentation team and the reviewers. It says what is "
               "measured, how the fires are built, and when a test is stopped. Questions go to the test director "
               "before the first ignition, not during a test."),
    ("Heading 1", "2 Scope"),
    ("Normal", "The series covers a compartment fire test with wood cribs and a propane burner as the fuel. Each fuel "
               "load is burned under the same sprinkler, with the door open. Tests with the door closed, with "
               "upholstered furniture or with more than one sprinkler are outside this plan."),
    ("Normal", "The measurements cover the first twenty minutes after ignition. Longer tests are run only if the "
               "test director asks for them in writing beforehand."),
    ("Heading 1", "3 Test Facility"),
    ("Normal", "The room is 3.6 m long, 2.4 m wide and 2.4 m high, with one opening, the doorway, on the short wall. "
               "The burn hall around the room is large enough that its air stays below 25 °C during a test."),
    ("Normal", "The walls are lined with two layers of gypsum board over a steel frame. The ceiling is a single "
               "sheet of calcium silicate board, so that the sprinkler sits on a flat, smooth surface."),
    ("Normal", "Smoke leaves the room through the doorway and rises into a large calorimetry hood. The hood draws "
               "the smoke into the exhaust duct, where the oxygen, carbon dioxide and flow are measured."),
    ("Heading 1", "4 Instrumentation"),
    ("Normal", "Gas temperatures are measured with thermocouple trees in the front and rear corners of the room. "
               "Each tree carries Type K thermocouples at eight heights, from 0.3 m to 2.3 m above the floor."),
    ("Normal", "Heat flux gauges sit on the floor at the centre of the room and at the doorway sill. Data are "
               "logged and recorded every second from five minutes before ignition to the end of the test."),
    ("Normal", "Each test is filmed from two cameras: one at the doorway and one looking down from the hood."),
    ("Heading 1", "5 Test Conditions"),
    ("Normal", "Table 1 lists the three tests. The fuel load and the target peak heat release rate change from test "
               "to test; the room, the sprinkler and the door stay the same."),
    ("T", ""),
    ("Heading 1", "6 Procedure"),
    ("Normal", "The crew checks every instrument, zeroes the gas analysers and starts the data logger. The fuel is "
               "ignited with a small propane torch at the centre of the crib."),
    ("Normal", "The sprinkler must activate within 5 minutes of ignition. If it has not activated by then, the test "
               "director stops the test and the crew puts the fire out by hand."),
    ("Normal", "After activation the sprinkler runs for approximately 20 minutes. The test ends when every "
               "thermocouple in the room reads below 50 °C."),
    ("Heading 1", "7 Acceptance Criteria"),
    ("Normal", "A test is valid when every channel recorded data for the whole test. The ceiling gas temperature "
               "must not exceed 600 °C at any thermocouple; a test above it is repeated with a smaller crib."),
    ("Normal", "The heat release rate from the hood must agree with the mass loss of the fuel within 15 percent."),
    ("Heading 1", "8 Safety"),
    ("Normal", "Nobody enters the room between ignition and the end of the test. Two people with extinguishers stand "
               "by the doorway throughout."),
    ("Normal", "Hearing protection and safety glasses are worn in the burn hall while a test runs."),
]
TABLE = [["Test", "Fuel", "Peak HRR", "Ventilation"],
         ["1", "Wood crib, 20 kg", "300 kW", "Door open"],
         ["2", "Wood crib, 30 kg", "500 kW", "Door open"],
         ["3", "Propane burner", "250 kW", "Door open"]]

# The reviewers' work. A tracked text edit: (author, find, replace); every find occurs once in the plan.
EDITS = [
    (A, "approximately", "about"),                       # a deletion and an insertion: 2 changes
    (A, "measured with thermocouple trees", "measured with three thermocouple trees"),  # 1 insertion
    (A, "logged and recorded every second", "logged every second"),                     # 1 deletion
    (A, "300 kW", "350 kW"),                             # in Table 1: a deletion and an insertion, 2 changes
    (B, "two cameras: one", "two cameras at 30 frames per second: one"),               # 1 insertion
    (B, "a large calorimetry hood", "a calorimetry hood"),                              # 1 deletion
]
NEW_PARAGRAPH = (A, "Two people with extinguishers stand by the doorway throughout.",
                 "A fire watch stays at the door until the room has cooled below 50 °C.")  # 1 insertion
BOLD = (B, "must not exceed 600 °C")                # 1 formatting change
COMMENTS = [  # (label, author, date, anchor, text, done)
    ("c1", A, A_DATE, "a compartment fire test series", "Is this one test or a series of three?", False),
    ("c2", B, B_DATE, "Type K thermocouples", "Give the bead diameter.", False),
    ("c3", A, A_DATE, "the doorway, on the short wall", "Door width?", True),
    ("c4", A, A_DATE, "the exhaust duct", "Which duct sizes?", False),
]
REPLIES = [("c5", B, B_DATE, "c1", "Three tests, one per fuel load (Table 1)."),
           ("c6", A, A_REPLY_DATE, "c1", "Then say so in the Scope.")]
# What the colleague changed in v2 (on the plan with every change accepted)
V2_REWORDED = ("The results will be used to check the sprinkler model.",
               "The results will be used to validate the sprinkler activation model.")
V2_DELETED = "The walls are lined with two layers of gypsum board over a steel frame. The ceiling is a single sheet " \
             "of calcium silicate board, so that the sprinkler sits on a flat, smooth surface."
V2_NEW_ROW = ["4", "Wood crib, 40 kg", "700 kW", "Door open"]
V2_NUMBER = ("The room is 3.6 m long", "The room is 3.8 m long")


def base_document(path):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    for style, text in PLAN:
        if style == "T":
            caption = doc.add_paragraph("Table 1. Test conditions")
            caption.runs[0].bold = True
            table = doc.add_table(rows=len(TABLE), cols=len(TABLE[0]))
            table.style = "Table Grid"
            for row, values in zip(table.rows, TABLE):
                for cell, value in zip(row.cells, values):
                    cell.text = value
            for cell in table.rows[0].cells:
                cell.paragraphs[0].runs[0].bold = True
            continue
        doc.add_paragraph(text, style=style)
    doc.save(path)


def locate(pkg, text):
    hits = [(p, accepted_text(p).index(text)) for p in paragraphs_of(pkg.document) if text in accepted_text(p)]
    assert len(hits) == 1, f"{text!r} occurs {len(hits)} times"
    p, at = hits[0]
    return p, at, at + len(text)


def reviewed(path):
    """The reviewers' comments first (ids 1-6, so 'comment 4' is c4), then their tracked changes."""
    pkg = dm.Package(path)
    writers = {(A, A_DATE): dm.Writer(pkg, A, "AR", A_DATE), (B, B_DATE): dm.Writer(pkg, B, "BC", B_DATE),
               (A, A_REPLY_DATE): dm.Writer(pkg, A, "AR", A_REPLY_DATE)}
    ids = {}
    for label, author, date, anchor, text, _ in COMMENTS:
        p, a, b = locate(pkg, anchor)
        ids[label] = dm.comment_on_text(pkg, p, a, b, text, writers[(author, date)])
        assert f"c{ids[label]}" == label, (label, ids[label])
    for label, author, date, parent, text in REPLIES:
        made = dm.add_reply(pkg, ids[parent], text, writers[(author, date)])
        assert f"c{made}" == label, (label, made)
    for label, author, date, _, _, done in COMMENTS:
        if done:
            dm.resolve_thread(pkg, ids[label], writers[(author, date)])
    for author, find, replace in EDITS:
        p, a, b = locate(pkg, find)
        dm.tracked_replace(p, a, b, replace, writers[(author, A_DATE if author == A else B_DATE)])
    author, after, text = NEW_PARAGRAPH
    p, _, _ = locate(pkg, after)
    dm.insert_paragraphs_after(p, [text], writers[(author, A_DATE)])
    author, phrase = BOLD
    p, a, b = locate(pkg, phrase)
    dm.track_formatting(p, a, b, ["b", "bCs"], writers[(author, B_DATE)])
    dm.set_track_changes(pkg, True)
    pkg.touch(pkg.names["document"])
    problems = dm.structure_problems(pkg)
    assert not problems, problems
    return pkg


def plan_lines(edits=(), new_paragraph=False):
    """The plan's paragraphs as plain text, table cells as their own lines, with some of the edits made: the key's
    expected texts, worked out with string replacements."""
    out = []
    for style, text in PLAN:
        if style == "T":
            out.append("Table 1. Test conditions")
            out += [cell for row in TABLE for cell in row]
            continue
        out.append(text)
        if new_paragraph and text.endswith(NEW_PARAGRAPH[1]):
            out.append(NEW_PARAGRAPH[2])
    for _, find, replace in edits:
        hits = [i for i, line in enumerate(out) if find in line]
        assert len(hits) == 1, find
        out[hits[0]] = out[hits[0]].replace(find, replace)
    return out


def version_two(v1, path):
    """v2: every change accepted, the comments gone, and the four planted differences, untracked."""
    pkg = dm.Package(v1)
    dm.resolve(pkg.document, lambda el: "accept")
    dm.remove_comments(pkg)
    dm.set_track_changes(pkg, False)
    pkg.touch(pkg.names["document"])
    tmp = path.with_name("_v2_tmp.docx")
    pkg.save(tmp)
    doc = Document(tmp)
    for p in doc.paragraphs:
        if V2_REWORDED[0] in p.text:
            p.text = p.text.replace(*V2_REWORDED)
        elif V2_NUMBER[0] in p.text:
            p.text = p.text.replace(*V2_NUMBER)
        elif p.text == V2_DELETED:
            p._p.getparent().remove(p._p)
    row = doc.tables[0].add_row()
    for cell, value in zip(row.cells, V2_NEW_ROW):
        cell.text = value
    doc.save(path)
    tmp.unlink()


def answer_key(v1_lines, accepted, b_only, v2_lines):
    a_kinds = {"insertion": 4, "deletion": 3}
    b_kinds = {"insertion": 1, "deletion": 1, "formatting": 1}
    return {
        "document": "workspace/review_demo/test_plan.docx",
        "revised": "workspace/review_demo/test_plan_v2.docx",
        "about": "The answer key for the Word review demo (workspace/make_review_demo.py writes both). Kept outside "
                 "workspace/ so no tool can show it to the LLM. Counts are as Word's Revisions.Count and "
                 "Comments.Count give them: a replacement is a deletion and an insertion; a whole inserted paragraph "
                 "is one insertion.",
        "census": {"changes": 10, "by_author": {A: 7, B: 3}, "by_author_kind": {A: a_kinds, B: b_kinds},
                   "comments": 6, "threads": 4, "open": 3, "resolved": 1, "replies": 2,
                   "comments_by_author": {A: 4, B: 2}},
        "threads": [{"label": label, "author": author, "anchor": anchor, "text": text, "done": done,
                     "replies": [r[0] for r in REPLIES if r[3] == label]}
                    for label, author, _, anchor, text, done in COMMENTS],
        "replies": [{"label": label, "author": author, "parent": parent, "text": text}
                    for label, author, _, parent, text in REPLIES],
        "v1_v2_differences": [
            {"what": "a reworded sentence", "v1": V2_REWORDED[0], "v2": V2_REWORDED[1]},
            {"what": "a deleted paragraph", "v1": V2_DELETED, "v2": None},
            {"what": "a new table row", "v1": None, "v2": " | ".join(V2_NEW_ROW)},
            {"what": "a changed number", "v1": V2_NUMBER[0], "v2": V2_NUMBER[1]},
        ],
        "texts": {"original": v1_lines, "accepted": accepted, "b_checker_accepted_rest_rejected": b_only,
                  "v2": v2_lines},
        "requests": {
            "redline": {"find": "within 5 minutes", "replace": "within 3 minutes"},
            "reply": {"to": "c4", "resolve": True},
            "accept": {"author": B, "rest": "reject"},
        },
    }


def main():
    FOLDER.mkdir(exist_ok=True)
    v1 = FOLDER / "test_plan.docx"
    base = FOLDER / "_base.docx"
    base_document(base)
    pkg = reviewed(base)
    pkg.save(v1)
    base.unlink()
    v2 = FOLDER / "test_plan_v2.docx"
    version_two(v1, v2)
    original = plan_lines()
    accepted = plan_lines(EDITS, new_paragraph=True)
    b_only = plan_lines([e for e in EDITS if e[0] == B])
    v2_lines = [line for line in accepted if line != V2_DELETED]
    v2_lines = [line.replace(*V2_REWORDED).replace(*V2_NUMBER) for line in v2_lines]
    at = v2_lines.index("Propane burner") + 3  # the new row after the last one (its cells: 3, Propane burner, ...)
    v2_lines[at:at] = V2_NEW_ROW
    key = answer_key(original, accepted, b_only, v2_lines)
    (PROJECT / "review_demo_answers.json").write_text(json.dumps(key, indent=2, ensure_ascii=False) + "\n",
                                                     encoding="utf-8")
    print(f"wrote {v1.relative_to(PROJECT)}, {v2.relative_to(PROJECT)} and review_demo_answers.json")


if __name__ == "__main__":
    main()
