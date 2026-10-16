"""
Builds the binary example files in this folder: data/fire_tests.sqlite (the tables of the FP520
lecture, for query_database), fire_calcs.xlsx (for update_excel), images/hrr_curves.png (a plot
with known numbers), project_update.pptx, old_format_deck.pptx, templates/workshop_template.pptx
(a plain 16:9 template for convert_presentation), report.docx and references/evaluation_results.pdf.
The other examples are plain text files. All of them are in the repository already; run this only
to get fresh copies. The fire-test database is rebuilt only if the FP520 lecture deck sits one folder
above the harness; otherwise the copy in data/ is kept.

    python workspace/make_examples.py

The harness never changes these files: every edit goes to a copy in outputs/run_.../ (see
harness/tools/files.py). Run this again any time to get fresh copies.
"""

import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FOX = HERE.parent / "outputs" / "example_fox_image_and_narration" / "image_01.png"  # made by the image tool


def make_pptx(path):
    prs = Presentation()  # python-pptx's built-in template (4:3)
    title_layout, content_layout = prs.slide_layouts[0], prs.slide_layouts[1]
    s = prs.slides.add_slide(title_layout)
    s.shapes.title.text = "Teaching agent harness"
    s.placeholders[1].text = "Project update, October 2026"
    for title, bullets in [
            ("What the harness does", ["Runs a small LLM (Llama 3.2 3B) with media and document tools",
                                       "Prints every stage of the data flow, color-coded",
                                       "Writes an HTML report and trace.json for every run"]),
            ("How long videos are made", ["To be written"]),
            ("Next steps", ["Keep storyboard beats visual", "Add more requests to evaluate.py"])]:
        s = prs.slides.add_slide(content_layout)
        s.shapes.title.text = title
        body = s.placeholders[1].text_frame
        body.text = bullets[0]
        for b in bullets[1:]:
            body.add_paragraph().text = b
    prs.save(path)


def _font(size, bold=False):
    try:
        return ImageFont.truetype("arialbd.ttf" if bold else "arial.ttf", size)
    except OSError:
        return ImageFont.load_default(size)


def _png(image):
    stream = io.BytesIO()
    image.save(stream, "PNG")
    stream.seek(0)
    return stream


def fox_photo():
    """The fox the image tool painted (outputs/example_fox_...), cropped to 4:3 like a photo in an old deck."""
    if FOX.exists():
        image = Image.open(FOX).convert("RGB")
    else:  # no example run on disk: a stand-in with photo-like noise
        image = Image.merge("RGB", [Image.effect_noise((512, 512), 60).point(lambda v, o=o: min(255, v // 3 + o))
                                    for o in (70, 120, 170)])
    w, h = image.size
    return image.crop((0, (h - w * 3 // 4) // 2, w, (h + w * 3 // 4) // 2))


def time_chart():
    """A bar chart in flat colours, like most charts pasted into old decks (numbers from a real fox run)."""
    bars = [("LLM (Llama 3.2 3B)", 7.0, "#0b8fb3"), ("Image + speech models", 6.4, "#b03bb0"),
            ("Classifiers", 0.9, "#c99a06"), ("Perception models", 0.9, "#2f9e44"), ("Plain Python", 0.03, "#7a7f87")]
    image = Image.new("RGB", (1200, 720), "white")
    draw = ImageDraw.Draw(image)
    left, top, scale = 420, 50, 90  # pixels per second
    for i, (label, seconds, colour) in enumerate(bars):
        y = top + i * 125
        draw.text((left - 20, y + 40), label, fill="#1d1f23", font=_font(30), anchor="rm")
        draw.rectangle([left, y + 5, left + max(4, seconds * scale), y + 75], fill=colour)
        draw.text((left + seconds * scale + 16, y + 40), f"{seconds:g} s", fill="#1d1f23", font=_font(30, True),
                  anchor="lm")
    draw.line([left, top - 10, left, top + 5 * 125 - 30], fill="#1d1f23", width=3)
    draw.text((left, 700), "seconds per kind of component, one request: a fox picture + narration",
              fill="#6b7078", font=_font(24), anchor="ls")
    return image


def _bullets(text_frame, items):
    """items: (level, text). Sized down from the theme's 32 pt so the slide fits, as people do by hand."""
    for i, (level, text) in enumerate(items):
        p = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
        p.text, p.level = text, level
        p.runs[0].font.size = Pt(24 if level == 0 else 20)


def make_old_deck(path):
    """A 4:3 deck in the plain Office look: the input for convert_presentation (title, bullets, picture,
    section, table, chart, speaker notes)."""
    prs = Presentation()  # python-pptx's built-in template: 4:3, Office theme
    layouts = {layout.name: layout for layout in prs.slide_layouts}

    s = prs.slides.add_slide(layouts["Title Slide"])
    s.shapes.title.text = "Teaching agent harness"
    s.placeholders[1].text_frame.text = "How a small LLM and plain Python share the work"
    s.placeholders[1].text_frame.add_paragraph().text = "Lab meeting, October 2026"
    s.notes_slide.notes_text_frame.text = ("Welcome. Twenty minutes: what the harness does, one tool up close, "
                                           "and what the evaluation showed.")

    s = prs.slides.add_slide(layouts["Title and Content"])
    s.shapes.title.text = "What the harness does"
    _bullets(s.placeholders[1].text_frame, [
        (0, "Llama 3.2 3B plans each request and calls tools"),
        (1, "It only reads and writes text"),
        (1, "Tools make the images, sounds, videos and documents"),
        (0, "Classifiers screen the request and pick the tools"),
        (0, "Perception models turn every file back into text"),
        (1, "Captions for images, transcripts for speech, labels for music"),
        (0, "Every stage is timed and color-coded in the trace"),
        (0, "Each run writes report.html and trace.json")])
    s.notes_slide.notes_text_frame.text = "The point: most of the work is plain Python, not AI."

    s = prs.slides.add_slide(layouts["Two Content"])
    s.shapes.title.text = "One tool up close: how the harness makes pictures"
    _bullets(s.placeholders[1].text_frame, [(0, "SDXL-Turbo paints the picture in one step"),
                                            (0, "BLIP writes a caption for the LLM"),
                                            (0, "A safety classifier checks it first")])
    right = s.placeholders[2]  # the old way: a picture dropped on top of the right-hand box
    s.shapes.add_picture(_png(fox_photo()), right.left, right.top + Inches(0.6), width=right.width)
    right._element.getparent().remove(right._element)
    s.notes_slide.notes_text_frame.text = "This fox came from the image tool: watercolor fox in a snowy forest."

    s = prs.slides.add_slide(layouts["Section Header"])
    s.shapes.title.text = "Results"
    s.placeholders[1].text = "What the evaluation suite measured"

    s = prs.slides.add_slide(layouts["Title Only"])
    s.shapes.title.text = "Evaluation scores"
    rows = [("Change", "Score", "What it fixed"),
            ("First system prompt", "1/5", "nothing yet"),
            ("One tool call per thing asked", "3/5", "speech read out four times"),
            ("Router scores each clause", "6/6", "the robot's missing music"),
            ("Document tools added", "11/11", "slides, Word, LaTeX and code")]
    table = s.shapes.add_table(len(rows), 3, Inches(0.5), Inches(1.9), Inches(9), Inches(3)).table
    for col, width in enumerate((3.6, 1.4, 4.0)):
        table.columns[col].width = Inches(width)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    s.notes_slide.notes_text_frame.text = "Each change was measured with evaluate.py, not judged by eye."

    s = prs.slides.add_slide(layouts["Title Only"])
    s.shapes.title.text = "Where the time goes"
    s.shapes.add_picture(_png(time_chart()), Inches(0.5), Inches(1.8), width=Inches(9))
    s.notes_slide.notes_text_frame.text = "The LLM and the generators take almost all the time."
    prs.save(path)


TEAL, NAVY, PALE, WHITE = "0F4C5C", "1B2A4A", "DCEBF5", "FFFFFF"


def _text_colour(shape, hex_colour, levels=3):
    """The colour a placeholder on a layout or master gives its text: written into its list style (a:lstStyle),
    which is what slide text inherits. The layout's own prompt text ("Click to add title") is not inherited."""
    from lxml import etree
    from pptx.oxml.ns import qn
    body = shape._element.find(qn("p:txBody"))
    lst = body.find(qn("a:lstStyle"))
    if lst is None:
        lst = etree.Element(qn("a:lstStyle"))
        body.insert(1, lst)  # after a:bodyPr
    for level in range(1, levels + 1):
        ppr = lst.find(qn(f"a:lvl{level}pPr"))
        if ppr is None:
            ppr = etree.SubElement(lst, qn(f"a:lvl{level}pPr"))
        rpr = ppr.find(qn("a:defRPr"))
        if rpr is None:
            rpr = etree.SubElement(ppr, qn("a:defRPr"))
        for old in rpr.findall(qn("a:solidFill")):
            rpr.remove(old)
        fill = etree.Element(qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr")).set("val", hex_colour)
        rpr.insert(0, fill)


def _copy_layout(prs, layout, name):
    """A second copy of a layout under a new name: python-pptx has no call for this, so copy the layout's XML
    into a new part and register it with the slide master."""
    from copy import deepcopy
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.packuri import PackURI
    from pptx.oxml.ns import qn
    from pptx.parts.slide import SlideLayoutPart
    master = prs.slide_master
    ids = master._element.sldLayoutIdLst
    next_id = max(int(e.get("id")) for e in ids) + 1
    part = SlideLayoutPart(PackURI(f"/ppt/slideLayouts/slideLayout{len(master.slide_layouts) + 1}.xml"),
                           layout.part.content_type, layout.part.package, deepcopy(layout.part._element))
    part.relate_to(master.part, RT.SLIDE_MASTER)
    entry = ids._add_sldLayoutId()
    entry.set("id", str(next_id))
    entry.set(qn("r:id"), master.part.relate_to(part, RT.SLIDE_LAYOUT))
    copy = master.slide_layouts[len(master.slide_layouts) - 1]
    copy._element.cSld.set("name", name)
    return copy


def make_workshop_template(path):
    """
    templates/workshop_template.pptx: a plain 16:9 template for convert_presentation, built from python-pptx's
    own template. Teal titles, a teal title slide, a dark section divider ("Section Header") and a light one
    ("Section Header Light"), so a converted deck looks visibly different from old_format_deck.pptx. The layout
    names are the standard ones, which convert_presentation's DEFAULT_LAYOUTS fall back to. Put your own
    template (.pptx or .potx) next to it and name it in a request: "Convert ... to the <name> template".
    """
    from pptx.dml.color import RGBColor
    prs = Presentation()  # 4:3, 10 x 7.5 in
    widen = 13.333 / 10
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9: stretch everything sideways
    master = prs.slide_master
    for owner in [master] + list(master.slide_layouts):
        for shape in owner.shapes:
            if shape._element.xpath("./p:spPr/a:xfrm"):  # a position of its own (others inherit the master's)
                shape.left, shape.width = int(shape.left * widen), int(shape.width * widen)
    for ph in master.placeholders:
        if ph.placeholder_format.type == 1:  # TITLE: every slide's title in teal
            _text_colour(ph, TEAL, levels=1)
    layouts = {layout.name: layout for layout in master.slide_layouts}
    light = _copy_layout(prs, layouts["Section Header"], "Section Header Light")
    for layout, background, text in [(layouts["Title Slide"], TEAL, WHITE), (layouts["Section Header"], NAVY, WHITE),
                                     (light, PALE, NAVY)]:
        layout.background.fill.solid()
        layout.background.fill.fore_color.rgb = RGBColor.from_string(background)
        for ph in layout.placeholders:
            _text_colour(ph, text)
    path.parent.mkdir(exist_ok=True)
    prs.save(path)


def make_docx(path):
    doc = Document()
    doc.add_heading("Teaching harness: progress report", level=0)
    for heading, text in [
            ("Summary", "This report describes a small agent harness built for teaching. It runs open models on "
                        "one laptop and records every stage of each request."),
            ("Methods", "Requests are routed by a zero-shot classifier, planned by Llama 3.2 3B and carried out by "
                        "tools. Perception models describe every file the tools make."),
            ("Results", "To be written."),
            ("Next steps", "Add evaluation requests for the document tools.")]:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(text)
    doc.save(path)


EVALUATION_TEX = r"""
\documentclass[11pt]{article}
\usepackage[margin=1in]{geometry}
\begin{document}
\section*{Evaluation results}
The script evaluate.py runs a fixed set of requests and checks which tools the LLM called.
Decoding is greedy, so the same harness gives the same score on every run, and a change in
score comes from a change in the harness.

\begin{tabular}{ll}
Change & Requests passed \\ \hline
First system prompt & 1 of 5 \\
Rule: one tool call per thing asked for & 3 of 5 \\
Hint about requests that need several tools (a sixth request added) & 3 of 6 \\
Score only the tool calls that actually ran & 5 of 6 \\
Router scores each clause of the request & 6 of 6 \\
\end{tabular}

\medskip
Each edit to the system prompt fixed one request and broke another. The final fix was in
the router, not in the prompt. After the completion check was added for the 40-second
narrated video, the score stayed at 6 of 6.
\end{document}
"""


def make_pdf(path):
    with tempfile.TemporaryDirectory() as tmp:
        tex = Path(tmp) / "evaluation_results.tex"
        tex.write_text(EVALUATION_TEX, encoding="utf-8")
        subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", tex.name], cwd=tmp,
                       check=True, capture_output=True, timeout=120)
        shutil.copy(Path(tmp) / "evaluation_results.pdf", path)


T_SQUARED = {"slow": 0.00293, "medium": 0.01172, "fast": 0.0469, "ultrafast": 0.1876}  # alpha, kW/s^2 (NFPA 72)


def make_hrr_plot(path, peak_kw=3000, end_s=600):
    """A plot with KNOWN numbers for digitize_plot: t-squared design fires, HRR = alpha * t^2, capped at 3 MW.
    The exact curves go into <name>_truth.csv next to it, so the digitized values can be checked."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    t = np.linspace(0, end_s, 601)
    curves = {name: np.minimum(alpha * t ** 2, peak_kw) for name, alpha in T_SQUARED.items()}
    fig, ax = plt.subplots(figsize=(8, 5), dpi=120)
    for name, hrr in curves.items():
        ax.plot(t, hrr, linewidth=2, label=name.capitalize())
    ax.set(xlim=(0, end_s), ylim=(0, 3500), xlabel="Time (s)", ylabel="Heat release rate (kW)",
           title="Design fires: t-squared growth to 3 MW")
    ax.grid(color="#dddddd")
    ax.legend(loc="center right")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)
    rows = ["series,time_s,hrr_kw"] + [f"{name.capitalize()},{x:g},{y:.1f}" for name, hrr in curves.items()
                                       for x, y in zip(t[::10], hrr[::10])]
    Path(path).with_name(Path(path).stem + "_truth.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")


LECTURE = HERE.parent.parent / "FP520_Introduction_Design_Fires.pptx"  # the source of the fire-test tables
# Each database table: (name, slide, the slide table's first header cell, columns as name + SQL type, what it holds,
# where the slide says the numbers come from). Rows are read from the lecture, never typed in here. Column names
# carry the units and the words people use ("measured" for the slides' "Exp", "predicted" for "Model"), so a
# small LLM can match them.
FIRE_TABLES = [
    ("vehicle_fires", 35, "Vehicles",
     "vehicle TEXT, measured_hrr_min_mw REAL, measured_hrr_max_mw REAL, representative_hrr_mw REAL, "
     "with_water_system_hrr_min_mw REAL, with_water_system_hrr_max_mw REAL",
     "Peak heat release rates of vehicle fires in road tunnels: measured ranges, the representative design value, "
     "and ranges with fixed water-based firefighting systems. NULL = no value on the slide.",
     "NFPA 502 representative HRRs; most test data from the EUREKA test series (Eureka-Project EU 499, 1995)"),
    ("tree_fires", 44, "Case",
     "tree TEXT, mass_kg REAL, moisture_pct REAL, predicted_peak_hrr_mw REAL, measured_peak_hrr_mw REAL",
     "Christmas-tree burns: peak heat release rate predicted by an empirical model vs. measured.",
     "NIST Fire Calorimetry Database, design fires: tall trees"),
    ("mining_vehicle_fires", 50, "Combustible",
     "vehicle TEXT, peak_hrr_kw REAL, diesel_fuel_l REAL, hydraulic_oil_l REAL, tires_kg REAL, hoses_kg REAL",
     "Fuel loads and peak heat release rates of a wheel loader, a drilling rig and a forklift.",
     "Hansen's model for mining vehicles, recreated in the lecture, fuel loads adjusted for the forklift"),
    ("luggage_cart_items", 54, "Item",
     "item TEXT, hrr_kw REAL, mass_kg REAL, hrr_per_mass_kw_per_kg REAL",
     "Peak heat release rate and mass of the items in a luggage-cart fire estimate (HRR scaled by mass).",
     "luggage data from the SFPE Handbook; forklift from Audouin and Coutin, SMiRT 24 (2017); the ATV assumed to be "
     "like a lightly loaded forklift"),
    ("sprinkler_activation", 26, "Software",
     "method TEXT, activation_time_min REAL, hrr_at_activation_kw REAL, ceiling_height_m REAL, spacing_m REAL, "
     "rti_m05_s05 REAL, activation_temp_c REAL, ambient_temp_c REAL, alpha_kw_s2 REAL",
     "One sprinkler problem (fast t-squared fire, inputs in the other columns) solved with four tools.",
     "the lecture's comparison of DETACT-T2 (NIST), two online DETACT calculators and a finite-difference solution "
     "with Alpert's ceiling-jet correlations"),
]


def _cell(text):
    """'13,000' -> 13000.0, '5-10' -> (5.0, 10.0), '-' -> None, 'Fast, 0.04689' -> 0.04689, 'Bus' -> 'Bus'."""
    text = " ".join(text.split())
    numbers = [float(n.replace(",", "")) for n in re.findall(r"\d[\d,]*\.?\d*", text)]
    if text in ("-", ""):
        return None
    if re.fullmatch(r"[\d,.]+\s*-\s*[\d,.]+", text):
        return tuple(numbers)
    if re.fullmatch(r"[\d,.]+", text) or text.lower().startswith(("fast,", "slow,", "medium,")):
        return numbers[-1]
    return text


def _rows(table, columns):
    """A slide table -> database rows. A slide cell that holds a range ("5-10") fills a _min_/_max_ pair of columns
    (a single number fills both); the forklift table is turned on its side (its vehicles are the slide's columns)."""
    cells = [[_cell(c.text) for c in row.cells] for row in table.rows]
    header, body = cells[0], cells[1:]
    if header[0] == "Combustible Fuel":
        return [[name] + [row[i] for row in body] for i, name in enumerate(header[1:], 1)]
    names = [c.split()[0] for c in columns.split(",")]
    ranges = [name.replace("_max_", "_min_") for name in names if "_max_" in name]
    rows = []
    for row in body:
        flat = []
        for value in row:
            if names[len(flat)] in ranges:
                value = value if isinstance(value, tuple) else (value, value)
            flat += list(value) if isinstance(value, tuple) else [value]
        rows.append(flat)
    return rows


def make_fire_tests_db(path, deck=LECTURE):
    """
    data/fire_tests.sqlite for query_database: the tables of the FP520 design-fires lecture, read from the deck with
    python-pptx. Every table is listed in `sources` with its slide and the reference the slide gives.
    """
    import sqlite3
    if not deck.exists():
        print(f"no {deck.name} one folder above the harness: " + (f"kept the existing {path.name}" if path.exists()
                                                                   else f"{path.name} not built (query_database needs it)"))
        return
    slides = list(Presentation(deck).slides)
    path.parent.mkdir(exist_ok=True)
    path.unlink(missing_ok=True)
    with sqlite3.connect(path) as db:
        db.execute("CREATE TABLE sources (table_name TEXT, slide INTEGER, slide_title TEXT, about TEXT, source TEXT)")
        for name, number, first_header, columns, about, source in FIRE_TABLES:
            slide = slides[number - 1]
            tables = [s.table for s in slide.shapes if s.has_table]
            table = next(t for t in tables if " ".join(t.cell(0, 0).text.split()).startswith(first_header))
            rows = _rows(table, columns)
            if name == "sprinkler_activation":  # the inputs are a second table on the same slide: repeat them per row
                inputs = next(t for t in tables if t.cell(0, 0).text.strip() == "Parameter")
                given = [_cell(row.cells[1].text) for row in list(inputs.rows)[1:]]  # H, spacing, RTI, T_act, T_amb, alpha
                rows = [row + given for row in rows]
            db.execute(f"CREATE TABLE {name} ({columns})")
            db.executemany(f"INSERT INTO {name} VALUES ({', '.join('?' * len(rows[0]))})", rows)
            db.execute("INSERT INTO sources VALUES (?, ?, ?, ?, ?)",
                       (name, number, slide.shapes.title.text_frame.text, about,
                        f"FP520 lecture (Introduction to Design Fires), slide {number}: {source}"))


def make_workbook(path):
    """fire_calcs.xlsx for update_excel: a time column and one parameter, alpha for a fast t-squared fire (NFPA 72),
    and empty columns for the LLM to fill with formulas. The time column's header is t and E1 is named alpha, so a
    formula can be written as the physics is: =alpha*t^2 (update_excel turns t into A2, A3, ...)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.workbook.defined_name import DefinedName
    book = Workbook()
    ws = book.active
    ws.title = "Design fire"
    ws["A1"], ws["D1"], ws["E1"] = "t", "alpha (kW/s^2), fast t-squared fire", T_SQUARED["fast"]
    ws["D2"] = "t in s; HRR = alpha * t^2 in kW"
    book.defined_names["alpha"] = DefinedName("alpha", attr_text="'Design fire'!$E$1")
    for row, seconds in enumerate(range(0, 601, 30), start=2):
        ws.cell(row=row, column=1, value=seconds)
    for cell in ("A1", "D1"):
        ws[cell].font = Font(bold=True)
    ws.column_dimensions["D"].width = 42
    book.save(path)


if __name__ == "__main__":
    make_fire_tests_db(HERE / "data" / "fire_tests.sqlite")
    make_workbook(HERE / "fire_calcs.xlsx")
    make_hrr_plot(HERE / "images" / "hrr_curves.png")
    make_pptx(HERE / "project_update.pptx")
    make_old_deck(HERE / "old_format_deck.pptx")
    make_workshop_template(HERE / "templates" / "workshop_template.pptx")
    make_docx(HERE / "report.docx")
    make_pdf(HERE / "references" / "evaluation_results.pdf")
    print("wrote fire_calcs.xlsx, images/hrr_curves.png (+ _truth.csv), project_update.pptx, "
          "old_format_deck.pptx, templates/workshop_template.pptx, report.docx, references/evaluation_results.pdf")
