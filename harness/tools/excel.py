"""
update_excel: the LLM writes FORMULAS into a spreadsheet, Excel calculates them, the harness reads the results back.

The LLM writes formulas with names, the way Excel tables do ([@time]); plain Python turns the names into addresses:

    columns: {"HRR": "=alpha*t^2"}   -> B1 = HRR, then B2 = alpha*A2^2, B3 = alpha*A3^2, ... B22, one per row of the
                                        table (t is column A's header; alpha is a named cell)
    cells:   {"E3": "=SUM(HRR)"}     -> E3 = SUM(B2:B22): in a single cell, a header means the whole column

Llama 3.2 3B writes "=alpha*t^2" happily; asked for addresses, it wrote E1 where $E$1 was needed (filled down, A2^2*E1
became A3^2*E2, A4^2*E3, ...), ran ranges past the data and copied B2 into B3. So the addressing moved into code.

  CHECK FORMULAS  names -> addresses, one formula per cell; refuses a formula that reaches outside the workbook
  WRITE CELLS     openpyxl (plain Python) writes into a working copy in the run folder
  RECALCULATE     openpyxl can't calculate anything: it stores "=alpha*A2^2" and leaves the cell's cached value empty.
                  Excel, started as a COM server (another program on this PC, with macros off), opens the copy,
                  recalculates every formula and saves it
  READ BACK       openpyxl reads the values Excel calculated for the changed cells. An error such as #NAME? goes back to
                  the LLM, with what it means, so it can fix the formula

The LLM writes formulas, never numbers it worked out itself, so the sheet stays live and every number in it can be
traced to a formula a person can read.
"""

import ast
import json
import os
import re
from contextlib import contextmanager
from pathlib import Path

from .base import Refusal, Tool
from .files import current_version, rel, resolve, working_copy

ADDRESS = re.compile(r"^([A-Z]{1,3})([1-9]\d{0,6})(?::([A-Z]{1,3})([1-9]\d{0,6}))?$")
HEADER = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")  # a header that can be used as a name in formulas
MAX_CELLS = 500   # one call writes at most this many cells
ERRORS = {  # what Excel shows in a cell it couldn't calculate -> what that means, for the LLM
    "#NAME?": "Excel does not know a name in the formula: use only the column headers, the named cells, Excel "
              "functions (SUM, SQRT, EXP, ...) and numbers",
    "#DIV/0!": "a division by zero or by an empty cell", "#REF!": "a reference to a cell that does not exist",
    "#VALUE!": "text where a number was expected", "#N/A": "a lookup found nothing", "#NUM!": "an impossible number",
    "#NULL!": "two ranges that do not meet", "#SPILL!": "a result that does not fit", "#CALC!": "a calculation error"}
# A formula may not reach outside the workbook: no web requests, no other files, no macro or DDE calls.
OUTSIDE = re.compile(r"WEBSERVICE|FILTERXML|IMAGE\(|\bCALL\(|REGISTER|\bEXEC\(|\bRTD\(|\[|\||https?:", re.I)


def table(ws):
    """Row 1's headers that work as names ({"t": "A"}), and the table's rows: 2 to the last row with a value in A."""
    headers = {str(c.value).strip(): c.column_letter for c in ws[1] if isinstance(c.value, str)
               and HEADER.match(c.value.strip()) and not ADDRESS.match(c.value.strip().upper())}
    last = max((c.row for c in ws["A"] if c.value not in (None, "")), default=1)
    return headers, range(2, max(last, 2) + 1)


def by_name(formula, headers, row=None, rows=None):
    """
    A formula written with column headers -> Excel addresses. In a column formula (row given) a header means the
    cell in that row: '=alpha*t^2' -> '=alpha*A2^2'. In a single cell it means the whole column: '=SUM(HRR)' ->
    '=SUM(B2:B22)'. Named cells (alpha), functions and plain addresses pass through unchanged.
    """
    if not (isinstance(formula, str) and formula.startswith("=") and headers):
        return formula
    lower = {name.lower(): col for name, col in headers.items()}
    names = "|".join(sorted(map(re.escape, headers), key=len, reverse=True))
    pattern = re.compile(rf"(?<![\w$.:!'\"])({names})(?![\w(!])", re.I)  # a whole word, not a function or a sheet
    col = lambda m: lower[m.group(1).lower()]
    if row is not None:
        return pattern.sub(lambda m: f"{col(m)}{row}", formula)
    return pattern.sub(lambda m: f"{col(m)}{rows.start}:{col(m)}{rows.stop - 1}", formula)


def expand(cells):
    """{"B2:B4": "=$E$1*A2^2", "D1": 5} -> [("B2:B4", "B2", "=$E$1*A2^2"), ("B2:B4", "B3", "=$E$1*A3^2"), ...,
    ("D1", "D1", 5)]: a range gets its formula filled down (or across) as Excel does it. openpyxl's Translator moves
    the relative references (A2 -> A3), and leaves absolute ones ($E$1) alone."""
    from openpyxl.formula.translate import Translator
    from openpyxl.utils import column_index_from_string, get_column_letter
    out = []
    for key, value in cells.items():
        match = ADDRESS.match(key.replace("$", ""))
        if not match:
            raise ValueError(f"'{key}' is not a cell such as E3. To fill a column, use columns with its header")
        col1, row1, col2, row2 = match.groups()
        if col2 and (column_index_from_string(col2) < column_index_from_string(col1) or int(row2) < int(row1)):
            raise ValueError(f"'{key}' runs backwards: a range goes from its first cell to its last")
        if not col2:
            out.append((key, f"{col1}{row1}", value))
            continue
        for col in range(column_index_from_string(col1), column_index_from_string(col2) + 1):
            for row in range(int(row1), int(row2) + 1):
                address = f"{get_column_letter(col)}{row}"
                formula = isinstance(value, str) and value.startswith("=")
                out.append((key, address, Translator(value, origin=f"{col1}{row1}").translate_formula(address)
                            if formula else value))
                if len(out) > MAX_CELLS:
                    raise ValueError(f"at most {MAX_CELLS} cells per call")
    return out


def as_cell_value(value):
    """'42' -> 42.0 (a number, so formulas can use it); '=A2*2' and 'hrr_kW' stay text."""
    if isinstance(value, str) and not value.startswith("="):
        try:
            return float(value.replace(",", ""))
        except ValueError:
            return value
    return value


def as_object(value):
    """A small LLM sometimes sends a JSON object as a string, or with single quotes."""
    if isinstance(value, str):
        for parse in (json.loads, ast.literal_eval):
            try:
                return parse(value)
            except (ValueError, SyntaxError):
                pass
    return value


def excel_available():
    """Windows, pywin32 installed, and Excel registered as a COM server."""
    if os.name != "nt":
        return False
    try:
        import win32com.client  # noqa: F401  (pywin32)
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "Excel.Application"))
        return True
    except (ImportError, OSError):
        return False


@contextmanager
def excel(started):
    """A NEW, hidden Excel process (DispatchEx), so a workbook you have open is never touched; closed afterwards.
    Its process id goes into `started`, so the harness can stop it if it hangs."""
    import pythoncom
    import win32com.client
    import win32process
    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("Excel.Application")
    started.append(win32process.GetWindowThreadProcessId(app.Hwnd)[1])
    app.Visible, app.DisplayAlerts = False, False
    app.AutomationSecurity = 3  # msoAutomationSecurityForceDisable: no macros run in files opened from code
    try:
        yield app
    finally:
        app.Quit()
        del app
        pythoncom.CoUninitialize()


def recalculate(path, timeout=60):
    """
    Excel opens the file, recalculates every formula in every sheet, and saves it. A hidden Excel that shows a dialog
    (the file is in use, a repair prompt) waits for a click that never comes, so the work runs in a thread, and an
    Excel still busy after `timeout` seconds is stopped: the harness never waits forever on another program.
    """
    import subprocess
    import threading
    started, failed = [], []

    def work():
        try:
            with excel(started) as app:
                book = app.Workbooks.Open(str(Path(path).resolve()), UpdateLinks=0, Notify=False)
                try:
                    app.CalculateFull()
                    book.Save()
                finally:
                    book.Close(SaveChanges=False)
        except Exception as exc:  # reported in the main thread
            failed.append(exc)

    worker = threading.Thread(target=work, daemon=True)
    worker.start()
    worker.join(timeout)
    if worker.is_alive():
        for pid in started:
            subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True)
        raise TimeoutError(f"Excel did not finish within {timeout} s (probably a dialog it could not show), so it was "
                           "stopped. The formulas are saved in the copy but not calculated.")
    if failed:
        raise failed[0]


def sample(plan, values):
    """The cells to show, group by group: a long column as its first three and last two cells and a marker such as
    "B5:B20": "...", because the context window is not free. plan: [(column or cell as written, address, value)]."""
    shown, groups = {}, {}
    for key, address, _ in plan:
        groups.setdefault(key, []).append(address)
    for addresses in groups.values():
        if len(addresses) <= 5:
            shown.update({a: values[a] for a in addresses})
        else:
            shown.update({a: values[a] for a in addresses[:3]})
            shown[f"{addresses[3]}:{addresses[-3]}"] = "..."
            shown.update({a: values[a] for a in addresses[-2:]})
    return shown


# =========================================================================================
class UpdateExcelTool(Tool):
    name = "update_excel"
    description = ("Add columns and cells with formulas to an Excel workbook (.xlsx) in the workspace; Excel calculates "
                   "them and the tool returns the values. Write formulas with the column headers of row 1 and the named "
                   "cells, not with cell addresses. columns: the header of each new column -> the formula for one row, "
                   "e.g. {\"mass_g\": \"=mass_kg*1000\"}; the tool writes it into every row of the table. cells: a "
                   "cell or a name for a result -> a value or a formula, where a header means the whole column, e.g. "
                   "{\"total_g\": \"=SUM(mass_g)\"}. Read the file first. Write formulas, never numbers you calculated "
                   "yourself. The edited workbook is saved as a copy in the run folder.")
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "The .xlsx file, e.g. workspace/fire_calcs.xlsx"},
            "sheet": {"type": "string", "description": "The sheet's name. Leave it out for the first sheet."},
            "columns": {"type": "object", "description": "New column header -> formula for one row, using headers "
                                                         "and named cells"},
            "cells": {"type": "object", "description": "Cell address, or a name for a result -> value or formula; "
                                                       "a header in a formula means its whole column"},
        },
        "required": ["path"],
    }
    model_id = "openpyxl writes and reads, Excel calculates (COM)"
    check_unused = False
    requires = ("read_file",)  # read the sheet first, to learn its headers and named cells
    how_it_works = ("No model here: openpyxl, a Python library, writes the cells into the file's XML. It can store a "
                    "formula but cannot calculate it, so Excel itself opens the copy, calculates and saves.")

    def validate(self, args):
        args = super().validate(args)
        columns, cells = as_object(args.get("columns") or {}), as_object(args.get("cells") or {})
        if isinstance(cells, list):  # [{"cell": "E3", "value": "=SUM(HRR)"}, ...]
            cells = {str(c.get("cell") or c.get("address")): c.get("value", c.get("formula"))
                     for c in cells if isinstance(c, dict)}
        if not isinstance(columns, dict) or not isinstance(cells, dict) or not (columns or cells):
            raise ValueError('give columns, cells or both, e.g. columns {"mass_g": "=mass_kg*1000"} and cells '
                             '{"total_g": "=SUM(mass_g)"}')
        args["columns"] = {str(k).strip(): v for k, v in columns.items()}
        # an address is written in capitals (e3 -> E3); a name such as total_energy keeps its spelling
        args["cells"] = {(str(k).strip().upper() if ADDRESS.match(str(k).strip().upper().replace("$", ""))
                          else str(k).strip()): v for k, v in cells.items()}
        return args

    def run(self, ctx, path, columns=None, cells=None, sheet=""):
        import openpyxl
        from openpyxl.utils import get_column_letter
        t, columns, cells = ctx.tracer, columns or {}, cells or {}
        source = current_version(resolve(path, ctx), ctx)
        if source.suffix.lower() != ".xlsx":
            raise ValueError(f"{rel(source)} is not an Excel workbook (.xlsx)")
        book = openpyxl.load_workbook(source)
        names = {name.lower(): name for name in book.sheetnames}
        if sheet and sheet.lower() not in names and len(book.sheetnames) > 1:
            raise ValueError(f"There is no sheet called '{sheet}'. The sheets are: {', '.join(book.sheetnames)}")
        if sheet and sheet.lower() not in names:  # one sheet only: no doubt which one was meant ("Sheet1")
            t.say(f"there is no sheet '{sheet}': using the only sheet, '{book.sheetnames[0]}'", kind="python")
        title = names[sheet.lower()] if sheet and sheet.lower() in names else book.sheetnames[0]

        with t.stage("CHECK FORMULAS", "python", "excel.by_name() + a regex (no ML)", "formulas with headers",
                     "one formula per cell, with addresses") as ev:
            t.say(self.how_it_works, kind="python")
            ws = book[title]
            headers, rows = table(ws)
            used = {c.column for row in ws.iter_rows() for c in row if c.value not in (None, "")}
            for header in columns:  # a new column goes into the first empty column
                if header.lower() not in {h.lower() for h in headers}:
                    free = next(i for i in range(1, 200) if i not in used)
                    used.add(free)
                    headers[header] = get_column_letter(free)
            plan = []
            for header, formula in columns.items():
                col = next(c for h, c in headers.items() if h.lower() == header.lower())
                plan.append((header, f"{col}1", header))
                plan += [(header, f"{col}{row}", by_name(formula, headers, row=row)) for row in rows]
                t.say(f"column {header}: {formula}  ->  {col}{rows.start}: {by_name(formula, headers, row=rows.start)}"
                      f"  ...  {col}{rows.stop - 1}", kind="python")
            for key in [k for k in cells if HEADER.match(k) and not ADDRESS.match(k.upper())]:
                # a named result ({"total_energy": "=SUM(HRR)"}): its name in row 1 of the next empty column, the
                # value under it. Llama sends names as often as addresses, and a label makes the sheet readable
                col = next((c for h, c in headers.items() if h.lower() == key.lower()), None)  # written before?
                if col is None:
                    free = next(i for i in range(1, 200) if i not in used)
                    used.add(free)
                    col = get_column_letter(free)
                plan.append((key, f"{col}1", key))
                cells[f"{col}2"] = cells.pop(key)
                t.say(f"named result {key}: written as {col}1 = {key}, {col}2 = {cells[f'{col}2']}", kind="python")
            for key, address, value in expand({k: by_name(v, headers, rows=rows) for k, v in cells.items()}):
                plan.append((key, address, value))
                if value != cells.get(key):
                    t.say(f"cell {address}: {cells.get(key)}  ->  {value}", kind="python")
            if len(plan) > MAX_CELLS:
                raise ValueError(f"at most {MAX_CELLS} cells per call")
            outside = [f"{a}: {v}" for _, a, v in plan if isinstance(v, str) and v.startswith("=") and OUTSIDE.search(v)]
            ev.summary = (f"REFUSED: {outside[0]} reaches outside the workbook" if outside else
                          f"headers {', '.join(f'{h}={c}' for h, c in headers.items())}, table rows {rows.start}-"
                          f"{rows.stop - 1} -> {len(plan)} cells; nothing reaches outside the workbook")
        if outside:  # raised outside the stage, so the trace shows the reason rather than "FAILED"
            raise Refusal(f"Refused by the harness: {outside[0]} would reach outside the workbook (a web address, "
                          "another file or a macro call). Use formulas on the cells of this workbook only.")

        with t.stage("WRITE CELLS", "python", "openpyxl (no ML)", "one formula or value per cell",
                     ".xlsx working copy") as ev:
            copy = working_copy(source, ctx)
            book = openpyxl.load_workbook(copy)
            ws = book[title]
            replaced = []
            for _, address, value in plan:
                old = ws[address].value
                if old not in (None, "") and str(old) != str(value):
                    replaced.append(f"{address}: {old!r} -> {value!r}")
                ws[address] = as_cell_value(value)
            book.save(copy)
            formulas = [a for _, a, v in plan if isinstance(v, str) and v.startswith("=")]
            written = {a: v for _, a, v in plan}
            for address, value in sample(plan, written).items():
                t.say(f"{address} = {value}", kind="python")
            for line in replaced:
                t.say(f"replaced {line}", kind="python")
            ev.summary = (f"{len(plan)} cell(s) in sheet '{title}' of {rel(copy)}: {len(formulas)} formula(s), "
                          f"{len(plan) - len(formulas)} value(s)" + (f"; {len(replaced)} replaced" if replaced else ""))
            ev.details = {"cells": written, "replaced": replaced}

        with t.stage("RECALCULATE", "server", "Excel as a COM server (another program on this PC)",
                     ".xlsx: formulas, no values", ".xlsx with calculated values") as ev:
            stored = openpyxl.load_workbook(copy, data_only=True)[title]
            empty = [a for a in formulas if stored[a].value is None]
            t.say(f"before: openpyxl stored {len(formulas)} formula(s) and calculated nothing: {len(empty)} cached "
                  "value(s) are empty", kind="server")
            if not formulas:
                ev.summary = "no formulas written: nothing to calculate"
            elif excel_available():
                recalculate(copy)
                ev.summary = f"Excel opened {rel(copy)}, recalculated every formula and saved it"
            else:
                ev.summary = "Excel is not installed: the formulas are saved, and Excel calculates them when the file opens"

        with t.stage("READ BACK", "python", "openpyxl, data_only=True (no ML)", ".xlsx",
                     "the changed cells' values") as ev:
            calculated = openpyxl.load_workbook(copy, data_only=True)[title]
            values = {a: calculated[a].value for _, a, _ in plan}
            values = {a: round(v, 6) if isinstance(v, float) else v for a, v in values.items()}
            errors = {a: v for a, v in values.items() if isinstance(v, str) and v in ERRORS}
            shown = sample(plan, values)
            for address, value in shown.items():
                t.say(f"{address} -> {value}", kind="python")
            ev.summary = (f"{len(values)} value(s) read back" + (f"; ERRORS in {', '.join(errors)}: "
                          f"{', '.join(sorted(set(errors.values())))}" if errors else "; no errors"))
        if errors:  # the agent's completion check asks the LLM once whether it wants to fix them
            ctx.pending[self.name] = (f"Excel could not calculate cell(s) {', '.join(list(errors)[:6])}: "
                                      + "; ".join(f"{e} means {ERRORS[e]}" for e in sorted(set(errors.values())))
                                      + ". Fix the formulas with update_excel")
        else:
            ctx.pending.pop(self.name, None)
        result = {"file": rel(copy), "sheet": title,
                  "formulas": {a: f for a, f in sample(plan, written).items() if str(f).startswith("=")},
                  "values": shown}
        if errors:
            result["errors"] = {e: f"{ERRORS[e]} (in {', '.join(a for a, v in errors.items() if v == e)[:80]})"
                                for e in sorted(set(errors.values()))}
            result["note"] = "Excel could not calculate these cells. Fix their formulas and call update_excel again."
        if replaced:
            result["replaced"] = replaced[:10]
        return {"artifacts": [{"type": "workbook", "path": copy, "note": f"edited copy of {rel(source)}"}],
                "observations": result}
