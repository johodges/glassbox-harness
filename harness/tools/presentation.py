"""
Presentation tools: move a deck into another template, and save a deck as PDF.

    convert_presentation   python-pptx reads every slide; plain-Python rules decide what kind of slide it
                           is (title, section, bullets, picture + text, table, picture); each kind goes to the
                           template layout made for it, and the text, pictures, tables and speaker notes are
                           put into that layout's placeholders. No model looks at a slide.
    export_pdf             PowerPoint (through COM, Windows) or LibreOffice draws the slides into a PDF.

The LLM's part is small on purpose: it names the two files, and it may pick a layout for a kind of slide
if the user asked for one. Llama 3.2 3B can't see slides and writes at most 400 tokens per call, so it could
never plan a whole deck. Checking the result is perception's job (READ BACK and RENDER in perception.py).
"""

import ast
import difflib
import io
import json
import re
import zipfile
from collections import Counter
from copy import deepcopy
from pathlib import Path

from PIL import Image

from ..config import WORKSPACE_DIR
from .base import Tool
from .deck_check import bulleted, save_as_pdf, style_value, text_fit
from .files import current_version, named_paths, rel, resolve, working_copy

TEMPLATES_DIR = WORKSPACE_DIR / "templates"
SLIDE_TYPES = ("title", "section", "bullets", "picture_text", "table", "picture", "drawn", "blank")

# Where each kind of slide goes: the first layout name the template has wins. The first names are
# webinar organisation's template's; the others are PowerPoint's built-in names, which workshop_template.pptx uses. A template
# with none of these names gets the first layout whose placeholders fit (see fits()).
DEFAULT_LAYOUTS = {
    "title": ["Cover 03", "Title Slide"],
    "section": ["Divider Dark Blue", "Section Header"],
    "bullets": ["One Column", "Title and Content"],
    "picture_text": ["Content Image", "Picture with Caption", "Two Content"],
    "table": ["One Column", "Title and Content", "Title Only"],
    "picture": ["Title Only White", "Title Only"],
    # a slide too busy for one placeholder per thing (videos, several pictures, arrows, text boxes...): the
    # template's title and text box, and everything else copied as drawn (see fill_drawn)
    "drawn": ["One Column", "Title and Content"],
    "blank": ["Blank White", "Blank"],
}
WORDS_BESIDE_A_PICTURE = 40  # more text than this next to a picture would be shrunk hard in a half-slide text box
# Placeholder types (python-pptx's PP_PLACEHOLDER numbers) -> what they hold.
KINDS = {1: "title", 3: "title", 4: "subtitle", 2: "text", 7: "text", 12: "text", 18: "picture"}
SLIDE_NUMBER = 13
CROP_ANY, CROP_PHOTO = 0.08, 0.35  # how much of a picture may be cut off to fill a box: any picture / a photo
FONT_SCALES = (0.925, 0.85, 0.775, 0.7, 0.625, 0.55)  # PowerPoint's own "shrink text on overflow" steps
TOO_FULL, TARGET = 0.98, 0.95  # the estimate is close but not exact (PowerPoint measured 102% for an estimated 100%)


def _norm(name):
    return " ".join(str(name).lower().split())


def shorten_list(items, keep=6):
    """A 3B model loses track in long tool results: show a few items and count the rest (the trace has all)."""
    return items if len(items) <= keep else items[:keep] + [f"... and {len(items) - keep} more (see the trace)"]


def open_deck(path):
    """python-pptx opens .pptx files only. A .potx template is the same ZIP with one different content
    type, so that one line is swapped in memory (the file itself is not touched)."""
    from pptx import Presentation
    path = Path(path)
    if path.suffix.lower() != ".potx":
        return Presentation(path)
    buffer = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
            zout.writestr(item, data)
    buffer.seek(0)
    return Presentation(buffer)


def find_template(text, ctx):
    """The LLM may write a path, a file name, or just 'the workshop template'. Plain Python finds the file."""
    try:
        return resolve(text, ctx)
    except FileNotFoundError:
        pass
    templates = sorted(p for p in WORKSPACE_DIR.rglob("*") if p.suffix.lower() in (".pptx", ".potx")
                       and (p.is_relative_to(TEMPLATES_DIR) or p.suffix.lower() == ".potx"))
    words = [w for w in re.findall(r"[a-z0-9]+", text.lower())
             if w not in ("the", "a", "template", "templates", "workspace", "pptx", "potx")]
    matches = [p for p in templates if all(w in p.stem.lower() for w in words)]
    if len(matches) == 1:
        ctx.tracer.say(f"template \"{text}\" -> {rel(matches[0])} (the only template whose name contains "
                       f"{', '.join(repr(w) for w in words) or 'anything'})", kind="python")
        return matches[0]
    raise FileNotFoundError(f"There is no template called {text}. Templates you can use: "
                            f"{', '.join(rel(p) for p in templates) or 'none (put one in workspace/templates)'}")


def roles_in_request(request):
    """
    Which deck the user called the template, and which one they want converted, read from their own words
    with plain string rules: "...Y.pptx template", "Y as the template", "the template Y" or a .potx file make Y
    the template; when two decks are named, the other one is the deck to convert.
    -> (source, template) as the user wrote them; None where the words don't say.
    """
    decks = [p for p in named_paths(request) if Path(p).suffix.lower() in (".pptx", ".potx")]
    templates = []
    for name in decks:
        start = request.find(name)
        before = re.findall(r"[a-z]+", request[:start].lower())[-2:]               # "the template (file)" Y
        after = re.findall(r"[a-z]+", request[start + len(name):].lower())[:3]    # Y "template", Y "as the template"
        if (Path(name).suffix.lower() == ".potx" or after[:1] == ["template"] or before[-1:] == ["template"]
                or before == ["template", "file"] or (after[:1] == ["as"] and "template" in after)):
            templates.append(name)
    if len(decks) == 2 and len(templates) == 1:
        return next(d for d in decks if d != templates[0]), templates[0]
    if len(decks) == 1:
        return (None, decks[0]) if templates else (decks[0], None)
    return None, None


def template_named(request):
    """
    "... to the workshop template" -> workspace/templates/workshop_template.pptx: the one template file whose name
    holds every word the user put between "the" and "template". None when the words name no template, or several.
    (A small LLM may copy the template path from the tool description's example instead of the one asked for.)
    """
    for match in re.finditer(r"\bthe\s+([a-z0-9][\w\- ]{0,40}?)\s+template\b", request, re.I):
        words = [w for w in re.findall(r"[a-z0-9]+", match.group(1).lower())
                 if w not in ("a", "new", "slide", "deck", "powerpoint", "pptx")]
        files = [p for p in TEMPLATES_DIR.glob("*") if p.suffix.lower() in (".pptx", ".potx")
                 and words and all(w in p.stem.lower() for w in words)]
        if len(files) == 1:
            return rel(files[0])
    return None


def _same_file(a, b, ctx):
    try:
        return resolve(a, ctx).resolve() == resolve(b, ctx).resolve()
    except Exception:  # not a file we can open (e.g. "the workshop template"): compare the names
        return Path(str(a).replace("\\", "/")).name.lower() == Path(str(b).replace("\\", "/")).name.lower()


# ---- 1. read the source deck ------------------------------------------------------------------------
def read_slides(prs):
    """Every slide as plain data: title, paragraphs (level, bullet or not, bold/italic runs), pictures (bytes
    + aspect ratio), tables (cell text + column widths), speaker notes, and what can't be carried over."""
    from pptx.shapes.group import GroupShape
    from pptx.shapes.picture import Movie, Picture
    sizes = prs.slide_master._element.xpath("./p:txStyles/p:otherStyle/a:lvl1pPr/a:defRPr/@sz")
    shared = {"slide_width": prs.slide_width, "default_sz": int(sizes[0]) if sizes else 1800,
              "cache": {}}  # for slides kept as drawn: their size, the default text size, files already copied
    slides = []
    for number, slide in enumerate(prs.slides, 1):
        info = {"number": number, "layout": slide.slide_layout.name, "title": "", "centred_title": False,
                "paragraphs": [], "pictures": [], "tables": [], "notes": "", "skipped": [], "media": 0,
                "free": 0, "slide": slide, **shared}  # free: shapes that are not placeholders, pictures or tables
        title = slide.shapes.title
        for shape in sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0)):  # reading order
            if title is not None and shape.shape_id == title.shape_id:
                info["title"] = shape.text_frame.text.strip()
                info["centred_title"] = shape.placeholder_format.type == 3  # CENTER_TITLE: a title slide
                continue
            if shape.is_placeholder and shape.placeholder_format.type in (SLIDE_NUMBER, 15, 16):  # number, footer, date
                continue
            if _empty(shape):
                continue
            if not (shape.is_placeholder or isinstance(shape, Picture) or getattr(shape, "has_table", False)):
                info["free"] += 1  # an arrow, a text box, a line, a group, a video, a chart...
            if isinstance(shape, Movie):  # not a Picture in python-pptx, so check it first
                info["media"] += 1
            elif isinstance(shape, Picture):
                try:
                    info["pictures"].append({"blob": shape.image.blob, "aspect": shape.width / shape.height,
                                             "name": shape.name})
                except Exception:  # a linked picture: the file is somewhere else
                    info["skipped"].append(f"linked picture {shape.name}")
            elif getattr(shape, "has_table", False):
                table = shape.table
                info["tables"].append({"rows": [[cell.text for cell in row.cells] for row in table.rows],
                                       "widths": [col.width for col in table.columns],
                                       "first_row": table.first_row, "banded": table.horz_banding})
            elif shape.has_text_frame and shape.text_frame.text.strip():
                for p in shape.text_frame.paragraphs:
                    if p.text.strip():
                        info["paragraphs"].append({
                            "level": p.level, "bullet": bulleted(shape, p.level, p),
                            "runs": [(r.text, bool(r.font.bold), bool(r.font.italic)) for r in p.runs if r.text]
                                    or [(p.text, False, False)]})
            elif getattr(shape, "has_chart", False):
                info["skipped"].append(f"chart {shape.name}")
            elif isinstance(shape, GroupShape):
                info["skipped"].append(f"group {shape.name}")
        if slide.has_notes_slide:
            info["notes"] = slide.notes_slide.notes_text_frame.text.strip()
        slides.append(info)
    return slides


def _empty(shape):
    """An empty text box or a shape with no size: PowerPoint leaves these behind; nobody sees them."""
    if not shape.width or not shape.height:
        return True
    return shape.shape_type == 17 and not shape.text_frame.text.strip()  # 17: TEXT_BOX


# ---- 2. classify each slide -----------------------------------------------------------------------------
def needs_drawing(slide):
    """Does this slide hold more than one placeholder per thing can carry? -> the reason, or None."""
    words = sum(len(text.split()) for p in slide["paragraphs"] for text, _, _ in p["runs"])
    pictures, tables = len(slide["pictures"]), len(slide["tables"])
    if slide["media"]:
        return f"it has {slide['media']} video(s)"
    if slide["free"]:
        return f"it has {slide['free']} shape(s) besides placeholders, pictures and tables (text boxes, arrows, lines, groups)"
    if pictures > 1 or tables > 1 or (pictures and tables):
        return f"it has {pictures} picture(s) and {tables} table(s): more than one thing to place"
    if tables and words:
        return f"a table and {words} words of text: keep each where it was, not both in the space below the title"
    if pictures and words > WORDS_BESIDE_A_PICTURE:
        return f"a picture and {words} words: a half-slide text box would shrink the text too much"
    return None


def classify(slide):
    """Plain-Python rules, first match wins. -> (slide type, the reason, in words)"""
    words = sum(len(text.split()) for p in slide["paragraphs"] for text, _, _ in p["runs"])
    content = slide["paragraphs"] or slide["pictures"] or slide["tables"] or slide["media"] or slide["free"]
    if not slide["title"] and not content:
        return "blank", "it has no title and nothing on it"
    if slide["centred_title"]:
        return "title", "its title is a centred title placeholder (a title-slide layout)"
    if slide["number"] == 1 and slide["title"]:
        return "title", "it is the first slide: the cover"
    drawn = needs_drawing(slide)
    if drawn:
        return "drawn", drawn
    if slide["tables"]:
        return "table", f"it has a table ({len(slide['tables'][0]['rows'])} rows)"
    if slide["pictures"] and slide["paragraphs"]:
        return "picture_text", f"it has a picture and {len(slide['paragraphs'])} paragraph(s) of text"
    if slide["pictures"]:
        return "picture", "it has a picture and no other text"
    if len(slide["paragraphs"]) <= 1 and words <= 12:
        return "section", f"only a title and {words} word(s) of text"
    return "bullets", f"{len(slide['paragraphs'])} paragraphs of text"


# ---- 3. pick a layout for each type ---------------------------------------------------------------------
def placeholder_kinds(layout):
    return Counter(KINDS.get(int(ph.placeholder_format.type)) for ph in layout.placeholders)


def fits(layout, slide_type):
    """Does this layout have the placeholders this type of slide needs?"""
    k = placeholder_kinds(layout)
    if slide_type == "blank":
        return not (k["title"] or k["text"] or k["picture"])
    if not k["title"]:
        return False
    if slide_type == "drawn":
        return bool(k["text"]) and not k["picture"]
    if slide_type == "title":
        return bool(k["subtitle"]) or any(int(ph.placeholder_format.type) == 3 for ph in layout.placeholders)
    if slide_type in ("section", "bullets"):
        return bool(k["text"]) and not k["picture"]
    if slide_type == "picture_text":
        return bool(k["text"]) and (bool(k["picture"]) or k["text"] >= 2)
    if slide_type == "table":
        return not k["picture"]
    return not k["text"]  # picture


def slide_type_of(key):
    """The LLM's name for a kind of slide -> one of SLIDE_TYPES, or None."""
    key = _norm(key).replace("-", " ").replace("_", " ")
    picture = any(w in key for w in ("pic", "image", "photo", "chart", "figure"))
    text = any(w in key for w in ("text", "bullet", "caption"))
    if picture:
        return "picture_text" if text else "picture"
    for slide_type, words in (("title", ("title", "cover", "first")), ("section", ("section", "divider", "header")),
                              ("table", ("table",)), ("bullets", ("bullet", "text", "content", "list"))):
        if any(w in key for w in words):
            return slide_type
    return None


def match_layout(prs, slide_type, wanted):
    """The layout the LLM named, matched forgivingly among the layouts that fit this type: the exact name,
    then names containing all its words ('light blue' -> 'Divider Light Blue'), then the closest spelling."""
    options = [layout for layout in prs.slide_layouts if fits(layout, slide_type)]
    names = [_norm(layout.name) for layout in options]
    words = _norm(wanted).split()
    for name in ([n for n in names if n == _norm(wanted)] + [n for n in names if all(w in n for w in words)]
                 + difflib.get_close_matches(_norm(wanted), names, n=1, cutoff=0.6)):
        return options[names.index(name)]
    raise ValueError(f"The template has no {slide_type} layout called {wanted}. Its {slide_type} layouts are: "
                     f"{', '.join(layout.name.strip() for layout in options[:12])}.")


def pick_layouts(prs, types, wanted):
    """-> {slide type: (layout, why)}"""
    by_name = {_norm(layout.name): layout for layout in prs.slide_layouts}
    plan = {}
    for slide_type in types:
        if slide_type in wanted:
            plan[slide_type] = (match_layout(prs, slide_type, wanted[slide_type]), f"the LLM asked for '{wanted[slide_type]}'")
            continue
        layout = next((by_name[_norm(n)] for n in DEFAULT_LAYOUTS[slide_type] if _norm(n) in by_name), None)
        if layout is not None:
            plan[slide_type] = (layout, "the default for this type")
        else:
            layout = next((l for l in prs.slide_layouts if fits(l, slide_type)), prs.slide_layouts[0])
            plan[slide_type] = (layout, "the first layout with the right placeholders")
    return plan


# ---- 4. build the new deck --------------------------------------------------------------------------------
def learn_from_samples(prs):
    """Before the template's sample slides are deleted, see how its designers made tables: the table style
    and the text size their sample tables use most."""
    styles, sizes = Counter(), Counter()
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                styles.update(shape._element.xpath(".//a:tblPr/a:tableStyleId/text()"))
                sizes.update(int(v) // 100 for v in shape._element.xpath(".//a:tc//a:rPr/@sz"))
    style = styles.most_common(1)[0][0] if styles else None
    names = {}
    try:
        from lxml import etree
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        xml = etree.fromstring(prs.part.part_related_by(RT.TABLE_STYLES).blob)
        names = {s.get("styleId"): s.get("styleName") for s in xml if s.get("styleId")}
    except (KeyError, ValueError):
        pass
    return {"table_style": style, "table_style_name": names.get(style, style),
            "table_text_pt": sizes.most_common(1)[0][0] if sizes else 14}


def remove_all_slides(prs):
    """A template ships with sample slides. Drop them from the slide list and their relationship; python-pptx
    only saves parts that are still connected, so their pictures go too."""
    slide_ids = prs.slides._sldIdLst
    count = len(slide_ids)
    for slide_id in list(slide_ids):
        slide_ids.remove(slide_id)
        prs.part.drop_rel(slide_id.rId)
    return count


def _placeholders(slide):
    """The new slide's placeholders by what they hold, in reading order (top to bottom, then left to right)."""
    found = {"title": [], "subtitle": [], "text": [], "picture": []}
    for ph in sorted(slide.placeholders, key=lambda p: (p.top or 0, p.left or 0)):
        kind = KINDS.get(int(ph.placeholder_format.type))
        if kind:
            found[kind].append(ph)
    return found


def _box(shape):
    return shape.left, shape.top, shape.width, shape.height


def _free_area(slide, prs):
    """The empty part of a layout: below the title, above the footer band (logo, slide number)."""
    width, height = prs.slide_width, prs.slide_height
    title = slide.shapes.title
    top = title.top + title.height + int(height * 0.04) if title is not None else int(height * 0.08)
    left, area_w = (title.left, title.width) if title is not None else (int(width * 0.05), int(width * 0.9))
    return left, top, area_w, int(height * 0.90) - top


def _fit_inside(aspect, box):
    left, top, width, height = box
    if width / height > aspect:  # box is wider than the picture: full height, centred
        w = int(height * aspect)
        return left + (width - w) // 2, top, w, height
    h = int(width / aspect)
    return left, top + (height - h) // 2, width, h


def looks_like_photo(blob):
    """Plain Python's guess, no vision model: a photo has tens of thousands of colours, a chart or a diagram a
    few hundred. (Every 2nd-4th pixel is enough to count; resizing would blend new colours in.)"""
    image = Image.open(io.BytesIO(blob)).convert("RGB")
    step = max(1, max(image.size) // 512)
    image = image.resize((image.width // step, image.height // step), Image.NEAREST)
    return image.getcolors(4096) is None


def _put_text(ph, paragraphs):
    """Paragraphs into a placeholder, keeping levels, bold and italic. The template decides fonts and colours.
    Levels shift if the template bullets different levels than the source did (ULRI: level 1 is a plain
    paragraph, level 2 a bullet), so bullets stay bullets. -> the shift used"""
    shift = max(range(3), key=lambda s: (sum(bulleted(ph, p["level"] + s) == p["bullet"] for p in paragraphs), -s))
    frame = ph.text_frame
    for i, para in enumerate(paragraphs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.level = min(8, para["level"] + shift)
        for text, bold, italic in para["runs"]:
            run = p.add_run()
            run.text = text
            if bold:
                run.font.bold = True
            if italic:
                run.font.italic = True
    return shift


def _split(paragraphs, n):
    """Cut a list of bullets into n columns as evenly as possible, only before a top-level bullet."""
    top = min(p["level"] for p in paragraphs)
    starts = [i for i, p in enumerate(paragraphs) if p["level"] == top and i > 0]
    cuts = sorted({min(starts, key=lambda i: abs(i - k * len(paragraphs) / n)) for k in range(1, n)}) if starts else []
    bounds = [0] + cuts + [len(paragraphs)]
    return [paragraphs[a:b] for a, b in zip(bounds, bounds[1:])]


def _put_picture(slide, picture, ph, box):
    """A picture into a picture placeholder (cropped to fill it, if little is lost), or fitted into a box.
    -> a sentence saying what was decided and why"""
    aspect = picture["aspect"]
    if ph is not None:
        box = _box(ph)
        lost = 1 - min(aspect / (box[2] / box[3]), (box[2] / box[3]) / aspect)
        photo = looks_like_photo(picture["blob"])
        pic = ph.insert_picture(io.BytesIO(picture["blob"]))  # python-pptx crops it to fill the placeholder
        if lost <= CROP_ANY or (photo and lost <= CROP_PHOTO):
            return f"cropped to fill the picture placeholder ({lost:.0%} cut off; {'a photo' if photo else 'a drawing'})"
        pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
        pic.left, pic.top, pic.width, pic.height = _fit_inside(aspect, box)
        return (f"fitted inside the picture placeholder, not cropped (filling it would cut off {lost:.0%} of "
                f"{'a photo' if photo else 'a drawing, maybe its labels'})")
    left, top, width, height = _fit_inside(aspect, box)
    slide.shapes.add_picture(io.BytesIO(picture["blob"]), left, top, width, height)
    return "fitted into the free space below the title (the layout has no picture placeholder)"


def _put_table(slide, table, box, learned):
    from pptx.util import Pt
    rows, cols = len(table["rows"]), max(len(r) for r in table["rows"])
    left, top, width, height = box
    row_height = Pt(learned["table_text_pt"] * 2.2)
    frame = slide.shapes.add_table(rows, cols, left, top, width, min(height, row_height * rows))
    new = frame.table
    total = sum(table["widths"]) or 1
    for col, w in zip(new.columns, table["widths"]):
        col.width = int(width * w / total)
    new.first_row, new.horz_banding = table["first_row"], table["banded"]
    if learned["table_style"]:
        new._tbl.xpath("./a:tblPr/a:tableStyleId")[0].text = learned["table_style"]
    for r, row in enumerate(table["rows"]):
        for c, text in enumerate(row):
            cell = new.cell(r, c)
            cell.text = text
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(learned["table_text_pt"])


def _add_slide_number(slide, layout):
    """python-pptx leaves out the slide-number placeholder when it adds a slide; copy it from the layout."""
    for ph in layout.placeholders:
        if int(ph.placeholder_format.type) == SLIDE_NUMBER:
            element = deepcopy(ph._element)
            element.xpath("./p:nvSpPr/p:cNvPr")[0].set("id", str(slide.shapes._next_shape_id))
            slide.shapes._spTree.append(element)


def fill_slide(prs, slide, info, slide_type, layout, learned):
    """Put one source slide's content into a new slide made from the chosen layout.
    -> (decisions in words, things that could not be carried over)"""
    if slide_type == "drawn":
        return fill_drawn(prs, slide, info, layout)
    if slide_type == "blank":
        _add_slide_number(slide, layout)
        return ["left blank, as it was"], []
    phs, said, lost = _placeholders(slide), [], list(info["skipped"])
    if slide_type == "title":  # the old cover's own design (background shapes, logos) is replaced by the template's
        old = len(info["pictures"]) + len(info["tables"]) + info["free"]
        if old:
            lost.append(f"the old cover's {old} design element(s) (pictures, shapes): the template's cover replaces them")
        info = {**info, "pictures": [], "tables": []}
    if phs["title"]:
        phs["title"][0].text_frame.text = info["title"]
    texts, paragraphs = phs["text"], info["paragraphs"]
    if slide_type == "title":
        slots = phs["subtitle"] + texts
        if slots and paragraphs:
            parts = [paragraphs] if len(slots) == 1 else [paragraphs[:1], paragraphs[1:]]
            for ph, part in zip(slots, parts):
                if part:
                    _put_text(ph, part)
        paragraphs = []
    pictures, tables = list(info["pictures"]), list(info["tables"])
    if tables:
        table = tables.pop(0)
        holder = texts.pop(0) if texts and not paragraphs else None  # the table takes the text placeholder's place
        _put_table(slide, table, _box(holder) if holder is not None else _free_area(slide, prs), learned)
        if holder is not None:
            holder._element.getparent().remove(holder._element)
        said.append(f"table re-created ({len(table['rows'])} rows) in the template's style "
                    f"'{learned['table_style_name'] or 'default'}', {learned['table_text_pt']} pt text")
    if pictures:
        picture = pictures.pop(0)
        if phs["picture"]:
            said.append("picture " + _put_picture(slide, picture, phs["picture"][0], None))
        elif len(texts) >= 2:  # e.g. 'Two Content': the second box holds the picture
            holder = texts.pop()
            said.append("picture " + _put_picture(slide, picture, None, _box(holder)))
            holder._element.getparent().remove(holder._element)
        else:
            said.append("picture " + _put_picture(slide, picture, None, _free_area(slide, prs)))
    if paragraphs:
        columns = [ph for ph in texts if (ph.top, ph.height) == (texts[0].top, texts[0].height)] if texts else []
        if slide_type == "bullets" and len(columns) > 1:  # e.g. 'Two Column': share the bullets out
            groups = _split(paragraphs, len(columns))
            for ph, group in zip(columns, groups):
                _put_text(ph, group)
            said.append(f"{len(paragraphs)} paragraphs split over {len(groups)} columns, before a top-level bullet")
        elif texts:
            shift = _put_text(texts[0], paragraphs)
            if shift:
                said.append(f"text levels moved down {shift} so bullets stay bullets (the template's first level "
                            "has no bullet)")
        else:
            lost.append(f"{len(paragraphs)} paragraph(s) of text (layout {layout.name.strip()} has no text box)")
    lost += [f"picture {p['name']}" for p in pictures] + ["a second table" for _ in tables]
    for ph in list(slide.placeholders):  # an empty placeholder shows 'Click to add text' in PowerPoint
        if ph._element.tag.endswith("}sp") and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)
    _add_slide_number(slide, layout)
    if info["notes"]:
        slide.notes_slide.notes_text_frame.text = info["notes"]
    return said, lost


# ---- 4b. busy slides, kept as drawn ------------------------------------------------------------------------
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


class _Unsupported(Exception):
    """A file that could not be copied with its own references intact: the shape is left out and reported."""


def _clone_part(src, package, cache):
    """
    Copy a file into the new deck (a picture, a video, a chart), once, however many shapes use it. A chart has
    files of its own (its spreadsheet, colours, style); they are copied too, and its relationship ids must come
    out the same, because the chart's XML is copied byte for byte and refers to them by id.
    """
    from pptx.opc.package import Part
    if src.partname in cache:
        return cache[src.partname]
    template = re.sub(r"\d*(\.\w+)$", r"%d\1", str(src.partname))  # /ppt/media/image12.png -> image%d.png
    new = cache[src.partname] = Part(package.next_partname(template), src.content_type, package, src.blob)
    for rid in sorted(src.rels, key=lambda r: int(re.sub(r"\D", "", r) or 0)):
        rel = src.rels[rid]
        target = rel.target_ref if rel.is_external else _clone_part(rel.target_part, package, cache)
        if new.relate_to(target, rel.reltype, is_external=rel.is_external) != rid:
            raise _Unsupported(f"{src.partname} refers to its own files by ids that could not be kept")
    return new


def _copy_relationships(element, src_part, dst_part, cache):
    """Every r:embed / r:link / r:id in copied XML names a relationship of the SOURCE slide (a picture, a video
    and its poster frame, a hyperlink). Give the new slide the same files and links, and rewrite the ids."""
    for el in element.iter():
        for attr, rid in list(el.attrib.items()):
            if not attr.startswith(R_NS) or not rid:
                continue
            if rid not in src_part.rels:
                del el.attrib[attr]
                continue
            rel = src_part.rels[rid]
            if rel.is_external:
                el.set(attr, dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True))
            else:
                el.set(attr, dst_part.relate_to(_clone_part(rel.target_part, dst_part.package, cache), rel.reltype))


def _set_geometry(el, box, original):
    """Write a shape's new position and size into its XML: spPr/xfrm (shapes, pictures, lines), grpSpPr/xfrm
    (groups, whose children keep their own coordinates) or p:xfrm (tables)."""
    from lxml import etree
    tag = etree.QName(el).localname
    if tag == "graphicFrame":
        xfrm = el.find(f"{P_NS}xfrm")
    else:
        holder = el.find(f"{P_NS}grpSpPr" if tag == "grpSp" else f"{P_NS}spPr")
        xfrm = holder.find(f"{A_NS}xfrm")
        if xfrm is None:  # a placeholder that inherited its position from the old layout
            xfrm = etree.Element(f"{A_NS}xfrm")
            holder.insert(0, xfrm)
    values = {"off": ("x", "y", box[0], box[1]), "ext": ("cx", "cy", box[2], box[3])}
    if tag == "grpSp":
        values.update({"chOff": ("x", "y", original[0], original[1]), "chExt": ("cx", "cy", original[2], original[3])})
    for name, (a, b, va, vb) in values.items():
        child = xfrm.find(f"{A_NS}{name}")
        if child is None:
            child = etree.SubElement(xfrm, f"{A_NS}{name}")
        elif name in ("chOff", "chExt"):
            continue  # a group's own coordinate system stays as it was
        child.set(a, str(int(va)))
        child.set(b, str(int(vb)))


def _scale_text(el, k, default_sz):
    """Shrink text with its shape. Sizes are in 1/100 pt; a run with no size of its own gets the source deck's
    default (scaled), unless its text box sets one."""
    from lxml import etree
    half_points = lambda sz: str(max(600, round(sz * k / 50) * 50))  # rounded to 0.5 pt, never below 6 pt
    for tag in ("rPr", "endParaRPr", "defRPr"):
        for rpr in el.iter(f"{A_NS}{tag}"):
            if "sz" in rpr.attrib:
                rpr.set("sz", half_points(int(rpr.get("sz"))))
    for body_pr in el.iter(f"{A_NS}bodyPr"):  # the margins inside a text box shrink with it too
        for side, default in (("lIns", 91440), ("rIns", 91440), ("tIns", 45720), ("bIns", 45720)):
            body_pr.set(side, str(int(int(body_pr.get(side, default)) * k)))
    for body in list(el.iter(f"{P_NS}txBody")) + list(el.iter(f"{A_NS}txBody")):
        if any("sz" in d.attrib for d in body.iter(f"{A_NS}defRPr")):
            continue
        for run in body.iter(f"{A_NS}r"):
            rpr = run.find(f"{A_NS}rPr")
            if rpr is None:
                rpr = etree.Element(f"{A_NS}rPr")
                run.insert(0, rpr)
            if "sz" not in rpr.attrib:
                rpr.set("sz", half_points(default_sz))
    for col in el.iter(f"{A_NS}gridCol"):  # a table's columns, rows and cell margins are sized in its own XML
        col.set("w", str(int(int(col.get("w")) * k)))
    for cell in el.iter(f"{A_NS}tcPr"):
        for side, default in (("marL", 91440), ("marR", 91440), ("marT", 45720), ("marB", 45720)):
            cell.set(side, str(int(int(cell.get(side, default)) * k)))
    for row in el.iter(f"{A_NS}tr"):
        row.set("h", str(int(int(row.get("h")) * k)))


def _freeze_sizes(shape, element, default_sz):
    """A copied shape leaves its old layout and master behind, and with them the text sizes it inherited from
    them. Write each run's inherited size into the copy first (then _scale_text scales it)."""
    from lxml import etree
    from pptx.shapes.group import GroupShape
    if isinstance(shape, GroupShape):
        kinds = ("sp", "pic", "grpSp", "graphicFrame", "cxnSp")
        children = [c for c in element if etree.QName(c).localname in kinds]
        for child_shape, child_element in zip(shape.shapes, children):
            _freeze_sizes(child_shape, child_element, default_sz)
        return
    body = element.find(f"{P_NS}txBody")
    if not shape.has_text_frame or body is None:
        return
    for p, p_element in zip(shape.text_frame.paragraphs, body.findall(f"{A_NS}p")):
        inherited = style_value(shape, p.level, "a:defRPr[@sz]", p)
        size = inherited.get("sz") if inherited is not None else str(default_sz)
        for run in p_element.findall(f"{A_NS}r"):
            rpr = run.find(f"{A_NS}rPr")
            if rpr is None:
                rpr = etree.Element(f"{A_NS}rPr")
                run.insert(0, rpr)
            if rpr.get("sz") is None:
                rpr.set("sz", size)


def _paragraphs_of(shape, default_sz):
    """A text box's paragraphs, blank ones included (on a busy slide they make room for a picture or a note
    beside the text), with each run's real size: its own, or the one it inherits up the style chain, times
    the shrink PowerPoint applied if the box was set to 'shrink text on overflow'."""
    scale = shape.text_frame._txBody.xpath("./a:bodyPr/a:normAutofit/@fontScale")
    scale = int(scale[0]) / 100000 if scale else 1.0
    paragraphs = []
    for p in shape.text_frame.paragraphs:
        inherited = style_value(shape, p.level, "a:defRPr[@sz]", p)
        size = int(inherited.get("sz")) if inherited is not None else default_sz
        runs = [(r.text, bool(r.font.bold), bool(r.font.italic)) for r in p.runs if r.text] or [(p.text, False, False)]
        sizes = [r.font.size.centipoints if r.font.size else size for r in p.runs if r.text] or [size]
        paragraphs.append({"level": p.level, "bullet": bulleted(shape, p.level, p), "runs": runs,
                           "sizes": [s * scale for s in sizes]})
    while paragraphs and not "".join(t for t, _, _ in paragraphs[-1]["runs"]).strip():
        paragraphs.pop()  # trailing blank lines hold nothing apart
    return paragraphs


def fill_drawn(prs, slide, info, layout):
    """
    A busy slide (videos, several pictures, arrows, text boxes...), kept as drawn. The template gives the title
    and its text box, with the template's fonts and bullets; every other shape is copied with the files it
    uses and placed as it was, scaled into the space below the template's title. The source slide's animations
    and video play settings come along if every shape they refer to was copied.
    -> (decisions in words, things that could not be carried over)
    """
    from pptx.shapes.picture import Movie
    source, said, lost = info["slide"], [], []
    phs = _placeholders(slide)
    if phs["title"]:
        phs["title"][0].text_frame.text = info["title"]
    old_title = source.shapes.title
    is_title = lambda sh: old_title is not None and sh.shape_id == old_title.shape_id
    content = [sh for sh in source.shapes if not is_title(sh) and not _empty(sh)
               and not (sh.is_placeholder and sh.placeholder_format.type in (SLIDE_NUMBER, 15, 16))]
    ids = {str(old_title.shape_id): str(phs["title"][0].shape_id)} if old_title is not None and phs["title"] else {}
    if content:
        x0, y0 = min(sh.left for sh in content), min(sh.top for sh in content)
        x1, y1 = max(sh.left + sh.width for sh in content), max(sh.top + sh.height for sh in content)
        left, top, area_w, area_h = _free_area(slide, prs)
        k = min(prs.slide_width / info["slide_width"], area_w / max(1, x1 - x0), area_h / max(1, y1 - y0))
        dx = left + (area_w - (x1 - x0) * k) / 2
        place = lambda sh: (dx + (sh.left - x0) * k, top + (sh.top - y0) * k, sh.width * k, sh.height * k)
        body = next((sh for sh in content if sh.is_placeholder and sh.has_text_frame and sh.text_frame.text.strip()), None)
        if body is not None and phs["text"]:  # the main text goes into the template's own text box
            from pptx.util import Pt
            box = phs["text"][0]
            paragraphs = _paragraphs_of(body, info["default_sz"])
            shift = _put_text(box, paragraphs)
            # the template's fonts and bullets, but the old slide's sizes, scaled like everything around them
            for para, p in zip(paragraphs, box.text_frame.paragraphs):
                for size, run in zip(para["sizes"], p.runs):
                    run.font.size = Pt(max(6, round(size * k / 50) / 2))
            box.left, box.top, box.width, box.height = (int(v) for v in place(body))
            ids[str(body.shape_id)] = str(box.shape_id)
            said.append("text in the template's text box, where it was on the old slide"
                        + (f" (levels moved down {shift})" if shift else ""))
            content = [sh for sh in content if sh.shape_id != body.shape_id]
        tree = slide.shapes._spTree
        next_id = max(int(i) for i in tree.xpath(".//p:cNvPr/@id")) + 1
        copied = videos = 0
        for sh in content:  # in the source's front-to-back order
            element = deepcopy(sh._element)
            try:
                _copy_relationships(element, source.part, slide.part, info["cache"])
            except _Unsupported as exc:
                lost.append(f"{sh.name} (it has files of its own: {exc})")
                continue
            for ph in element.xpath("./*/p:nvPr/p:ph"):  # a placeholder in the old layout: an ordinary shape here
                ph.getparent().remove(ph)
            for c in element.xpath(".//p:cNvPr"):
                ids[c.get("id")] = str(next_id)
                c.set("id", str(next_id))
                next_id += 1
            _set_geometry(element, place(sh), (sh.left, sh.top, sh.width, sh.height))
            _freeze_sizes(sh, element, info["default_sz"])
            _scale_text(element, k, info["default_sz"])
            tree.append(element)
            copied += 1
            videos += isinstance(sh, Movie)
        said.append(f"kept as drawn: {copied} shape(s) copied at {k:.0%} of their size"
                    + (f", {videos} video(s) with their files" if videos else ""))
    timing = source._element.find(f"{P_NS}timing")
    if timing is not None:
        timing = deepcopy(timing)
        targets = [el for el in timing.iter() if "spid" in el.attrib]
        if all(el.get("spid") in ids for el in targets):
            for el in targets:
                el.set("spid", ids[el.get("spid")])
            slide._element.find(f"{P_NS}clrMapOvr").addnext(timing)
            said.append("animations and video play settings copied")
        else:
            lost.append("animations (they point at shapes that were not copied)")
    for ph in list(slide.placeholders):  # an empty placeholder shows 'Click to add text' in PowerPoint
        if ph._element.tag.endswith("}sp") and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)
    _add_slide_number(slide, layout)
    if info["notes"]:
        slide.notes_slide.notes_text_frame.text = info["notes"]
    return said, lost


def shrink_to_fit(shape):
    """If the estimate says the text is too long for its box, turn on PowerPoint's 'shrink text on overflow'
    with the first font scale that leaves some room. -> (ratio before, scale) or None if it fits already."""
    from pptx.enum.text import MSO_AUTO_SIZE
    before = text_fit(shape)
    if before is None or before <= TOO_FULL:
        return None
    scale = next((s for s in FONT_SCALES if text_fit(shape, s, 0.1) <= TARGET), FONT_SCALES[-1])
    shape.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    autofit = shape.text_frame._txBody.bodyPr.xpath("./a:normAutofit")[0]
    autofit.set("fontScale", str(round(scale * 100000)))
    autofit.set("lnSpcReduction", "10000")
    return before, scale


# =========================================================================================
class ConvertPresentationTool(Tool):
    name = "convert_presentation"
    description = ("Move a PowerPoint deck into another template: every slide gets the template's matching layout "
                   "(title, section, bullets, picture with text, table, picture), its fonts and colours, with the "
                   "same text, pictures, tables and speaker notes. The converted deck is saved as a copy in the run "
                   "folder and checked slide by slide.")
    parameters = {
        "type": "object",
        "properties": {
            "source": {"type": "string", "description": "The deck whose slides the user wants converted (their "
                                                         "content), e.g. workspace/old_format_deck.pptx"},
            "template": {"type": "string",
                         "description": "The deck or .potx the user calls the template (the new design), e.g. "
                                        "workspace/templates/workshop_template.pptx"},
            "layouts": {"type": "object",
                        "description": "Optional, only if the user asked for a layout: slide type -> layout name, "
                                       "for example section -> Section Header Light. Slide types: title, section, "
                                       "bullets, picture_text, table, picture."},
        },
        "required": ["source", "template"],
    }
    model_id = "python-pptx + if/else rules (no ML)"
    how_it_works = ("No model here: python-pptx reads every slide into plain data, if/else rules give each slide a "
                    "type, each type gets the template layout made for it, and python-pptx fills that layout's "
                    "placeholders. The LLM only named the files (and maybe a layout).")

    def validate(self, args):
        """
        layouts is the one argument that isn't text, and Llama 3.2 3B sends it as a string: "{}", or a Python
        dict "{'section': 'Divider Light Blue'}" (single quotes are not JSON). Read it in plain Python: JSON,
        then a Python literal, then "section: Divider Light Blue" pairs.
        """
        clean = super().validate(args)
        wanted = clean.get("layouts") or {}
        if isinstance(wanted, str) and _norm(wanted) in ("", "default", "none", "auto"):
            wanted = {}
        elif isinstance(wanted, str):
            try:
                wanted = json.loads(wanted)
            except json.JSONDecodeError:
                try:
                    wanted = ast.literal_eval(wanted)
                except (ValueError, SyntaxError):
                    wanted = dict(re.findall(r"([\w ]+?)\s*(?::|->|=)\s*([^,;{}]+)", wanted.strip("{} "))) or wanted
        if not isinstance(wanted, dict):
            raise ValueError("layouts must be an object such as {\"section\": \"Section Header Light\"}, or left out.")
        layouts = {}
        for key, value in wanted.items():
            slide_type = slide_type_of(key)
            if slide_type is None:
                raise ValueError(f"layouts: {key} is not a slide type. Use title, section, bullets, picture_text, "
                                 "table or picture.")
            if value and _norm(value) not in ("default", "none", "auto"):
                layouts[slide_type] = str(value)
        clean["layouts"] = layouts
        return clean

    def check_request(self, args, ctx):
        """
        Two paths that look alike ("..\\a.pptx", "..\\b.pptx") are easy for a small LLM to put in the wrong
        slots, and a swapped call still "works": it pours the template's sample slides into the user's design.
        The user's own words say which deck is which, so plain Python checks. -> (args, summary, note or None)
        """
        source, template = roles_in_request(ctx.request)
        template = template or template_named(ctx.request)  # "to the workshop template": a name, not a path
        if source is None and template is None:
            return args, "the request doesn't say which deck is the template: the LLM's choice stands", None
        wanted = dict(args)
        if source and template:
            wanted["source"], wanted["template"] = source, template
        elif source and _same_file(args["template"], source, ctx):  # the deck to convert sits in the template slot
            wanted["source"], wanted["template"] = source, args["source"]
        elif template and _same_file(args["source"], template, ctx):
            wanted["source"], wanted["template"] = args["template"], template
        elif template:  # the user named one template, the LLM wrote another (often the description's example)
            wanted["template"] = template
        user = f"the user's words: convert {source or '(not named)'}, template {template or '(not named)'}"
        changed = [key for key in ("source", "template") if not _same_file(wanted[key], args[key], ctx)]
        if not changed:
            return args, f"{user} -> the LLM's arguments match", None
        swapped = _same_file(wanted["source"], args["template"], ctx)
        ctx.tracer.say(f"{user}; the LLM wrote source={args['source']}, template={args['template']}", kind="python")
        note = (f"Your call had source and template swapped. The user asked to convert {wanted['source']} to the "
                f"{wanted['template']} template, so the harness ran it that way." if swapped else
                f"The harness changed {' and '.join(changed)} to match the user's words: source {wanted['source']}, "
                f"template {wanted['template']}.")
        return wanted, ("source and template were SWAPPED -> put back" if swapped
                        else f"{' and '.join(changed)} changed to the user's file") + f" ({user})", note

    def run(self, ctx, source, template, layouts=None):
        t, layouts = ctx.tracer, layouts or {}
        original = resolve(source, ctx)
        if original.suffix.lower() != ".pptx":
            raise ValueError(f"convert_presentation converts .pptx decks, not {original.name}.")
        deck = current_version(original, ctx)  # if this request already edited the deck, convert the edited copy
        template_path = find_template(template, ctx)
        if template_path.resolve() == deck.resolve():
            raise ValueError("source and template are the same file. Give the deck to convert as source and the "
                             "template as template.")

        with t.stage("READ SOURCE", "python", "python-pptx (no ML)", ".pptx file", "slides as data") as ev:
            t.say(self.how_it_works, kind="python")
            source_prs = open_deck(deck)
            slides = read_slides(source_prs)
            shape = "4:3" if abs(source_prs.slide_width / source_prs.slide_height - 4 / 3) < 0.01 else \
                "16:9" if abs(source_prs.slide_width / source_prs.slide_height - 16 / 9) < 0.01 else "custom size"
            ev.summary = (f"{rel(deck)}: {len(slides)} slides, {shape}; {sum(len(s['pictures']) for s in slides)} "
                          f"picture(s), {sum(len(s['tables']) for s in slides)} table(s), "
                          f"{sum(s['media'] for s in slides)} video(s), {sum(s['free'] for s in slides)} other "
                          f"shape(s), speaker notes on {sum(bool(s['notes']) for s in slides)}")
            ev.details = {"slides": [{k: (len(v) if k == "pictures" else v) for k, v in s.items()
                                      if k not in ("slide", "cache", "slide_width", "default_sz")} for s in slides]}

        with t.stage("CLASSIFY SLIDES", "python", "if/else rules (no ML)", "slides as data", "a type per slide") as ev:
            types = []
            for s in slides:
                slide_type, why = classify(s)
                types.append(slide_type)
                t.say(f"slide {s['number']}: {slide_type:<13} because {why}   \"{s['title'][:50]}\"", kind="python")
            ev.summary = ", ".join(f"{i} {kind}" for i, kind in enumerate(types, 1))

        with t.stage("PICK LAYOUTS", "python", "layout names + placeholder check (no ML)",
                     "slide types (+ the LLM's choices)", "a layout per type") as ev:
            prs = open_deck(template_path)
            plan = pick_layouts(prs, list(dict.fromkeys(types)), layouts)
            for slide_type, (layout, why) in plan.items():
                t.say(f"{slide_type:<13} -> '{layout.name.strip()}'   ({why})", kind="python")
            ev.summary = (f"{rel(template_path)} has {len(prs.slide_layouts)} layouts; "
                          + ", ".join(f"{k} -> {layout.name.strip()}" for k, (layout, _) in plan.items()))

        out = working_copy(original, ctx, name=f"{original.stem}_{template_path.stem}.pptx")
        with t.stage("BUILD DECK", "python", "python-pptx (no ML)", "slides as data + template", ".pptx file") as ev:
            learned = learn_from_samples(prs)
            removed = remove_all_slides(prs)
            t.say(f"removed the template's {removed} sample slides; " + (
                  f"its sample tables use style '{learned['table_style_name']}' at {learned['table_text_pt']} pt"
                  if learned["table_style"] else "it has no sample tables, so tables keep PowerPoint's default "
                  f"style at {learned['table_text_pt']} pt"), kind="python")
            not_carried = []
            for s, slide_type in zip(slides, types):
                layout = plan[slide_type][0]
                new = prs.slides.add_slide(layout)
                said, lost = fill_slide(prs, new, s, slide_type, layout, learned)
                for line in said:
                    t.say(f"slide {s['number']}: {line}", kind="python")
                not_carried += [f"slide {s['number']}: {item}" for item in lost]
            ev.summary = (f"{len(slides)} slides on the template's layouts, speaker notes copied"
                          + (f"; not carried over: {', '.join(not_carried)}" if not_carried else ""))

        with t.stage("FIT TEXT", "python", "font metrics from the .ttf files (PIL, no ML)", "text boxes",
                     "a font scale where text is too long") as ev:
            shrunk = []
            for number, slide in enumerate(prs.slides, 1):
                for shape in slide.shapes:  # placeholders, and the text boxes copied onto busy slides
                    if not shape.has_text_frame or (shape.is_placeholder
                                                    and int(shape.placeholder_format.type) == SLIDE_NUMBER):
                        continue
                    result = shrink_to_fit(shape)
                    if result:
                        shrunk.append(f"slide {number} {shape.name}: needed {result[0]:.0%} of its box -> text at "
                                      f"{result[1]:.1%}")
                        t.say(shrunk[-1], kind="python")
            ev.summary = "; ".join(shrunk) if shrunk else "every text box fits at the template's font sizes"
            prs.save(out)
            ev.artifacts.append(str(out))

        observations = {"converted": rel(original), "template": rel(template_path), "slides": len(slides),
                        "layouts_used": {k: layout.name.strip() for k, (layout, _) in plan.items()}}
        drawn = [s for s, kind in zip(slides, types) if kind == "drawn"]
        if drawn:
            videos = sum(s["media"] for s in drawn)
            observations["kept_as_drawn"] = (f"{len(drawn)} busy slide(s) kept as drawn on the template's layout"
                                             + (f", with {videos} video(s)" if videos else ""))
        if not_carried:
            observations["not_carried_over"] = shorten_list(not_carried)
        if shrunk:
            observations["text_shrunk_to_fit"] = shorten_list([s.split(":")[0] for s in shrunk])
        note = f"{rel(original)} converted to the {template_path.stem} template; the original is unchanged"
        return {"artifacts": [{"type": "deck", "path": out, "note": note}], "observations": observations}


# =========================================================================================
class ExportPdfTool(Tool):
    name = "export_pdf"
    description = ("Save a PowerPoint deck (.pptx) as a PDF. If the deck was converted or edited earlier in this "
                   "request, that new version is exported. The PDF is saved in the run folder.")
    parameters = {
        "type": "object",
        "properties": {"path": {"type": "string", "description": "The .pptx deck, e.g. workspace/old_format_deck.pptx"}},
        "required": ["path"],
    }
    model_id = "PowerPoint through COM, or LibreOffice (no ML)"
    how_it_works = ("No model here: Python asks another program to do it. On Windows it drives PowerPoint through "
                    "COM, an API Office has had since the 1990s, and PowerPoint saves the deck as PDF. Without "
                    "PowerPoint it runs LibreOffice with --headless --convert-to pdf.")

    def run(self, ctx, path):
        deck = current_version(resolve(path, ctx), ctx)
        if deck.suffix.lower() != ".pptx":
            raise ValueError(f"export_pdf exports .pptx decks, not {deck.name}.")
        pdf = ctx.run_dir / f"{deck.stem}.pdf"
        with ctx.tracer.stage("TOOL: export_pdf", "python", self.model_id, ".pptx file", "PDF") as ev:
            ctx.tracer.say(self.how_it_works, kind="python")
            program = save_as_pdf(deck.resolve(), pdf.resolve())
            ev.summary = f"{rel(deck)} -> {pdf.name}, made by {program}"
            ev.artifacts.append(str(pdf))
        return {"artifacts": [{"type": "pdf", "path": pdf, "note": f"PDF of {rel(deck)}"}]}
