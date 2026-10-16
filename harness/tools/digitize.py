"""
digitize_plot: the numbers back out of a picture of a plot.

A plot in a report or on a slide is often only pixels. Getting the data back is split between plain
Python and the vision-language model (VLM), so each does what it is good at:

    FIND PLOT AREA    python   the axes are the lowest and leftmost long straight lines -> the plot box
    FIND LABELS       python   blobs of dark ink under / left of the box -> where each tick label sits (pixels)
    READ LABELS       VLM      a crop of each row of labels -> "0 100 200 300" (what the numbers say)
    CALIBRATE         python   label pixel + label value -> a straight line per axis, checked
    FIND CURVES       python   coloured pixels, grouped by colour -> one pixel path per curve
    READ LEGEND       VLM      a crop next to each legend swatch -> the curve's name
    PIXELS -> NUMBERS python   each path through the axis lines -> x, y values -> CSV
    CHECK             python   redraw the curves from the numbers: how many curve pixels were matched?

The VLM is never asked to estimate a value ("how high is the red curve at 250 s?"): a 2-billion-parameter
model reads printed text well and measures positions badly. Python measures; the VLM reads; Python checks
what the VLM read (the labels of an axis must fall on one straight line).

A .pptx slide can hold the plot as a picture (pixels: all of the above) or as a native chart. A native chart
keeps its numbers in the file, so READ CHART DATA copies them exactly and nothing is digitized.
"""

import csv
import io
import re
from dataclasses import dataclass, field

import numpy as np
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.shapes.picture import Picture
from scipy import ndimage

from ..config import MODELS
from .base import Tool
from .files import rel, resolve

MAX_POINTS = 200    # points per curve in the CSV; one per pixel column would be more than anyone reads
WHITE = 235         # a pixel whose darkest channel is above this is background
LINE_SHARE = 0.4    # a straight run of ink this long (share of the picture's width / height) is an axis or grid line
SAME_COLOR = 40     # RGB distance: this close to a curve's colour (or to a paler, anti-aliased version of it)
MIN_CURVE = 0.08    # a curve piece spans at least this share of the plot's width (shorter: legend, text, noise)
TEXT_PX = 40        # crops for the VLM are enlarged until their text is about this many pixels high
IMAGES = (".png", ".jpg", ".jpeg", ".webp", ".bmp")
# Next to the CSV's path in the result: Llama 3.2 3B, asked to calculate with the curves, gave the script the
# .png instead of the CSV until the result said which file holds the numbers.
CSV_NOTE = "the numbers of every curve: to calculate with them, give this CSV file (not the picture) to a script"
NUMBER = re.compile(r"-?\d+(?:,\d{3})*(?:\.\d+)?(?:[eE][-+]?\d+)?")


@dataclass
class Blob:
    """A patch of ink in pixel coordinates: a tick label, a word, a legend swatch."""
    top: int
    bottom: int
    left: int
    right: int

    @property
    def cx(self):
        return (self.left + self.right) / 2

    @property
    def cy(self):
        return (self.top + self.bottom) / 2

    @property
    def h(self):
        return self.bottom - self.top + 1

    @property
    def w(self):
        return self.right - self.left + 1


@dataclass
class Axis:
    """pixel = a + b * value (linear axis) or a + b * log10(value) (log axis), fitted to the tick labels."""
    scale: str
    a: float
    b: float
    worst_px: float                              # how far the worst label sits from the fitted line
    ignored: list = field(default_factory=list)  # labels that did not fit the others (misread?)

    def value(self, pixel):
        v = (np.asarray(pixel, float) - self.a) / self.b
        return 10 ** v if self.scale == "log" else v

    def describe(self):
        unit = "per decade" if self.scale == "log" else "per unit"
        return (f"{self.scale}, {abs(self.b):.4g} px {unit}, worst label {self.worst_px:.1f} px off the line"
                + (f", ignored {self.ignored}" if self.ignored else ""))


class DigitizePlotTool(Tool):
    name = "digitize_plot"
    description = ("Turn a plot (a line chart) back into numbers. It finds the axes and every coloured curve in the "
                   "picture, reads the axis numbers and the legend with a vision-language model, and writes a CSV "
                   "file with the x and y values of each curve. Works on an image file, or on a slide of a "
                   "PowerPoint file (a picture of a plot, or a native chart). Use it when the user wants the data, "
                   "values or curves from a graph.")
    parameters = {
        "type": "object",
        "properties": {
            "path": {"type": "string", "description": "The plot: an image (.png or .jpg), or a PowerPoint file "
                                                      "(.pptx) together with slide."},
            "slide": {"type": "integer", "minimum": 1,
                      "description": "Only for a .pptx: the number of the slide the plot is on."},
        },
        "required": ["path"],
    }
    model_id = f"numpy + scipy, and {MODELS['vlm']} for the text"
    how_it_works = ("Python measures and the VLM reads: Python finds the axes, the tick labels and the coloured curves "
                    "as pixels; the vision-language model only reads the printed numbers and names in small crops; "
                    "Python fits pixel <-> value lines through the labels and checks they agree.")

    def __init__(self, eyes):
        self.eyes = eyes  # the look_at_image tool: its VLM reads the labels (one model, loaded once, shared)

    @property
    def vram_gb(self):  # the harness frees GPU memory first only if the VLM still has to go there
        return 0.0 if self.eyes.on_gpu() else self.eyes.vram_gb

    # ---- the tool ---------------------------------------------------------------------------
    def run(self, ctx, path, slide=None):
        source = resolve(path, ctx)
        if source.suffix.lower() not in IMAGES + (".pptx",):
            raise ValueError(f"digitize_plot needs an image (.png, .jpg) or a .pptx and a slide, not {source.name}.")
        ctx.tracer.say(self.how_it_works, kind="python")
        out = ctx.new_path("digitized", ".csv")
        crops = out.with_name(out.stem + "_crops")
        crops.mkdir(exist_ok=True)
        if source.suffix.lower() != ".pptx":
            return self.digitize(ctx, Image.open(source), out, crops, rel(source))
        if not slide:
            raise ValueError("Which slide is the plot on? Call digitize_plot again with the slide number.")
        chart, picture = self.open_slide(ctx, source, slide, crops)
        what = f"{rel(source)}, slide {slide}"
        if chart is not None:
            return self.copy_chart_data(ctx, chart, out, what)
        return self.digitize(ctx, picture, out, crops, what)

    def open_slide(self, ctx, source, number, crops):
        with ctx.tracer.stage("OPEN SLIDE", "python", "python-pptx (no ML)", ".pptx slide", "a chart or a picture") as ev:
            prs = Presentation(source)
            if not 1 <= number <= len(prs.slides):
                raise ValueError(f"{source.name} has {len(prs.slides)} slides, so there is no slide {number}.")
            shapes = list(_walk(prs.slides[number - 1].shapes))
            area = lambda s: int(s.width or 0) * int(s.height or 0)
            charts = [s for s in shapes if getattr(s, "has_chart", False) and s.has_chart]
            pictures = [s for s in shapes if isinstance(s, Picture)]
            if charts:
                chart = max(charts, key=area).chart
                ev.summary = (f"slide {number}: a native chart ({chart.chart_type}) -> its numbers are stored in the "
                              f"file, nothing to digitize")
                return chart, None
            if not pictures:
                raise ValueError(f"slide {number} of {source.name} has no chart and no picture to digitize.")
            pic = max(pictures, key=area)
            try:
                image = _apply_crop(Image.open(io.BytesIO(pic.image.blob)), pic)
            except Exception:
                raise ValueError(f"the picture on slide {number} is a {pic.image.content_type} drawing that PIL "
                                 f"cannot open; save it as a .png and digitize that.") from None
            path = crops / f"slide{number}_picture.png"
            image.save(path)
            ev.artifacts.append(str(path))
            ev.summary = (f"slide {number}: no native chart, {len(pictures)} picture(s) -> the largest, "
                          f"{image.width}x{image.height} px, taken out of the file (more pixels than a screenshot)")
            return None, image

    def copy_chart_data(self, ctx, chart, out, what):
        with ctx.tracer.stage("READ CHART DATA", "python", "python-pptx: the chart's own XML (no ML)", "chart part",
                              "numbers -> CSV") as ev:
            series = chart_series(chart)
            titles = chart_axis_titles(chart)
            with open(out, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["series", titles.get("x") or "x", titles.get("y") or "y"])
                for name, xs, ys in series:
                    writer.writerows([name, x, y] for x, y in zip(xs, ys) if y is not None)
            part = chart.part.partname.lstrip("/")
            ev.summary = (f"{len(series)} series, {sum(len(ys) for _, _, ys in series)} points copied exactly from "
                          f"{part}: no pixels, no VLM")
            ev.artifacts.append(str(out))
        observations = {"plot": what, "method": f"a native chart: its numbers were copied from {part}, not digitized",
                        "x_axis": titles.get("x", ""), "y_axis": titles.get("y", ""),
                        "series": [_stats(name, xs, ys) for name, xs, ys in series]}
        return {"artifacts": [{"type": "table", "path": str(out), "note": CSV_NOTE}], "observations": observations}

    def digitize(self, ctx, image, out, crops, what):
        t = ctx.tracer
        say = lambda msg, kind="python": t.say(msg, kind=kind)
        rgb = _rgb(image)
        H, W = rgb.shape[:2]

        with t.stage("FIND PLOT AREA", "python", "numpy: the longest straight runs of ink", "pixels", "plot box") as ev:
            box, lines, grid = find_plot_area(rgb)
            x0, y0, x1, y1 = box
            ev.summary = (f"{W}x{H} px picture; x axis = the lowest long line (row {y1}), y axis = the leftmost "
                          f"(column {x0}) -> plot box {x1 - x0}x{y1 - y0} px"
                          + (f"; grid lines {_hex(grid)}" if grid is not None else ""))

        with t.stage("FIND LABELS", "python", "scipy.ndimage.label: blobs of dark ink", "pixels",
                     "label positions") as ev:
            ink = erase_ticks(dark_ink(rgb) & ~lines, box)
            blobs = find_blobs(ink)
            xl, yl = x_labels(blobs, box, W), y_labels(blobs, box, H)
            titles = find_titles(blobs, box, xl, yl, W)
            ev.summary = (f"{len(blobs)} blobs of ink: {len(xl)} labels under the x axis, {len(yl)} left of the y "
                          f"axis" + (f"; text for: {', '.join(titles)}" if titles else ""))
            ev.details = {"x_label_pixels": [round(b.cx) for b in xl], "y_label_pixels": [round(b.cy) for b in yl]}

        text_h = float(np.median([b.h for b in xl + yl]))
        scale = float(np.clip(TEXT_PX / text_h, 1, 4))
        if self.eyes.model is None:
            self.eyes.load(ctx)
        with t.stage("READ LABELS", "perception", self.eyes.model_id, "crops of the labels",
                     "numbers + titles") as ev:
            x_values = self.read_numbers(ctx, rgb, xl, "from left to right", scale, crops / "x_labels.png", ev)
            y_values = self.read_numbers(ctx, rgb, yl, "from top to bottom", scale, crops / "y_labels.png", ev)
            text = {}
            for key, blob in titles.items():
                crop = _crop(rgb, [blob], scale, rotate=key == "y")  # the y title reads bottom-to-top: turn it
                crop.save(crops / f"{key}_title.png")
                ev.artifacts.append(str(crops / f"{key}_title.png"))
                text[key] = _clean(self.eyes.ask(ctx, crop, "What text is written here? Write only that text."))
                say(f'{key} title: "{text[key]}"', "perception")
            ev.summary = (f"x: {[_g(v) for v in x_values]}; y: {[_g(v) for v in y_values]}"
                          + "".join(f"; {k} title: {v!r}" for k, v in text.items()))

        # ---- Python checks the VLM: evenly spaced labels must fall on one straight line ----
        values, axes = {"x": x_values, "y": y_values}, {}
        for name, blobs, centre in (("x", xl, lambda b: b.cx), ("y", yl, lambda b: b.cy)):
            with t.stage("CALIBRATE", "python", "numpy.polyfit: pixel = a + b * value", f"{name} label pixels + values",
                         f"pixel <-> {name} value") as ev:
                try:
                    axes[name] = fit_axis([centre(b) for b in blobs], values[name], name)
                    ev.summary = f"{name}: {axes[name].describe()}"
                except ValueError as exc:
                    ev.summary = f"{exc} -> ask the VLM again, one label at a time"
            if name in axes:
                continue
            with t.stage("READ LABELS AGAIN", "perception", self.eyes.model_id, f"one crop per {name} label",
                         "numbers") as ev:
                values[name] = self.read_each(ctx, rgb, blobs, scale)
                ev.summary = f"{name}: {[_g(v) if v is not None else None for v in values[name]]}"
            with t.stage("CALIBRATE", "python", "numpy.polyfit: pixel = a + b * value", f"{name} label pixels + values",
                         f"pixel <-> {name} value") as ev:
                axes[name] = fit_axis([centre(b) for b in blobs], values[name], name)  # fails again: the LLM is told
                ev.summary = f"{name}: {axes[name].describe()}"
        (x_values, fx), (y_values, fy) = (values["x"], axes["x"]), (values["y"], axes["y"])

        with t.stage("FIND CURVES", "python", "numpy + scipy.ndimage: pixels grouped by colour", "pixels",
                     "one pixel path per curve") as ev:
            series, ignored = find_curves(rgb, box, grid)
            if not series:
                raise ValueError("found no coloured curves in the plot box. (Black or grey curves, bars and plots "
                                 "made only of markers are not handled.)")
            for s in series:
                s["columns"], s["rows"] = trace(s["mask"])
            ev.summary = (f"{len(series)} curve colour(s): "
                          + ", ".join(f"{_hex(s['color'])} ({int(s['mask'].sum())} px"
                                      + (", legend swatch" if s["swatch"] else "") + ")" for s in series)
                          + (f"; ignored colours (no long line of their own): {ignored}" if ignored else ""))

        with t.stage("READ LEGEND", "perception", self.eyes.model_id, "crops next to the legend swatches",
                     "curve names") as ev:
            spacing = _spacing([s["swatch"] for s in series if s["swatch"]])
            for k, s in enumerate(series, 1):
                s["name"] = f"series {k} ({_hex(s['color'])})"
                where = name_box(ink, s["swatch"], text_h, spacing) if s["swatch"] else None
                if where is None:
                    continue
                crop = _crop(rgb, [where], scale)
                crop.save(crops / f"legend_{k}.png")
                ev.artifacts.append(str(crops / f"legend_{k}.png"))
                name = _clean(self.eyes.ask(ctx, crop, "What text is written here? Write only that text."))
                say(f'swatch {_hex(s["color"])} -> "{name}"', "perception")
                s["name"] = name or s["name"]
            ev.summary = ", ".join(f"{_hex(s['color'])} = {s['name']!r}" for s in series)

        with t.stage("PIXELS -> NUMBERS", "python", "the two axis lines from CALIBRATE", "pixel paths",
                     "x, y values -> CSV") as ev:
            headers = ["series", text.get("x") or "x", text.get("y") or "y"]
            with open(out, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for s in series:
                    s["x"], s["y"] = fx.value(s["columns"]), fy.value(s["rows"])
                    writer.writerows([s["name"], f"{x:.5g}", f"{y:.5g}"] for x, y in zip(s["x"], s["y"]))
            ev.summary = (f"{sum(len(s['x']) for s in series)} points ({', '.join(str(len(s['x'])) for s in series)}"
                          f" per curve, at most {MAX_POINTS}) -> {rel(out)}")
            ev.artifacts.append(str(out))

        with t.stage("CHECK", "python", "PIL: redraw the curves from the numbers", "numbers + the original pixels",
                     "share of curve pixels matched") as ev:
            for s in series:
                s["matched"] = coverage(s["mask"], s["columns"], s["rows"])
            overlay = out.with_name(out.stem + "_overlay.png")
            draw_overlay(rgb, box, xl, yl, series, overlay)
            ev.artifacts.append(str(overlay))
            ev.summary = ("curve pixels on the digitized line (within half a line width): "
                          + ", ".join(f"{s['name']} {s['matched']:.0%}" for s in series))

        notes = [f"{s['name']}: only {s['matched']:.0%} of its pixels lie on the digitized line; check the overlay"
                 for s in series if s["matched"] < 0.9]
        notes += [f"{axis} axis: ignored labels {a.ignored}" for axis, a in (("x", fx), ("y", fy)) if a.ignored]
        observations = {
            "plot": what, "method": "pixels: Python measured the curves, the VLM read the labels",
            "x_axis": f"{text.get('x', 'x')} ({fx.scale}), labels read: {[_g(v) for v in x_values]}",
            "y_axis": f"{text.get('y', 'y')} ({fy.scale}), labels read: {[_g(v) for v in y_values]}",
            "series": [{**_stats(s["name"], s["x"], s["y"]), "pixels_matched": f"{s['matched']:.0%}"} for s in series],
        }
        if text.get("title"):
            observations["title"] = text["title"]
        if notes:
            observations["notes"] = notes
        return {"artifacts": [{"type": "table", "path": str(out), "note": CSV_NOTE},
                              {"type": "preview", "path": str(overlay),
                               "note": "the digitized points drawn over the original plot"}],
                "observations": observations}

    def read_numbers(self, ctx, rgb, blobs, order, scale, path, ev):
        """One VLM question for a whole row of labels. Python knows how many labels there are, so it can tell when
        the answer is wrong; then it asks again, one label at a time."""
        crop = _crop(rgb, blobs, scale)
        crop.save(path)
        ev.artifacts.append(str(path))
        question = (f"This picture shows the {len(blobs)} numbers along one axis of a chart. Write them {order}, "
                    f"separated by spaces. Write only the numbers.")
        answer = self.eyes.ask(ctx, crop, question)
        values = parse_numbers(answer)
        ctx.tracer.say(f'{path.stem}: "{question}" -> "{answer}"', kind="perception")
        if len(values) == len(blobs):
            return values
        ctx.tracer.say(f"{path.stem}: read {len(values)} numbers, but Python found {len(blobs)} labels -> asking "
                       f"about each label on its own", kind="python")
        values = self.read_each(ctx, rgb, blobs, scale)
        ctx.tracer.say(f"{path.stem}: one at a time -> {values}", kind="perception")
        return values

    def read_each(self, ctx, rgb, blobs, scale):
        """Slower but surer: one small crop, one number, per label."""
        values = []
        for blob in blobs:
            found = parse_numbers(self.eyes.ask(ctx, _crop(rgb, [blob], scale),
                                                "What number is written here? Write only the number."))
            values.append(found[0] if found else None)
        return values


# ---- pixels ------------------------------------------------------------------------------------
def _rgb(image):
    """The picture as an H x W x 3 array of numbers. Transparent parts (charts pasted from Excel) become white."""
    image = image.convert("RGBA")
    white = Image.new("RGBA", image.size, (255, 255, 255, 255))
    return np.asarray(Image.alpha_composite(white, image).convert("RGB")).astype(np.int16)


def _longest_runs(mask):
    """For every row of a True/False array: (length, first, last) of its longest unbroken run of True."""
    steps = np.diff(np.pad(mask, ((0, 0), (1, 1))).astype(np.int8), axis=1)
    runs = []
    for row in steps:
        starts, ends = np.flatnonzero(row == 1), np.flatnonzero(row == -1)
        if len(starts) == 0:
            runs.append((0, 0, 0))
            continue
        i = int(np.argmax(ends - starts))
        runs.append((int(ends[i] - starts[i]), int(starts[i]), int(ends[i] - 1)))
    return runs


def find_plot_area(rgb):
    """Axes and grid lines are the long straight runs of ink. The lowest horizontal one is the x axis, the leftmost
    vertical one the y axis: together they give the plot box. -> (box, mask of the lines, grid-line colour)"""
    H, W = rgb.shape[:2]
    ink = rgb.min(axis=2) <= WHITE
    lines = np.zeros((H, W), bool)
    rows, cols = [], []
    for y, (n, a, b) in enumerate(_longest_runs(ink)):
        if n >= LINE_SHARE * W and 3 <= y < H - 3:  # a frame round the whole picture is not an axis
            rows.append((y, a, b))
            lines[y, a:b + 1] = True
    for x, (n, a, b) in enumerate(_longest_runs(ink.T)):
        if n >= LINE_SHARE * H and 3 <= x < W - 3:
            cols.append((x, a, b))
            lines[a:b + 1, x] = True
    if not rows or not cols:
        raise ValueError("found no axes (long straight lines) in the picture. Is it a plot with x and y axes?")
    y1, _, x1 = max(rows)  # the lowest long horizontal line: the x axis, and where it ends
    x0, y0, _ = min(cols)  # the leftmost long vertical line: the y axis, and where it starts
    if x1 - x0 < 0.2 * W or y1 - y0 < 0.2 * H:
        raise ValueError(f"the axes found make a plot box of only {x1 - x0}x{y1 - y0} px; this does not look like a plot.")
    c = rgb.astype(np.float32)
    light_grey = (c.min(axis=2) > 150) & (np.ptp(c, axis=2) < 30)
    inner = np.zeros((H, W), bool)
    inner[y0 + 3:y1 - 2, x0 + 3:x1 - 2] = True
    grid_pixels = c[lines & inner & light_grey]
    grid = np.median(grid_pixels, axis=0) if len(grid_pixels) > 50 else None
    return (x0, y0, x1, y1), lines, grid


def dark_ink(rgb):
    """Text, axes and frames: dark and not coloured. (A dark blue curve is coloured, so it is not ink.)"""
    return (rgb.max(axis=2) < 150) & (np.ptp(rgb, axis=2) < 60)


def erase_ticks(ink, box):
    """Tick marks are short lines touching an axis. Removed, so that a tick and its label are not one blob."""
    x0, y0, x1, y1 = box
    labels, _ = ndimage.label(ink, structure=np.ones((3, 3)))
    ink = ink.copy()
    for i, (ys, xs) in enumerate(ndimage.find_objects(labels), 1):
        h, w = ys.stop - ys.start, xs.stop - xs.start
        x_tick = ys.start <= y1 + 2 and ys.stop > y1 and w <= 3 and h <= 15
        y_tick = xs.stop >= x0 - 2 and xs.start < x0 + 1 and h <= 3 and w <= 15
        if x_tick or y_tick:
            ink[ys, xs] &= labels[ys, xs] != i
    return ink


def find_blobs(ink, join=None):
    """Connected patches of ink. A sideways smear joins the characters of one number or word: first a small one to
    measure the text, then one half a character high (the gap after a "1" can be wide: "14" must stay one blob)."""
    grown = ndimage.binary_dilation(ink, structure=np.ones((1, join or 3), bool))
    labels, _ = ndimage.label(grown, structure=np.ones((3, 3)))
    blobs = []
    for i, (ys, xs) in enumerate(ndimage.find_objects(labels), 1):
        rr, cc = np.nonzero(ink[ys, xs] & (labels[ys, xs] == i))
        if len(rr) >= 3:
            blobs.append(Blob(ys.start + rr.min(), ys.start + rr.max(), xs.start + cc.min(), xs.start + cc.max()))
    if join is None:
        heights = [b.h for b in blobs if b.h >= 6]
        return find_blobs(ink, join=max(5, round(0.5 * float(np.median(heights)))) if heights else 5)
    return blobs


def x_labels(blobs, box, W):
    """The first row of blobs under the x axis, left to right."""
    x0, y0, x1, y1 = box
    below = [b for b in blobs if b.top > y1 + 1 and x0 - 0.05 * W <= b.cx <= x1 + 0.05 * W]
    if len(below) < 2:
        raise ValueError("found fewer than two numbers under the x axis.")
    first = min(below, key=lambda b: b.top)
    return sorted((b for b in below if abs(b.cy - first.cy) <= 0.6 * first.h), key=lambda b: b.cx)


def y_labels(blobs, box, H):
    """The column of blobs nearest the y axis, on its left, top to bottom."""
    x0, y0, x1, y1 = box
    left = [b for b in blobs if b.right < x0 - 1 and y0 - 0.05 * H <= b.cy <= y1 + 0.05 * H]
    if len(left) < 2:
        raise ValueError("found fewer than two numbers left of the y axis.")
    nearest = max(left, key=lambda b: b.right)
    return sorted((b for b in left if b.left <= nearest.right and b.right >= nearest.left), key=lambda b: b.cy)


def find_titles(blobs, box, xl, yl, W):
    """Text beyond the labels: the x-axis title below them, the y-axis title left of them, a title above the box."""
    x0, y0, x1, y1 = box
    found = {}
    below = [b for b in blobs if b.top > max(b.bottom for b in xl) + 1 and x0 - 0.05 * W <= b.cx <= x1 + 0.05 * W]
    if below:
        first = min(below, key=lambda b: b.top)
        found["x"] = _union([b for b in below if b.top <= first.bottom])
    left = [b for b in blobs if b.right < min(b.left for b in yl) - 1 and y0 <= b.cy <= y1]
    if left:
        nearest = max(left, key=lambda b: b.right)
        found["y"] = _union([b for b in left if b.left <= nearest.right and b.right >= nearest.left])
    above = [b for b in blobs if b.bottom < y0 - 1 and x0 <= b.cx <= x1]
    if above:
        last = max(above, key=lambda b: b.bottom)
        found["title"] = _union([b for b in above if b.bottom >= last.top])
    return found


def fit_axis(pixels, values, name):
    """Least squares: pixel = a + b * value. Evenly spaced labels give a perfect line, so a label far off the
    line was misread. Tries a log axis (labels 1, 10, 100) too, and once more without each label in turn."""
    pairs = [(p, v) for p, v in zip(pixels, values) if v is not None]
    if len(pairs) < 2 or len({v for _, v in pairs}) < 2:
        raise ValueError(f"the {name} axis needs at least two different numbers, but the VLM read {values}.")
    p, v = (np.array(column, float) for column in zip(*pairs))
    tolerance = max(2.0, 0.1 * float(np.median(np.diff(np.sort(p)))))
    attempts = [list(range(len(p)))]
    if len(p) >= 4:
        attempts += [[j for j in range(len(p)) if j != i] for i in range(len(p))]  # one misread label?
    for keep in attempts:
        for scale in ("linear", "log"):
            if scale == "log" and (len(keep) < 3 or (v[keep] <= 0).any()):
                continue
            x = v[keep] if scale == "linear" else np.log10(v[keep])
            b, a = np.polyfit(x, p[keep], 1)
            worst = float(np.abs(a + b * x - p[keep]).max())
            if worst <= tolerance:
                ignored = [_g(v[i]) for i in range(len(p)) if i not in keep]
                ignored += [f"label {i + 1} unreadable" for i, val in enumerate(values) if val is None]
                return Axis(scale, float(a), float(b), worst, ignored)
    raise ValueError(f"the {name}-axis numbers do not fall on a straight line, so some were misread: the VLM read "
                     f"{[_g(x) if x is not None else None for x in values]} at pixels {[round(x) for x in pixels]}.")


def _toward_white(pixels, color):
    """Distance from each pixel to a curve colour, or to a paler version of it: the anti-aliased edge of a line
    is the line's colour mixed with the white background."""
    color = np.asarray(color, np.float32)
    whiten = 255.0 - color
    t = np.clip(((pixels - color) @ whiten) / max(float(whiten @ whiten), 1.0), 0.0, 0.6)
    return np.linalg.norm(pixels - (color + np.asarray(t)[..., None] * whiten), axis=-1)


def pick_colors(pixels, min_count):
    """The curve colours: the most frequent colours among the coloured pixels in the plot box. The core of a line is
    one colour, many pixels of it; its anti-aliased edges are spread over many paler shades, each one rare."""
    if len(pixels) == 0:
        return []
    bins = (pixels // 8).astype(int) @ np.array([1024, 32, 1])
    ids, counts = np.unique(bins, return_counts=True)
    picked = []
    for i in ids[np.argsort(-counts)][:60]:
        color = pixels[bins == i].mean(axis=0)
        if any(np.linalg.norm(color - p) < SAME_COLOR or _toward_white(color, p) < SAME_COLOR / 2 for p in picked):
            continue  # a colour already picked, or a paler shade of one
        if (np.linalg.norm(pixels - color, axis=1) < SAME_COLOR).sum() >= min_count:
            picked.append(color)
    return picked


def find_curves(rgb, box, grid):
    """Coloured pixels inside the plot box, grouped by colour. Each colour's long connected pieces are its curve;
    a short flat piece of the same colour is its legend swatch. -> ([{color, mask, swatch}], ignored colours)"""
    x0, y0, x1, y1 = box
    H, W = rgb.shape[:2]
    c = rgb.astype(np.float32)
    candidate = (c.min(axis=2) <= WHITE) & ~((c.max(axis=2) < 100) & (np.ptp(c, axis=2) < 40))  # not paper, not ink
    if grid is not None:
        candidate &= np.linalg.norm(c - grid, axis=2) >= SAME_COLOR
    inside = np.zeros((H, W), bool)
    inside[y0:y1 + 1, x0:x1 + 1] = True
    colors = pick_colors(c[candidate & inside], MIN_CURVE * (x1 - x0))
    if not colors:
        return [], []
    distance = np.stack([_toward_white(c, color) for color in colors])
    nearest = distance.argmin(axis=0)
    assigned = candidate & (distance.min(axis=0) < SAME_COLOR)
    near_box = np.zeros((H, W), bool)
    near_box[max(0, y0 - 3):y1 + 4, max(0, x0 - 3):x1 + 4] = True
    found, ignored = [], []
    for k, color in enumerate(colors):
        mask = assigned & (nearest == k)
        grown = ndimage.binary_dilation(mask, iterations=2)  # bridge small gaps: dashes, where curves cross
        labels, _ = ndimage.label(grown, structure=np.ones((3, 3)))
        curve, short, swatches = np.zeros((H, W), bool), [], []
        for i, (ys, xs) in enumerate(ndimage.find_objects(labels), 1):
            rr, cc = np.nonzero(mask[ys, xs] & (labels[ys, xs] == i))
            if len(rr) == 0:
                continue
            rr, cc = rr + ys.start, cc + xs.start
            piece = Blob(rr.min(), rr.max(), cc.min(), cc.max())
            if near_box[rr, cc].mean() > 0.8 and piece.w >= MIN_CURVE * (x1 - x0):
                curve[rr, cc] = True
            else:
                short.append((piece, rr, cc))
        # a short piece right next to the curve, in columns the curve does not cover yet, is the curve: another
        # curve was drawn across it and cut it off. A short SOLID bar elsewhere is the legend swatch.
        if curve.any():
            thickness = float(np.median(curve.sum(axis=0)[curve.any(axis=0)]))
            joined = True
            while joined:  # again after each pass: a piece can be next to a piece that just joined
                joined = False
                distance = ndimage.distance_transform_edt(~curve)  # pixels from the curve
                covered = curve.any(axis=0)                         # pixel columns the curve already has
                for i, (piece, rr, cc) in reversed(list(enumerate(short))):
                    if (near_box[rr, cc].mean() > 0.8 and distance[rr, cc].min() <= max(12, 3 * thickness)
                            and (~covered[np.unique(cc)]).mean() > 0.7):
                        curve[rr, cc] = covered[cc] = True
                        del short[i]
                        joined = True
            swatches = [piece for piece, rr, cc in short
                        if piece.w >= 3 * piece.h and piece.w >= 12 and len(rr) >= 0.7 * piece.w * piece.h]
        # a colour that only exists along the edge of a curve already found is two lines blended, not a curve
        others = ndimage.binary_dilation(np.any([s["mask"] for s in found], axis=0), iterations=2) if found else None
        if curve.any() and (others is None or others[curve].mean() < 0.7):
            found.append({"color": color, "mask": curve, "swatches": swatches})
        else:
            ignored.append(_hex(color))
    for s in found:  # the biggest solid bar of a colour is its legend swatch
        s["swatch"] = max(s["swatches"], key=lambda b: b.w * b.h) if s["swatches"] else None
    found.sort(key=lambda s: (s["swatch"] is None, s["swatch"].cy if s["swatch"] else 0,
                              s["swatch"].cx if s["swatch"] else 0))  # legend order
    return found, ignored


def trace(mask):
    """One point per pixel column: the middle row of the curve's pixels there. At most MAX_POINTS, evenly spread."""
    rows, columns = np.nonzero(mask)
    order = np.argsort(columns, kind="stable")
    rows, columns = rows[order], columns[order]
    columns, first = np.unique(columns, return_index=True)
    rows = np.array([np.median(r) for r in np.split(rows, first[1:])])
    if len(columns) > MAX_POINTS:
        keep = np.unique(np.linspace(0, len(columns) - 1, MAX_POINTS).round().astype(int))
        columns, rows = columns[keep], rows[keep]
    return columns, rows


def _spacing(swatches):
    """The distance between legend entries (stacked legends), or None."""
    steps = np.diff(sorted(b.cy for b in swatches))
    steps = steps[steps > 2]
    return float(np.median(steps)) if len(steps) else None


def name_box(ink, swatch, text_h, spacing):
    """Where a legend name is: the ink to the right of its swatch, up to a wide gap or a frame line."""
    half = min(text_h, 0.45 * spacing) if spacing else text_h
    top, bottom = max(0, int(swatch.cy - half)), int(swatch.cy + half)
    band = ink[top:bottom + 1, swatch.right + 1:]
    start = end = None
    gap = 0
    for j in range(band.shape[1]):
        column = band[:, j]
        if column.mean() > 0.8:  # ink as tall as the band: the legend's frame, not a letter
            break
        if column.any():
            start = j if start is None else start
            end, gap = j, 0
        elif start is None and j > 3 * text_h:
            return None  # nothing written next to this swatch
        elif start is not None:
            gap += 1
            if gap > text_h:
                break
    if start is None:
        return None
    return Blob(top, bottom, swatch.right + 1 + start, swatch.right + 1 + end)


def coverage(mask, columns, rows):
    """Share of the curve's own pixels near the line drawn through the digitized points: within half the curve's
    thickness (its median height per pixel column) plus one pixel."""
    thickness = float(np.median(mask.sum(axis=0)[mask.any(axis=0)]))
    canvas = Image.new("1", (mask.shape[1], mask.shape[0]), 0)
    ImageDraw.Draw(canvas).line(list(zip(columns.tolist(), rows.tolist())), fill=1, width=1)
    near = ndimage.binary_dilation(np.asarray(canvas, bool), iterations=max(2, round(thickness / 2) + 1))
    return float(near[mask].mean())


def draw_overlay(rgb, box, xl, yl, series, path):
    """The original, faded, with what Python found: the plot box and label positions (blue) and the points."""
    base = Image.fromarray(rgb.astype(np.uint8))
    faded = Image.blend(base, Image.new("RGB", base.size, "white"), 0.6)
    draw = ImageDraw.Draw(faded)
    x0, y0, x1, y1 = box
    blue = (0, 90, 255)
    draw.rectangle([x0, y0, x1, y1], outline=blue, width=2)
    for b in xl:
        draw.line([(b.cx, y1 - 8), (b.cx, y1 + 8)], fill=blue, width=3)
    for b in yl:
        draw.line([(x0 - 8, b.cy), (x0 + 8, b.cy)], fill=blue, width=3)
    for s in series:
        color = tuple(int(v) for v in s["color"])
        points = list(zip(s["columns"].tolist(), s["rows"].tolist()))
        draw.line(points, fill=(0, 0, 0), width=1)
        for x, y in points[::max(1, len(points) // 40)]:
            draw.ellipse([x - 4, y - 4, x + 4, y + 4], fill=color, outline=(0, 0, 0))
        if s["swatch"] is not None:
            b = s["swatch"]
            draw.rectangle([b.left - 3, b.top - 3, b.right + 3, b.bottom + 3], outline=blue, width=2)
    faded.save(path)


def _crop(rgb, blobs, scale, rotate=False, pad=3):
    """A crop for the VLM holding only the ink of these blobs, on white, enlarged so the text is easy to read."""
    H, W = rgb.shape[:2]
    whole = _union(blobs)
    left, top = max(0, whole.left - pad), max(0, whole.top - pad)
    right, bottom = min(W, whole.right + pad + 1), min(H, whole.bottom + pad + 1)
    canvas = np.full((bottom - top, right - left, 3), 255, np.uint8)
    for b in blobs:
        t, l = max(0, b.top - pad), max(0, b.left - pad)
        r, d = min(W, b.right + pad + 1), min(H, b.bottom + pad + 1)
        canvas[t - top:d - top, l - left:r - left] = rgb[t:d, l:r]
    image = Image.fromarray(canvas)
    if scale > 1:
        image = image.resize((round(image.width * scale), round(image.height * scale)), Image.LANCZOS)
    if rotate:
        image = image.rotate(-90, expand=True)  # text that reads bottom-to-top, turned to read left-to-right
    framed = Image.new("RGB", (image.width + 32, image.height + 32), "white")
    framed.paste(image, (16, 16))
    return framed


def _union(blobs):
    return Blob(min(b.top for b in blobs), max(b.bottom for b in blobs),
                min(b.left for b in blobs), max(b.right for b in blobs))


# ---- native charts: the numbers are already in the file ----------------------------------------------
def _walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)


def _apply_crop(image, picture):
    """A picture on a slide can be cropped; digitize what the audience sees."""
    W, H = image.size
    l, t, r, b = (max(0.0, float(v)) for v in (picture.crop_left, picture.crop_top, picture.crop_right,
                                                picture.crop_bottom))
    return image.crop((round(l * W), round(t * H), round(W - r * W), round(H - b * H))) if l or t or r or b else image


def chart_series(chart):
    """[(name, xs, ys)] from the chart part's XML: the cached values PowerPoint draws the chart from."""
    out = []
    for plot in chart.plots:
        try:
            categories = [_number(c) for c in plot.categories]
        except Exception:
            categories = []
        for k, series in enumerate(plot.series, 1):
            ys = list(series.values)
            x_values = series._element.find(qn("c:xVal"))  # scatter charts keep their x values here
            if x_values is not None:
                points = {int(pt.get("idx")): pt.findtext(qn("c:v")) for pt in x_values.iter(qn("c:pt"))}
                xs = [_number(points.get(i)) for i in range(len(ys))]
            else:
                xs = categories or list(range(1, len(ys) + 1))
            out.append((series.name or f"series {k}", xs, ys))
    return out


def chart_axis_titles(chart):
    titles = {}
    for axis in chart._chartSpace.iter(qn("c:catAx"), qn("c:valAx"), qn("c:dateAx")):
        position = axis.find(qn("c:axPos"))
        side = "x" if position is not None and position.get("val") in ("b", "t") else "y"
        title = axis.find(qn("c:title"))
        text = "".join(t.text or "" for t in title.iter(qn("a:t"))) if title is not None else ""
        titles[side] = titles.get(side) or text
    return titles


# ---- small helpers -----------------------------------------------------------------------------
def parse_numbers(text):
    """'0 100 200' / '0, 100, 200' / '1,000' -> floats. Plain regex, no ML."""
    text = text.replace("−", "-").replace("–", "-")
    return [float(m.replace(",", "")) for m in NUMBER.findall(text)]


def _number(text):
    try:
        return float(text)
    except (TypeError, ValueError):
        return text


def _clean(text):
    return text.strip().strip("\"'").strip().rstrip(".|").strip()


def _g(value):
    """A number rounded to 4 significant digits, as a plain float (for the LLM and the trace)."""
    return float(f"{value:.4g}") if isinstance(value, (int, float, np.floating)) else value


def _hex(color):
    return "#" + "".join(f"{int(round(v)):02x}" for v in color)


def _stats(name, xs, ys):
    """What the LLM gets per curve: enough to answer "when does it peak?" without reading the CSV."""
    pairs = [(x, y) for x, y in zip(xs, ys) if y is not None]
    if not pairs:
        return {"name": name, "points": 0}
    top = max(range(len(pairs)), key=lambda i: pairs[i][1])
    return {"name": name, "points": len(pairs), "x_from": _g(pairs[0][0]), "x_to": _g(pairs[-1][0]),
            "max_y": _g(pairs[top][1]), "x_at_max_y": _g(pairs[top][0])}
