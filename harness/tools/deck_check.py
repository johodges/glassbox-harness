"""
Checking a deck without eyes. Llama 3.2 3B can't see a slide, so the harness checks the layout in code
and tells the LLM what it found, as text. Two ways to ask "does this text fit its box?":

    estimate   plain Python: follow PowerPoint's style inheritance (paragraph -> shape -> layout ->
               master) to find each paragraph's font and size, measure every word with that font's real
               letter widths (PIL reads the .ttf file), wrap the lines, add up their heights and compare
               the total with the box. Needs nothing but the file, and it is only an estimate.
    measure    PowerPoint itself, through COM (Windows only): it lays the text out and reports its height
               (TextRange.BoundHeight). Exact. The same connection draws every slide as a PNG and saves a
               deck as PDF. Without PowerPoint, LibreOffice (soffice --headless) can still make the PDF.

PowerPoint may be the user's own, with their deck open (OFFICE_COLLAB_DESIGN.md, F5). So every COM job:
  - runs in a CHILD PROCESS with a time limit (config OFFICE["timeout"]): a hung export never hangs the harness;
  - notes whether POWERPNT.EXE was running before. Only one PowerPoint can run, so it may be the user's: the job
    opens a COPY of the deck read-only without a window, closes only what it opened, and calls Quit() only if it
    started PowerPoint itself. On a timeout it stops the PowerPoint it started, and never the user's;
  - never changes a setting of the application.
When PowerPoint is busy (a dialog, an edit in progress) or absent, the job raises PowerPointUnavailable and the
caller says "not rendered" in its answer.

No neural network anywhere in this file.
"""

import csv
import functools
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from contextlib import contextmanager
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont

EMU_PER_PT = 12700
LINE = 1.2  # PowerPoint's single line spacing is about 1.2 x the font size
FONT_FILES = {"calibri light": "calibril", "arial narrow": "arialn", "segoe ui": "segoeui",
              "times new roman": "times", "courier new": "cour"}  # family -> .ttf name, where it isn't obvious


# ---- the estimate ----------------------------------------------------------------------------------
def _chain(shape):
    """The XML elements a shape's text styles come from: the shape, the layout placeholder it inherits from,
    and the master placeholder above that (a text box that is not a placeholder inherits from nothing)."""
    elements, current = [shape._element], shape
    for _ in range(2):
        try:
            current = current._base_placeholder
        except (AttributeError, NotImplementedError):
            break
        if current is None:
            break
        elements.append(current._element)
    return elements


def _master(shape):
    return shape.part.slide_layout.slide_master


def _text_style(shape):
    """Which of the master's three text styles applies: titles, other placeholders, or plain text boxes."""
    if not shape.is_placeholder:
        return "otherStyle"
    return "titleStyle" if shape.placeholder_format.type in (1, 3) else "bodyStyle"  # 1 TITLE, 3 CENTER_TITLE


def style_value(shape, level, path, paragraph=None):
    """
    The first value of `path` (an XPath below <a:lvlNpPr>) on the way up the inheritance chain, in the
    order PowerPoint looks: the paragraph itself, the shape, its layout placeholder, the master placeholder,
    then the master's text styles. None if nobody sets it.
    """
    tag = f"a:lvl{level + 1}pPr"
    if paragraph is not None:
        found = paragraph._p.xpath(f"./a:pPr/{path}")
        if found:
            return found[0]
    for element in _chain(shape):
        found = element.xpath(f"./p:txBody/a:lstStyle/{tag}/{path}")
        if found:
            return found[0]
    found = _master(shape)._element.xpath(f"./p:txStyles/p:{_text_style(shape)}/{tag}/{path}")
    return found[0] if found else None


BULLETS = "*[self::a:buNone or self::a:buChar or self::a:buAutoNum or self::a:buBlip]"


def bulleted(shape, level, paragraph=None):
    """Does this level of this shape show a bullet? Templates disagree: the Office theme bullets every level,
    ULRI's first level is a plain paragraph."""
    found = style_value(shape, level, BULLETS, paragraph)
    return found is not None and not found.tag.endswith("buNone")


@functools.lru_cache(maxsize=None)
def _theme_fonts(master_part):
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    theme = etree.fromstring(master_part.part_related_by(RT.THEME).blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    major = theme.xpath("//a:majorFont/a:latin/@typeface", namespaces=ns)
    minor = theme.xpath("//a:minorFont/a:latin/@typeface", namespaces=ns)
    return {"+mj-lt": major[0] if major else "Arial", "+mn-lt": minor[0] if minor else "Arial"}


@functools.lru_cache(maxsize=None)
def _ttf(face, bold):
    """The font file behind a family name, loaded at 100 pt so widths scale to any size. None if not installed."""
    base = FONT_FILES.get(face.lower(), face.lower().replace(" ", ""))
    for name in ([base + "bd.ttf", base + "b.ttf"] if bold else []) + [base + ".ttf"]:
        try:
            return ImageFont.truetype(name, 100)
        except OSError:
            continue
    return None


def _wrapped_lines(text, width, measure):
    """Greedy word wrap, as PowerPoint does it: how many lines does this text take in this width (points)?"""
    lines, current, space = 1, 0.0, measure(" ")
    for word in text.split():
        w = measure(word)
        if current and current + space + w > width:
            lines, current = lines + 1, w
        else:
            current += (space if current else 0) + w
        if w > width:  # one word wider than the box breaks over several lines
            lines += math.ceil(w / width) - 1
    return lines


def _spacing(shape, level, tag, paragraph, size):
    points = style_value(shape, level, f"a:{tag}/a:spcPts/@val", paragraph)
    if points is not None:
        return int(points) / 100
    percent = style_value(shape, level, f"a:{tag}/a:spcPct/@val", paragraph)
    return int(percent) / 100000 * size if percent is not None else 0.0


def _inset(shape, side, default):
    for element in _chain(shape):
        found = element.xpath(f"./p:txBody/a:bodyPr/@{side}")
        if found:
            return int(found[0])
    return default


def text_fit(shape, font_scale=None, line_reduction=None):
    """
    How full is this text box? 1.0 = exactly full, 1.25 = the text needs 25% more height than the box has.
    font_scale / line_reduction try out a shrink (PowerPoint's "shrink text on overflow") before applying it.
    None if the shape has no text or no size.
    """
    if not shape.has_text_frame or not shape.text_frame.text.strip() or not shape.width or not shape.height:
        return None
    room_w = (shape.width - _inset(shape, "lIns", 91440) - _inset(shape, "rIns", 91440)) / EMU_PER_PT
    room_h = (shape.height - _inset(shape, "tIns", 45720) - _inset(shape, "bIns", 45720)) / EMU_PER_PT
    autofit = shape.text_frame._txBody.bodyPr.xpath("./a:normAutofit")
    if font_scale is None:
        font_scale = int(autofit[0].get("fontScale", "100000")) / 100000 if autofit else 1.0
    if line_reduction is None:
        line_reduction = int(autofit[0].get("lnSpcReduction", "0")) / 100000 if autofit else 0.0
    fonts = _theme_fonts(_master(shape).part)
    default_face = "+mj-lt" if _text_style(shape) == "titleStyle" else "+mn-lt"
    needed = 0.0
    for i, p in enumerate(shape.text_frame.paragraphs):
        level = p.level
        sizes = [r.font.size.pt for r in p.runs if r.font.size is not None]
        size = (max(sizes) if sizes else int(style_value(shape, level, "a:defRPr/@sz") or 1800) / 100) * font_scale
        face = next((r.font.name for r in p.runs if r.font.name), None) \
            or style_value(shape, level, "a:defRPr/a:latin/@typeface") or default_face
        face = fonts.get(face, face)
        bold = any(r.font.bold for r in p.runs) or style_value(shape, level, "a:defRPr/@b") in ("1", "true")
        font = _ttf(face, bold)
        measure = (lambda s, f=font, z=size: f.getlength(s) * z / 100) if font else (lambda s, z=size: len(s) * z * 0.5)
        indent = int(style_value(shape, level, "@marL", p) or 0) / EMU_PER_PT
        line_pct = int(style_value(shape, level, "a:lnSpc/a:spcPct/@val", p) or 100000) / 100000
        lines = sum(_wrapped_lines(part, room_w - indent, measure) for part in p.text.split("\v"))
        before = _spacing(shape, level, "spcBef", p, size) if i else 0.0  # no space above the first paragraph
        needed += before + lines * size * LINE * line_pct * (1 - line_reduction) + _spacing(shape, level, "spcAft", p, size)
    return needed / room_h if room_h > 0 else None


def check_deck(path):
    """Read a saved deck back: each slide's layout, title and text, and how full each text box looks."""
    from pptx import Presentation
    from pptx.shapes.picture import Picture
    slides = []
    for number, slide in enumerate(Presentation(path).slides, 1):
        title = slide.shapes.title
        text, fit, pictures, tables = [], {}, 0, 0
        for shape in slide.shapes:
            pictures += isinstance(shape, Picture)
            tables += bool(getattr(shape, "has_table", False))
            if not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            if shape.is_placeholder and shape.placeholder_format.type == 13:  # SLIDE_NUMBER
                continue
            if title is None or shape.shape_id != title.shape_id:
                text += [p.text.replace("\v", " ") for p in shape.text_frame.paragraphs if p.text.strip()]
            ratio = text_fit(shape)
            if ratio is not None:
                fit[shape.shape_id] = (shape.name, round(ratio, 2))
        slides.append({"slide": number, "layout": slide.slide_layout.name.strip(),
                       "title": title.text_frame.text.replace("\v", " ") if title is not None else "",
                       "text": text, "pictures": pictures, "tables": tables, "fit": fit})
    return slides


def estimated_fit(path, positions):
    """The estimate above for the slides at these 1-based positions -> {(position, shape id): (name, ratio)}."""
    from pptx import Presentation
    found = {}
    for number, slide in enumerate(Presentation(path).slides, 1):
        if number not in positions:
            continue
        for shape in slide.shapes:
            if shape.is_placeholder and shape.placeholder_format.type == 13:  # SLIDE_NUMBER
                continue
            ratio = text_fit(shape) if shape.has_text_frame else None
            if ratio is not None:
                found[(number, shape.shape_id)] = (shape.name, round(ratio, 2))
    return found


# ---- keep-out zones: a template profile -------------------------------------------------------------
# workspace/template_profiles/<template>.json names the boxes a template keeps for itself (a logo, the band below a
# footer rule), in cm from the slide's top-left corner, like the "stay above 6.3 in" rule of the ULRI and Spyro
# templates. A deck is matched to a profile by its template's names, else by layouts only that template has.
TOUCH_CM = 0.05  # shapes that only touch a zone's edge don't count


def template_profiles():
    from ..config import WORKSPACE_DIR
    folder = WORKSPACE_DIR / "template_profiles"
    return [dict(json.loads(p.read_text(encoding="utf-8")), file=p) for p in sorted(folder.glob("*.json"))] \
        if folder.is_dir() else []


def profile_for(template, deck_path=None):
    """deck_view's template info -> (the profile, how it matched) or (None, why not)."""
    names = [n.lower() for n in template["masters"] + template["themes"] + [template.get("app_template", "")]
             + ([Path(deck_path).stem] if deck_path else []) if n]
    for profile in template_profiles():
        match = profile.get("match", {})
        wanted = [n.lower() for n in match.get("names", [])]
        hit = next((n for n in wanted if any(n in have for have in names)), None)
        if hit:
            return profile, f"the template's name ({hit})"
        layouts, size = match.get("layouts", []), match.get("slide_size_cm")
        same_size = not size or (template.get("slide_size_cm") and all(abs(a - b) < 0.1 for a, b in
                                                                       zip(size, template["slide_size_cm"])))
        if layouts and same_size and all(name in template["layouts"] for name in layouts):
            return profile, f"its layouts ({', '.join(layouts)}) and slide size"
    return None, "no profile in workspace/template_profiles matches this deck's template"


def overlap(box, zone):
    """(left, top, width, height) in cm -> the overlap's (width, height), or None if they only touch or miss."""
    w = min(box[0] + box[2], zone[0] + zone[2]) - max(box[0], zone[0])
    h = min(box[1] + box[3], zone[1] + zone[3]) - max(box[1], zone[1])
    return (round(w, 2), round(h, 2)) if w > TOUCH_CM and h > TOUCH_CM else None


def keep_out_problems(slides, profile, boxes=None):
    """
    Shapes that cross a profile's keep-out zones -> [{"slide", "shape", "zones", "text"}]. slides: deck_view slides;
    boxes: {(position, shape id): (l, t, w, h)} as PowerPoint laid them out (a table that grew rows), where rendered.
    The template's own date, footer and slide-number placeholders belong in the footer band and are allowed.
    """
    allowed = set(profile.get("allowed_placeholders", ["sldNum", "ftr", "dt"]))
    zones = [(z["zone"], (z["left"], z["top"], z["width"], z["height"])) for z in profile.get("keep_out", [])]
    problems = []
    for slide in slides:
        for shape in slide["shapes"]:
            if shape.get("group") or shape.get("hidden") or shape.get("ph") in allowed:
                continue
            box = (boxes or {}).get((slide["position"], shape["id"])) or shape.get("box")
            if not box or shape["kind"] == "placeholder":  # an empty placeholder shows nothing in a slide show
                continue
            crossed = [(name, overlap(box, zone)) for name, zone in zones if overlap(box, zone)]
            if crossed:
                where = " and ".join(f"the {name} ({w:.1f} x {h:.1f} cm)" for name, (w, h) in crossed)
                problems.append({"slide": slide["position"], "shape": shape["name"], "zones": [n for n, _ in crossed],
                                 "text": f"slide {slide['position']}: \"{shape['name']}\" covers {where}"})
    return problems


# ---- stale references: words that may point at the wrong slide after a move, a delete or an import ------------
REFERENCE = re.compile(
    r"\bslides?\s+(\d+)(?:\s*(?:-|–|to|and)\s*(\d+))?\b"
    r"|\b(?:the\s+)?(?:next|previous|following|preceding|last|earlier|coming)\s+"
    r"(?:slides?|sections?|parts?|chapters?)\b"
    r"|\bas\s+(?:shown|seen|discussed|described|mentioned|noted)\s+(?:above|below|earlier|before|previously|later)\b",
    re.I)


def _texts(slide):
    """(where, text) for every piece of text on a slide and in its notes."""
    out = []
    for shape in slide["shapes"]:
        for p in shape.get("paragraphs", []):
            out.append((f"\"{shape['name']}\"", "".join(r["t"] for r in p["runs"])))
        for row in shape.get("table", []):
            out += [(f"\"{shape['name']}\"", cell) for cell in row]
    if slide.get("notes"):
        out += [("notes", line) for line in slide["notes"].split("\n")]
    return out


def _context(text, start, end, width=40):
    left, right = max(0, start - width), min(len(text), end + width)
    return ("..." if left else "") + text[left:right].strip() + ("..." if right < len(text) else "")


def stale_references(slides, before=None):
    """
    After a move, delete or import: every 'slide 12', 'the next section', 'as shown above' in the text and notes.
    before: {"order": [slide ids before the edit], "titles": {id: title}} -> a reference to slide N is stale when the
    slide that was N is now elsewhere or gone; words such as "the next section" are reported on slides whose
    neighbours changed. Reported, never changed.
    """
    now = [s["id"] for s in slides]
    old = [int(i) for i in (before or {}).get("order", [])]
    titles = {int(k): v for k, v in (before or {}).get("titles", {}).items()}
    found = []
    for slide in slides:
        position = slide["position"]
        sid = slide["id"]
        was = old.index(sid) if sid in old else None
        neighbours_changed = was is None or (old[was - 1] if was > 0 else None) != (now[position - 2] if position > 1
                                                                                      else None) \
            or (old[was + 1] if was + 1 < len(old) else None) != (now[position] if position < len(now) else None)
        for where, text in _texts(slide):
            for m in REFERENCE.finditer(text):
                if m.group(1):  # "slide 7": which slide was 7 then, and where is it now?
                    number = int(m.group(1))
                    if not old or not 1 <= number <= len(old):
                        continue
                    meant = old[number - 1]
                    if meant in now and now.index(meant) + 1 == number:
                        continue  # still the same slide
                    where_now = f"now slide {now.index(meant) + 1}" if meant in now else "now deleted"
                    why = f"slide {number} was \"{titles.get(meant, '?')}\", {where_now}"
                elif neighbours_changed:
                    why = "the slides around it changed"
                else:
                    continue
                found.append({"slide": position, "where": where, "words": m.group(0),
                              "text": f"slide {position} {where}: \"{_context(text, m.start(), m.end())}\" ({why})"})
    return found


def term_hits(slides, terms):
    """Every place a term the user named occurs (case ignored), in text and notes -> [{"term", "slide", "text"}]."""
    hits = []
    for term in [t for t in terms if t.strip()]:
        pattern = re.compile(re.escape(term.strip()), re.I)
        for slide in slides:
            for where, text in _texts(slide):
                for m in pattern.finditer(text):
                    hits.append({"term": term, "slide": slide["position"],
                                 "text": f"slide {slide['position']} {where}: "
                                         f"\"{_context(text, m.start(), m.end())}\""})
    return hits


# ---- PowerPoint and LibreOffice --------------------------------------------------------------------
class PowerPointUnavailable(RuntimeError):
    """PowerPoint could not do the job: not installed, busy with the user's own work, or too slow. Edits made in
    Python stand; only the render (or the PDF) is skipped, and the answer says why."""


def powerpoint_available():
    """Windows, pywin32 installed, and PowerPoint registered as a COM server."""
    if os.name != "nt":
        return False
    try:
        import win32com.client  # noqa: F401  (pywin32)
        import winreg
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, "PowerPoint.Application"))
        return True
    except (ImportError, OSError):
        return False


def powerpoint_pids():
    """The process ids of every running POWERPNT.EXE (Windows' tasklist; no extra package)."""
    if os.name != "nt":
        return set()
    out = subprocess.run(["tasklist", "/FI", "IMAGENAME eq POWERPNT.EXE", "/FO", "CSV", "/NH"], capture_output=True,
                         text=True, timeout=30)
    return {int(row[1]) for row in csv.reader(out.stdout.splitlines()) if len(row) > 1
            and row[0].upper() == "POWERPNT.EXE" and row[1].isdigit()}


@contextmanager
def powerpoint(status=None):
    """
    PowerPoint as a COM server: another program on this PC that Python drives through an API. Runs inside the
    child process (powerpoint_job). Yields (app, opened): put every presentation you open in `opened`, and only those
    are closed. PowerPoint itself is quit only if it was not running before and nobody opened anything in it since;
    a PowerPoint you have open is left alone. `status` (a file) records whether it was running and which process
    was started, so the parent can stop that one, and only that one, after a timeout.
    """
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    before = powerpoint_pids()
    app = win32com.client.Dispatch("PowerPoint.Application")
    started = sorted(powerpoint_pids() - before) if not before else []
    if status:
        Path(status).write_text(json.dumps({"was_running": bool(before), "started": started}), encoding="utf-8")
    opened, quit_ = [], False
    try:
        yield app, opened
    finally:
        for pres in opened:
            try:
                pres.Close()
            except Exception:  # already closed, or PowerPoint went away: nothing of the user's is touched
                pass
        try:
            if not before and app.Presentations.Count == 0:
                app.Quit()
                quit_ = True
        finally:
            del app
            pythoncom.CoUninitialize()
        deadline = time.time() + 15  # Quit() returns before PowerPoint has gone: the next job must not find it
        while quit_ and started and set(started) & powerpoint_pids() and time.time() < deadline:
            time.sleep(0.25)


def _busy(exc):
    """COM's 'call rejected' / 'retry later': PowerPoint is showing a dialog or the user is typing in it."""
    codes = {-2147418111, -2147417846}  # RPC_E_CALL_REJECTED, RPC_E_SERVERCALL_RETRYLATER
    return bool(codes & set(getattr(exc, "args", ())[:1])) or "rejected" in str(exc).lower()


def _open(app, opened, deck):
    pres = app.Presentations.Open(str(deck), True, False, False)  # read-only, not untitled, no window
    opened.append(pres)
    return pres


def _job_render(app, opened, job):
    """Draw the slides as PNGs and measure every shape: where PowerPoint put it (a table grows rows on its own) and,
    for text, the height of its laid-out text / the height inside the box."""
    pres = _open(app, opened, job["deck"])
    count = pres.Slides.Count
    width = job.get("width", 1280)
    height = round(width * pres.PageSetup.SlideHeight / pres.PageSetup.SlideWidth)
    folder = Path(job["folder"])
    folder.mkdir(parents=True, exist_ok=True)
    pngs, shapes = [], []
    for n in [n for n in (job.get("slides") or range(1, count + 1)) if 1 <= n <= count]:
        slide = pres.Slides(n)
        png = folder / f"slide_{n:02d}.png"
        slide.Export(str(png), "PNG", width, height)
        pngs.append(str(png))
        for shape in slide.Shapes:
            ratio = None
            if shape.HasTextFrame and shape.TextFrame.HasText:
                frame = shape.TextFrame
                room = shape.Height - frame.MarginTop - frame.MarginBottom
                if room > 0:
                    ratio = round(frame.TextRange.BoundHeight / room, 3)
            shapes.append([n, shape.Id, shape.Name, shape.Left, shape.Top, shape.Width, shape.Height, ratio])
    return {"slide_count": count, "pngs": pngs, "shapes": shapes,
            "slide_size_pt": [pres.PageSetup.SlideWidth, pres.PageSetup.SlideHeight]}


def _job_save(app, opened, job):
    """Save a copy in another format: 32 = PDF, 24 = .pptx (PowerPoint's own writing of the same deck)."""
    pres = _open(app, opened, job["deck"])
    count = pres.Slides.Count
    if job["format"] == 24:
        pres.SaveCopyAs(job["out"], 24)
    else:
        pres.SaveAs(job["out"], job["format"])
    return {"slide_count": count}


def _job_count(app, opened, job):
    """Does PowerPoint open it, and how many slides does it see?"""
    return {"slide_count": _open(app, opened, job["deck"]).Slides.Count}


JOBS = {"render": _job_render, "save": _job_save, "count": _job_count}


def _job_main(job_file):
    """The child process: one job, its result written to a file for the parent."""
    job = json.loads(Path(job_file).read_text(encoding="utf-8"))
    result = {"ok": False}
    try:
        with powerpoint(job["status"]) as (app, opened):
            result.update(JOBS[job["action"]](app, opened, job))
        result["ok"] = True
    except Exception as exc:  # reported by the parent, in words
        result.update(error=f"{type(exc).__name__}: {exc}"[:500], busy=_busy(exc))
    Path(job["result"]).write_text(json.dumps(result), encoding="utf-8")


# The child loads this file by its path, so harness/tools/ (with its code.py, image.py ...) never lands on sys.path
CHILD = ("import importlib.util, sys; spec = importlib.util.spec_from_file_location('glassbox_powerpoint_job', "
         "sys.argv[1]); module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module); "
         "module._job_main(sys.argv[2])")


def powerpoint_job(action, deck, timeout=None, **job):
    """
    Run one PowerPoint job (render, save, count) in a child process, on a copy of `deck` with a name of its own: never
    the user's file, never a name PowerPoint already has open. -> the job's result (a dict). Raises
    PowerPointUnavailable when PowerPoint is absent, busy, failing or too slow.
    """
    from ..config import OFFICE
    from .office_common import lock_file
    timeout = timeout or OFFICE["timeout"]
    if not powerpoint_available():
        raise PowerPointUnavailable("PowerPoint is not installed here")
    work = Path(tempfile.mkdtemp(prefix="glassbox_ppt_"))
    try:
        copy = work / f"glassbox_{os.getpid()}_{time.time_ns() % 10 ** 9}{Path(deck).suffix}"
        shutil.copy2(deck, copy)
        job.update(action=action, deck=str(copy), status=str(work / "status.json"), result=str(work / "result.json"))
        (work / "job.json").write_text(json.dumps(job), encoding="utf-8")
        flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        child = [sys.executable, "-I", "-c", CHILD, str(Path(__file__).resolve()), str(work / "job.json")]
        proc = subprocess.Popen(child, stdout=subprocess.DEVNULL, stderr=subprocess.PIPE, creationflags=flags)
        try:
            _, err = proc.communicate(timeout=timeout)
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.communicate()
            status = json.loads((work / "status.json").read_text(encoding="utf-8")) \
                if (work / "status.json").is_file() else {}
            for pid in status.get("started", []):  # the PowerPoint this job started, never the user's
                subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True)
            raise PowerPointUnavailable(
                f"PowerPoint did not finish within {timeout} s (config OFFICE['timeout']), so the job was stopped"
                + ("; the PowerPoint it had started was stopped too" if status.get("started") else
                   "; PowerPoint is yours (it was running before), so it was left running" if status.get("was_running")
                   else "")) from None
        if not (work / "result.json").is_file():
            raise PowerPointUnavailable("the PowerPoint job stopped without a result: "
                                        + err.decode(errors="replace")[-300:])
        result = json.loads((work / "result.json").read_text(encoding="utf-8"))
        if not result["ok"]:
            raise PowerPointUnavailable(("PowerPoint is busy (a dialog is open, or someone is typing in it): "
                                         if result.get("busy") else "PowerPoint failed: ") + result["error"])
        result["read_from_copy"] = bool(lock_file(deck))  # open in PowerPoint (a ~$ file): read from the copy
        return result
    finally:
        shutil.rmtree(work, ignore_errors=True)


def render_slides(path, folder, width=1280, slides=None, timeout=None):
    """
    PowerPoint draws the slides (all, or the 1-based positions in `slides`) as PNGs in `folder`, and measures every
    shape. -> {"slide_count", "pngs": [paths], "measured": {(slide, shape id): ratio},
              "boxes": {(slide, shape id): (left, top, width, height) in cm, as PowerPoint laid it out}}
    """
    result = powerpoint_job("render", path, timeout, folder=str(Path(folder).resolve()), width=width,
                            slides=list(slides) if slides else None)
    cm_per_pt = 2.54 / 72
    result["pngs"] = [Path(p) for p in result["pngs"]]
    result["measured"] = {(n, sid): ratio for n, sid, _, _, _, _, _, ratio in result["shapes"] if ratio is not None}
    result["boxes"] = {(n, sid): tuple(round(v * cm_per_pt, 2) for v in (left, top, w, h))
                       for n, sid, _, left, top, w, h, _ in result["shapes"]}
    result["names"] = {(n, sid): name for n, sid, name, *_ in result["shapes"]}
    return result


def render_deck(path, folder, width=1280):
    """
    PowerPoint opens a copy of the deck read-only and without a window, draws every slide as a PNG, and measures
    every text box: the height of its laid-out text / the height inside the box.
    -> ([png paths], {(slide number, shape id): ratio})
    """
    result = render_slides(path, folder, width)
    return result["pngs"], {key: round(ratio, 2) for key, ratio in result["measured"].items()}


def find_soffice():
    found = shutil.which("soffice") or shutil.which("libreoffice")
    for folder in (os.environ.get("ProgramFiles", ""), os.environ.get("ProgramFiles(x86)", "")):
        candidate = Path(folder) / "LibreOffice" / "program" / "soffice.exe"
        if not found and folder and candidate.exists():
            found = str(candidate)
    return found


def save_as_pdf(deck, pdf):
    """Turn a deck into a PDF with whichever program is installed. -> the name of the program that did it."""
    soffice = find_soffice()
    if powerpoint_available():
        try:
            powerpoint_job("save", deck, out=str(Path(pdf).resolve()), format=32)  # 32 = ppSaveAsPDF
            return "PowerPoint (COM)"
        except PowerPointUnavailable:
            if not soffice:
                raise
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf.parent), str(deck)],
                       check=True, capture_output=True, timeout=180)
        made = pdf.parent / (Path(deck).stem + ".pdf")
        if made != pdf:
            made.replace(pdf)
        return "LibreOffice (soffice --headless)"
    raise RuntimeError("Neither PowerPoint nor LibreOffice is installed here, so the deck cannot be made into a PDF.")


def contact_sheet(pngs, path, columns=3, thumb=420, labels=None):
    """Every slide in one picture, for report.html (plain PIL). labels: one per picture (default 'slide 1', ...)."""
    images = [Image.open(p).convert("RGB") for p in pngs]
    labels = labels or [f"slide {i + 1}" for i in range(len(images))]
    height = round(thumb * images[0].height / images[0].width)
    gap, label = 12, 22
    rows = math.ceil(len(images) / columns)
    sheet = Image.new("RGB", (gap + columns * (thumb + gap), gap + rows * (label + height + gap)), "#d9dce1")
    draw = ImageDraw.Draw(sheet)
    try:
        font = ImageFont.truetype("arial.ttf", 15)
    except OSError:
        font = ImageFont.load_default()
    for i, image in enumerate(images):
        x, y = gap + (i % columns) * (thumb + gap), gap + (i // columns) * (label + height + gap)
        draw.text((x, y + 2), labels[i], fill="#1d1f23", font=font)
        sheet.paste(image.resize((thumb, height), Image.LANCZOS), (x, y + label))
    sheet.save(path)
    return Path(path)
