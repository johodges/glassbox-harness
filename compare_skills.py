"""
With and without: what AGENTS.md and the skills change.

workspace/AGENTS.md and workspace/skills/*/SKILL.md are instructions kept in files (see
harness/instructions.py). This script runs a few fast requests twice, first with neither, then with
both, and measures what came out: the voice that spoke, the words on the slide, the layout of the
section slide, the checks that ran, and what the extra instructions cost in prompt tokens.

    python compare_skills.py                    # all five requests
    python compare_skills.py speech question    # only some of them (names from TASKS)

About four minutes once the models are downloaded. Every run is kept, in
outputs/skills_comparison_<time>/<request>_<without|with>/, next to comparison.md (the table below).
"""

import re
import shutil
import sys
import time

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation

from evaluate import tool_calls
from harness import Agent, HarnessConfig
from harness.config import OUTPUT_DIR


def successful(result, tool):
    """(arguments, result) of every call to `tool` that worked."""
    found = []
    for i, message in enumerate(result.messages):
        for call in message.get("tool_calls", []):
            reply = result.messages[i + 1]["content"] if i + 1 < len(result.messages) else {}
            if call["function"]["name"] == tool and isinstance(reply, dict) and reply.get("status") == "ok":
                found.append((call["function"]["arguments"], reply))
    return found


def deck(result, word):
    return next(Presentation(f) for f in result.files if f.endswith(".pptx") and word in f)


# ---- one check per request: what to measure in the output ------------------------------------------------
def voice(result):
    spoken = successful(result, "generate_speech")
    return f"voice {spoken[0][0].get('voice') or 'af_heart (the default)'}" if spoken else "nothing spoken"


def picture(result):
    drawn = successful(result, "generate_image")
    if not drawn:
        return "no picture"
    style = "flat vector" in drawn[0][0]["prompt"].lower()
    text = f"{len(drawn)} picture(s), {'house style' if style else 'no house style'} in the prompt"
    looked = successful(result, "look_at_image")
    return text + (f"; checked: \"{looked[0][1]['answer'][:60]}\"" if looked else "; not checked")


def bullets(result):
    written = successful(result, "update_powerpoint")
    if not written:
        return "slide not written"
    number = int(written[-1][0].get("slide") or 0)  # the slide the tool really changed (0 = a new last slide)
    prs = deck(result, "project_update")
    slide = prs.slides[number - 1] if number else prs.slides[-1]
    lines = [p.text.strip() for shape in slide.shapes if shape.has_text_frame and shape != slide.shapes.title
             for p in shape.text_frame.paragraphs if p.text.strip()]
    if not lines:
        return "slide not written"
    words = [len(line.split()) for line in lines]
    # a bullet "keeps a fact" if it has a number that is in the passages search_references returned
    notes = " ".join(p["text"] for _, reply in successful(result, "search_references") for p in reply["passages"])
    facts = sum(any(n in notes for n in re.findall(r"\d+(?:\.\d+)?", line)) for line in lines)
    first = next((m["tool_calls"][0]["function"]["name"] for m in result.messages if m.get("tool_calls")), None)
    order = "searched first" if first == "search_references" else "wrote before searching"
    return f"{order}; {len(lines)} bullets of {min(words)}-{max(words)} words, {facts} with a number from the notes"


def dividers(result):
    layouts = [s.slide_layout.name.strip() for s in deck(result, "old_format_deck").slides
               if "Divider" in s.slide_layout.name or "Section" in s.slide_layout.name]
    return f"section slide on {', '.join(layouts) or 'no divider'}"


def nothing(result):
    return result.answer.strip()[:50]


TASKS = [  # (short name, request, what changes it, the check)
    ("speech", "Read aloud: hello class", "AGENTS.md (the team's voice)", voice),
    ("slide", "Using the reference notes, update slide 2 of workspace/project_update.pptx with three bullet "
              "points on how speech is fitted to a length.", "skill: slide-bullets", bullets),
    ("convert", "Convert workspace/old_format_deck.pptx to the workshop template.", "skill: deck-conversion",
     dividers),
    ("picture", "Draw a red bicycle leaning on a brick wall.", "AGENTS.md (house style) + skill: image-check",
     picture),
    ("question", "What is the capital of France?", "nothing matches: AGENTS.md only", nothing),
]


def run(agent, request, with_files):
    agent.config.use_agents_md = agent.config.use_skills = with_files
    result = agent.run(request)
    prompt_tokens = next(ev.details["num_tokens"] for ev in result.events if ev.stage == "PROMPT ASSEMBLY")
    skills = next((ev.details["loaded"] for ev in result.events if ev.stage == "SKILLS"), [])
    seconds = sum(ev.seconds for ev in result.events if ev.stage != "LOAD MODEL")  # loading happens once
    return result, prompt_tokens, skills, seconds


def main():
    config = HarnessConfig(verbosity=0, color=False)
    if hasattr(config, "ask_before_acting"):
        config.ask_before_acting = False  # nothing here runs code, but never wait for a y/n
    agent = Agent(config)
    folder = OUTPUT_DIR / time.strftime("skills_comparison_%Y%m%d_%H%M%S")
    folder.mkdir(parents=True)
    rows = []
    wanted = sys.argv[1:] or [task[0] for task in TASKS]
    for name, request, cause, check in (task for task in TASKS if task[0] in wanted):
        row = {"name": name, "request": request, "cause": cause}
        for setting, with_files in (("without", False), ("with", True)):
            result, tokens, skills, seconds = run(agent, request, with_files)
            try:
                outcome = check(result)
            except Exception as exc:  # a check that can't find its file is a result too
                outcome = f"could not check: {type(exc).__name__}: {exc}"
            _, executed = tool_calls(result)
            row[setting] = {"outcome": outcome, "tokens": tokens, "skills": skills, "tools": executed,
                            "seconds": seconds}
            shutil.move(str(result.run_dir), str(folder / f"{name}_{setting}"))
            print(f"{name:<9} {setting:<8} {tokens:>5} tokens  {seconds:5.1f}s  skills {skills or '-'}  "
                  f"tools {executed}  ->  {outcome}", flush=True)
        rows.append(row)

    lines = ["| request | what changes it | without AGENTS.md and skills | with them | prompt tokens |",
             "|---|---|---|---|---|"]
    for r in rows:
        lines.append(f"| {r['request']} | {r['cause']} | {r['without']['outcome']} | {r['with']['outcome']} | "
                     f"{r['without']['tokens']} -> {r['with']['tokens']} |")
    (folder / "comparison.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n" + "\n".join(lines) + f"\n\nRuns and comparison.md: {folder}")


if __name__ == "__main__":
    main()
