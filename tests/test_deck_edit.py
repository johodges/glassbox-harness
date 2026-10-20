"""deck_edit.py's edit_deck: the ops a small model writes, the run-preserving replace, VERIFY (untouched slides keep
their fingerprints, an op that breaks that is undone), APPROVAL for a delete, CHECK ARGS, copies never sharing a
chart, the package checks, and the version ledger. No GPU, no LLM, no PowerPoint."""

import hashlib
import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from pptx import Presentation
from test_deck_model import DEMO, small_deck

from harness.tools import deck_edit
from harness.tools.base import Refusal
from harness.tools.deck_edit import (DiffDecksTool, EditDeckTool, package_problems, parse_ops, same_words,
                                     shared_parts)
from harness.tools.deck_model import deck_view, marked

BRIEFING = "workspace/deck_demo/briefing.pptx"
RESULTS = "workspace/deck_demo/results.pptx"


@pytest.fixture
def deck_ctx(ctx, tmp_path):
    ctx.facts["versions_dir"] = str(tmp_path / "versions")  # the version ledger goes to the test's folder
    return ctx


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def edit(ctx, path, ops):
    tool = EditDeckTool()
    return tool.run(ctx, **tool.validate({"path": path, "ops": ops}))


def untouched_same(before, after, changed_ids):
    """The test's own check, independent of VERIFY: every other slide still has its fingerprint."""
    old = {s["id"]: s["fingerprint"] for s in before["slides"]}
    new = {s["id"]: s["fingerprint"] for s in after["slides"]}
    return all(new[i] == f for i, f in old.items() if i in new and i not in changed_ids)


def test_ops_in_the_shapes_a_small_model_writes():
    ops = parse_ops('[{"op": "move", "slide": "5", "after": "slide 2"}, {"notes": 3, "text": "Hi there"}, '
                    '{"delete": 7}, {"slide": 4, "find": "12 tests", "replace": "14 tests"}, '
                    '{"op": "duplicate", "slide": 4}, {"op": "move", "slide": 6, "before": 3}, "nonsense"]')
    assert [op["op"] for op in ops] == ["move", "notes", "delete", "text", "duplicate", "move", "unreadable"]
    assert (ops[0]["slide"], ops[0]["after"]) == (5, 2) and ops[1]["mode"] == "append"
    assert ops[4]["after"] == 4 and ops[5]["after"] == 2
    texts = parse_ops(["move slide 5 after slide 2", 'slide 4: "12 tests" -> "14 tests"', "delete slide 7",
                       "notes 3: Say why", "import 2-3 from workspace/deck_demo/results.pptx after 6"])
    assert texts[1] == {"op": "text", "slide": 4, "find": "12 tests", "replace": "14 tests", "shape": None}
    assert texts[4]["slides"] == [2, 3] and texts[4]["after"] == 6 and texts[3]["text"] == "Say why"
    one = EditDeckTool().validate({"path": BRIEFING, "op": "move", "slide": 5, "after": 2})  # no ops list at all
    assert one["ops"] == [{"op": "move", "slide": 5, "after": 2}]
    with pytest.raises(ValueError):
        EditDeckTool().validate({"path": BRIEFING, "ops": "gibberish"})


def test_same_words_trims_whole_words_only():
    assert same_words("12 tests", "14 tests") == (0, 6)
    assert same_words("has a bold", "has one bold") == (4, 5)
    assert same_words("a bold", "an emphasised") == (0, 0)


def test_replace_keeps_the_formatting_of_the_words_around(deck_ctx):
    deck = small_deck(deck_ctx.run_dir / "plain.pptx")
    out = edit(deck_ctx, str(deck), [{"op": "text", "slide": 1, "find": "has a bold", "replace": "has one bold"},
                                     {"op": "text", "slide": 2, "find": "a bold word", "replace": "an odd one"},
                                     {"op": "text", "slide": 3, "find": "Second point", "replace": "**Key** point"},
                                     {"op": "text", "slide": 4, "find": "Point 4", "replace": "Point four"}])
    view = deck_view(out["artifacts"][0]["path"])
    first = lambda n, k=0: marked(view["slides"][n]["shapes"][1]["paragraphs"][k]["runs"])
    assert first(0) == "Point 1 has one **bold** word"   # only "a" changed; "bold" stays bold
    assert first(1) == "Point 2 has an odd one"          # the first replaced run's format (plain) for every word
    assert first(2, 1) == "**Key** point of slide 3"     # **...** in the new words makes them bold
    assert first(3) == "Point four has a **bold** word"
    assert "several formats" in " ".join(out["observations"]["done"])


def test_move_and_notes_on_the_demo_leave_the_rest_alone(deck_ctx):
    original = sha(BRIEFING)
    before = deck_view(BRIEFING)
    out = edit(deck_ctx, BRIEFING, [{"op": "move", "slide": 5, "after": 2},
                                    {"op": "notes", "slide": 3, "text": "Start with the hood."}])
    new = Path(out["artifacts"][0]["path"])
    after = deck_view(new)
    assert [s["title"] for s in after["slides"]][:4] == ["Compartment fire tests", "Test programme", "Observations",
                                                         "Instrumentation"]
    assert after["slides"][3]["notes"] == "Slide 7 has the full table of peak values.\nStart with the hood."
    assert untouched_same(before, after, {before["slides"][2]["id"]})
    assert package_problems(new) == [] and sha(BRIEFING) == original  # a new file; the user's deck is untouched
    assert out["observations"]["changed_slides"] == [4] and new.parent == deck_ctx.run_dir


def test_text_op_keeps_bold_and_the_next_tool_sees_the_edit(deck_ctx):
    out = edit(deck_ctx, BRIEFING, [{"op": "text", "slide": 4, "find": "12 tests", "replace": "14 tests"}])
    view = deck_view(out["artifacts"][0]["path"])
    assert marked(view["slides"][3]["shapes"][1]["paragraphs"][0]["runs"]) == \
        "We have completed **14 tests** of the 16 planned"
    assert deck_edit.deck_file(BRIEFING, deck_ctx)[1] == Path(out["artifacts"][0]["path"])  # check_deck reads this


def test_import_restyles_and_copies_the_chart(deck_ctx):
    out = edit(deck_ctx, BRIEFING, [{"op": "import", "from": RESULTS, "slides": "2-3", "after": 6}])
    new = Path(out["artifacts"][0]["path"])
    view = deck_view(new)
    assert [s["title"] for s in view["slides"]][5:9] == ["Design fire curves", "Peak heat release rates, series B",
                                                         "Time to flashover, series B", "Peak values by test"]
    assert view["slides"][6]["layout"] == "Title and Content" and view["template"]["slide_size_cm"][0] == 33.87
    chart = next(sh for sh in view["slides"][7]["shapes"] if sh["kind"] == "chart")
    assert chart["chart_text"] == ["Time to flashover (s)"]
    assert not any(sh.get("ph") == "sldNum" for sh in view["slides"][6]["shapes"])  # the deck shows no numbers
    # duplicating the imported chart slide copies the chart: two charts, not one shared by two slides
    deck_ctx.copies.clear()
    again = edit(deck_ctx, str(new), [{"op": "duplicate", "slide": 8}])
    twin = deck_view(again["artifacts"][0]["path"])
    assert twin["slides"][8]["fingerprint"] == twin["slides"][7]["fingerprint"]
    assert shared_parts(again["artifacts"][0]["path"], [s["id"] for s in twin["slides"]]) == []


def test_delete_asks_a_person_and_no_means_no(deck_ctx):
    deck_ctx.config.ask_before_acting = True
    asked = []
    deck_ctx.agent = SimpleNamespace(config=deck_ctx.config, ask_user=lambda q: asked.append(q) or "n")
    with pytest.raises(Refusal):
        edit(deck_ctx, BRIEFING, [{"op": "delete", "slide": 7}])
    assert "Peak values by test" in asked[0]
    out = edit(deck_ctx, BRIEFING, [{"op": "delete", "slide": 7}, {"op": "notes", "slide": 1, "text": "Hello."}])
    assert out["observations"]["not_done"] == ["delete slide 7: a person said no"]
    assert deck_view(out["artifacts"][0]["path"])["slide_count"] == 8
    deck_ctx.agent.ask_user = lambda q: "y"
    out = edit(deck_ctx, BRIEFING, [{"op": "delete", "slide": 7}])
    assert "Peak values by test" not in [s["title"] for s in deck_view(out["artifacts"][0]["path"])["slides"]]


def test_verify_undoes_an_op_that_touches_another_slide(deck_ctx, monkeypatch):
    real = deck_edit.APPLY["notes"]

    def sloppy(batch, plan):  # a buggy op: writes the notes, and also changes another slide's title
        real(batch, plan)
        batch.prs.slides[0].shapes.title.text_frame.text = "Oops"

    monkeypatch.setitem(deck_edit.APPLY, "notes", sloppy)
    out = edit(deck_ctx, BRIEFING, [{"op": "notes", "slide": 3, "text": "More."},
                                    {"op": "move", "slide": 5, "after": 2}])
    view = deck_view(out["artifacts"][0]["path"])
    assert view["slides"][0]["title"] == "Compartment fire tests" and view["slides"][2]["title"] == "Observations"
    assert out["observations"]["done"] == ["move slide 5 after slide 2"]
    assert "changed, but no op named it" in out["observations"]["not_done"][0]


def test_an_op_that_crashes_half_way_is_left_out(deck_ctx, monkeypatch):
    def crash(batch, plan):  # changes another slide, then fails
        batch.prs.slides[1].shapes.title.text_frame.text = "Half done"
        raise KeyError("something python-pptx did not expect")

    monkeypatch.setitem(deck_edit.APPLY, "notes", crash)
    out = edit(deck_ctx, BRIEFING, [{"op": "notes", "slide": 3, "text": "More."},
                                    {"op": "move", "slide": 5, "after": 2}])
    view = deck_view(out["artifacts"][0]["path"])
    assert view["slides"][1]["title"] == "Test programme" and view["slides"][2]["title"] == "Observations"
    assert "KeyError" in out["observations"]["not_done"][0]


def test_words_must_occur_once_and_quoted_words_win(deck_ctx):
    out = edit(deck_ctx, BRIEFING, [{"op": "text", "slide": 2, "find": "tests", "replace": "trials"},
                                    {"op": "text", "slide": 4, "find": "99 tests", "replace": "x"},
                                    {"op": "notes", "slide": 1, "text": "Fine."}])
    assert [n.split(":")[0] for n in out["observations"]["not_done"]] == ["slide 2", "slide 4"]
    tool = EditDeckTool()
    deck_ctx.request = ('In workspace/deck_demo/briefing.pptx, move slide 5 after slide 2, and add speaker notes to '
                        'slide 3: "Start with the hood." Then change "12 tests" to "14 tests".')
    args = tool.validate({"path": BRIEFING, "ops": [{"op": "move", "slide": 2, "after": 5},
                                                    {"op": "notes", "slide": 4, "text": "Start with hood"},
                                                    {"op": "text", "slide": 4, "find": "12", "replace": "14"},
                                                    {"op": "delete", "slide": 7}]})
    fixed, summary, note = tool.check_request(args, deck_ctx)
    assert fixed["ops"] == [{"op": "move", "slide": 5, "after": 2},
                            {"op": "notes", "slide": 3, "text": "Start with the hood.", "mode": "append"},
                            {"op": "text", "slide": 4, "find": "12 tests", "replace": "14 tests", "shape": None}]
    assert note and "delete" in summary


def test_insert_duplicate_and_package_checks(deck_ctx, tmp_path):
    out = edit(deck_ctx, BRIEFING, [{"op": "insert", "after": 8, "title": "Questions",
                                     "bullets": ["What did we miss?", "  Which **test** next?"]},
                                    {"op": "insert", "after": 1, "layout": "section header light", "title": "Part 1"},
                                    {"op": "duplicate", "slide": 6}])
    view = deck_view(out["artifacts"][0]["path"])
    titles = [s["title"] for s in view["slides"]]
    assert titles[1] == "Part 1" and view["slides"][1]["layout"] == "Section Header Light"
    assert titles[-1] == "Questions" and titles[6:9] == ["Design fire curves", "Design fire curves",
                                                         "Peak values by test"]
    body = view["slides"][-1]["shapes"][1]["paragraphs"]
    assert body[1]["level"] == 1 and marked(body[1]["runs"]) == "Which **test** next?"
    assert package_problems(DEMO / "briefing.pptx") == []
    shutil.copy(DEMO / "briefing.pptx", tmp_path / "broken.pptx")
    import zipfile
    with zipfile.ZipFile(DEMO / "briefing.pptx") as zin, zipfile.ZipFile(tmp_path / "broken.pptx", "w") as zout:
        for item in zin.infolist():
            if item.filename != "ppt/media/image2.png":
                zout.writestr(item, zin.read(item.filename))
    assert any("not in the package" in p for p in package_problems(tmp_path / "broken.pptx"))


def test_a_deck_open_in_powerpoint_is_read_from_a_copy(deck_ctx):
    deck = small_deck(deck_ctx.run_dir / "open.pptx")
    (deck_ctx.run_dir / "~$open.pptx").write_bytes(b"owner")  # PowerPoint's lock file: the deck is open
    original, readable = deck_edit.deck_file(str(deck), deck_ctx)
    assert original == deck.resolve() and readable == deck_ctx.run_dir / "inputs" / "open.pptx"
    before = sha(deck)
    out = edit(deck_ctx, str(deck), [{"op": "notes", "slide": 1, "text": "Fine."}])
    assert sha(deck) == before and Path(out["artifacts"][0]["path"]).parent == deck_ctx.run_dir


def test_your_edits_become_the_latest_version(deck_ctx):
    """'Use mine as the latest': the harness's version, edited by the user, is diffed against it and edited next."""
    deck = small_deck(deck_ctx.run_dir / "talk.pptx")
    first = edit(deck_ctx, str(deck), [{"op": "notes", "slide": 2, "text": "From the harness."}])
    shutil.copy(first["artifacts"][0]["path"], deck)  # the user takes the harness's version ...
    prs = Presentation(deck)                           # ... and edits a title in PowerPoint
    prs.slides[3].shapes.title.text_frame.text = "Slide four, renamed"
    prs.save(deck)
    deck_ctx.copies.clear()  # a new request
    diff = DiffDecksTool().run(deck_ctx, new=str(deck))["observations"]
    assert "the version ledger" in diff["compared"] and "Slide four, renamed" in diff["changed"][0]
    second = edit(deck_ctx, str(deck), [{"op": "notes", "slide": 1, "text": "More."}])
    kept = second["observations"]["your_edits_kept"]
    assert "new latest" in kept[0] and "Slide four, renamed" in " ".join(kept)
    assert deck_view(second["artifacts"][0]["path"])["slides"][3]["title"] == "Slide four, renamed"
