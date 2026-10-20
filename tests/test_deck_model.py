"""deck_model.py: deck_view sees everything (groups, mc:AlternateContent, marks, inherited boxes, broken pictures),
and a slide's fingerprint changes with its content, never with part names or a save. No GPU, no LLM, no PowerPoint.
small_deck() is shared with the other test_deck_*.py files."""

import io
import zipfile
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Cm, Pt

from harness.tools.deck_model import Package, canonical, deck_view, marked, plain

DEMO = Path(__file__).resolve().parent.parent / "workspace" / "deck_demo"


def png(colour="red", size=(40, 30)):
    stream = io.BytesIO()
    Image.new("RGB", size, colour).save(stream, "PNG")
    stream.seek(0)
    return stream


def small_deck(path, slides=4, notes=True):
    """A plain deck: slide k has the title 'Slide k', two bullets (the first with a bold word) and notes."""
    prs = Presentation()
    for k in range(1, slides + 1):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Slide {k}"
        frame = s.placeholders[1].text_frame
        p = frame.paragraphs[0]
        for words, bold in ((f"Point {k} has a ", False), ("bold", True), (" word", False)):
            run = p.add_run()
            run.text, run.font.size = words, Pt(24)
            run.font.bold = bold or None
        frame.add_paragraph().text = f"Second point of slide {k}"
        if notes:
            s.notes_slide.notes_text_frame.text = f"Notes for slide {k}."
    prs.save(path)
    return Path(path)


def test_view_reads_marks_groups_and_alternate_content(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    s.shapes.title.text = "Marks"
    box = s.shapes.add_textbox(Cm(1), Cm(5), Cm(10), Cm(3))
    p = box.text_frame.paragraphs[0]
    for words, mark in (("plain ", None), ("bold", "b"), (" and ", None), ("italic", "i"), (" m", None), ("2", "sup")):
        run = p.add_run()
        run.text = words
        if mark == "b":
            run.font.bold = True
        elif mark == "i":
            run.font.italic = True
        elif mark == "sup":
            run.font._rPr.set("baseline", "30000")
    link = box.text_frame.add_paragraph().add_run()
    link.text = "a link"
    link.hyperlink.address = "https://example.org/fire"
    group = s.shapes.add_group_shape()
    group.shapes.add_textbox(Cm(2), Cm(12), Cm(4), Cm(1)).text_frame.text = "in a group"
    group._element.xpath("./p:grpSpPr/a:xfrm")[0]  # python-pptx writes the group's own box
    s.shapes._spTree.append(parse_xml(  # an equation as PowerPoint saves one: a Choice with a14, a picture Fallback
        f'<mc:AlternateContent {nsdecls("p", "a")} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility'
        f'/2006" xmlns:a14="http://schemas.microsoft.com/office/drawing/2010/main"><mc:Choice Requires="a14">'
        f'<p:sp><p:nvSpPr><p:cNvPr id="90" name="Equation 1"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr><p:spPr>'
        f'<a:xfrm><a:off x="360000" y="360000"/><a:ext cx="720000" cy="360000"/></a:xfrm></p:spPr><p:txBody>'
        f'<a:bodyPr/><a:p><a14:m><a:t>x</a:t></a14:m></a:p></p:txBody></p:sp></mc:Choice><mc:Fallback/>'
        f'</mc:AlternateContent>'))
    prs.save(tmp_path / "marks.pptx")
    slide = deck_view(tmp_path / "marks.pptx")["slides"][0]
    shapes = {sh["name"]: sh for sh in slide["shapes"]}
    text = next(sh for sh in slide["shapes"] if sh["kind"] == "text" and sh.get("ph") is None and not sh.get("group"))
    assert marked(text["paragraphs"][0]["runs"]) == "plain **bold** and *italic* m^2^"
    assert text["paragraphs"][1]["runs"][0]["link"] == "https://example.org/fire"
    assert text["box"] == [1.0, 5.0, 10.0, 3.0]
    child = next(sh for sh in slide["shapes"] if sh.get("group"))
    assert plain(child["paragraphs"][0]["runs"]) == "in a group" and child["box"][1] == 12.0  # slide coordinates
    assert shapes["Equation 1"]["kind"] == "equation" and shapes["Equation 1"]["alternate"]
    title = shapes["Title 1"]
    assert title["ph"] == "title" and title["inherited_box"] and title["box"]  # the layout's position


def test_broken_picture_is_reported(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(png(), Cm(1), Cm(1))
    prs.save(tmp_path / "pic.pptx")
    assert deck_view(tmp_path / "pic.pptx")["slides"][0]["shapes"][0]["image"] not in ("broken", "none")
    with zipfile.ZipFile(tmp_path / "pic.pptx") as zin, zipfile.ZipFile(tmp_path / "broken.pptx", "w") as zout:
        for item in zin.infolist():  # PowerPoint sometimes saves a picture whose file is missing
            if not item.filename.startswith("ppt/media/"):
                zout.writestr(item, zin.read(item.filename))
    assert deck_view(tmp_path / "broken.pptx")["slides"][0]["shapes"][0]["image"] == "broken"


def test_fingerprint_follows_content_not_part_names(tmp_path):
    deck = small_deck(tmp_path / "a.pptx")
    before = {s["id"]: s["fingerprint"] for s in deck_view(deck)["slides"]}
    prs = Presentation(deck)  # move the last slide first: python-pptx renames slide4.xml to slide1.xml ...
    ids = prs.slides._sldIdLst
    last = ids[-1]
    ids.remove(last)
    ids.insert(0, last)
    prs.part.rename_slide_parts([e.get(qn("r:id")) for e in ids])
    prs.save(tmp_path / "b.pptx")
    moved = deck_view(tmp_path / "b.pptx")
    assert moved["slides"][0]["part"] == "/ppt/slides/slide1.xml" and moved["slides"][0]["title"] == "Slide 4"
    assert {s["id"]: s["fingerprint"] for s in moved["slides"]} == before
    prs = Presentation(deck)
    prs.slides[1].placeholders[1].text_frame.paragraphs[0].runs[1].font.italic = True  # one run, one mark
    prs.save(tmp_path / "c.pptx")
    after = {s["id"]: s["fingerprint"] for s in deck_view(tmp_path / "c.pptx")["slides"]}
    assert [k for k in before if before[k] != after[k]] == [list(before)[1]]
    prs = Presentation(deck)
    prs.slides[2].notes_slide.notes_text_frame.text = "Other notes."  # notes are part of a slide's fingerprint
    prs.save(tmp_path / "d.pptx")
    notes = {s["id"]: s["fingerprint"] for s in deck_view(tmp_path / "d.pptx")["slides"]}
    assert [k for k in before if before[k] != notes[k]] == [list(before)[2]]


def test_canonical_ignores_prefixes_attribute_order_and_resolves_ids():
    a = parse_xml(f'<a:p {nsdecls("a")}><a:r><a:rPr b="1" lang="en-US"/><a:t>x</a:t></a:r></a:p>')
    b = parse_xml('<d:p xmlns:d="http://schemas.openxmlformats.org/drawingml/2006/main"><d:r>'
                  '<d:rPr lang="en-US" b="1"/><d:t>x</d:t></d:r></d:p>')
    assert canonical(a) == canonical(b)
    pic = lambda rid: parse_xml(f'<a:blip {nsdecls("a", "r")} r:embed="{rid}"/>')
    same = {"rId2": "image:abc", "rId7": "image:abc"}
    assert canonical(pic("rId2"), same.get) == canonical(pic("rId7"), same.get)


def test_demo_decks_view(tmp_path):
    view = deck_view(DEMO / "briefing.pptx")
    assert view["slide_count"] == 8 and view["template"]["slide_size_cm"] == [33.87, 19.05]
    assert [s["title"] for s in view["slides"]][4:7] == ["Observations", "Design fire curves", "Peak values by test"]
    table = next(sh for sh in view["slides"][6]["shapes"] if sh["kind"] == "table")
    assert table["table"][0][0] == "Test" and len(table["table"]) == 7
    assert "**12 tests**" in marked(view["slides"][3]["shapes"][1]["paragraphs"][0]["runs"])
    chart = next(sh for sh in deck_view(DEMO / "results.pptx")["slides"][2]["shapes"] if sh["kind"] == "chart")
    assert chart["chart_text"] == ["Time to flashover (s)"] and chart["chart_data"]
    assert Package(DEMO / "briefing.pptx").content_type("/ppt/slides/slide1.xml").endswith("slide+xml")
