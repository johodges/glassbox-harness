"""docx_markup.py, the writer (OFFICE_COLLAB_DESIGN.md, W5 and W6): the round trip of a word-level redline, accept and
reject against the reader's prediction, paragraph marks and table rows, ids, comments with threads, and the zip rules.
No GPU, no LLM, no Word."""

import copy
import zipfile

import pytest
from docx import Document
from lxml import etree
from test_word_model import hand_written

from harness.tools import docx_markup as dm
from harness.tools import word_model as wm
from harness.tools.word_model import w

ME = "Llama 3.2 3B Instruct, via. Glass-box harness (for Test Person)"


def make_doc(path, extra=None):
    """A small document with mixed formatting in one paragraph, a table, and an end."""
    doc = Document()
    doc.add_heading("Plan", 1)
    p = doc.add_paragraph("The sprinkler must activate within ")
    p.add_run("5").bold = True
    p.add_run(" minutes of ignition, or the test stops.")
    doc.add_paragraph("Data are logged and recorded every second.")
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["Test", "HRR"], ["1", "300 kW"]]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    doc.add_paragraph("Last paragraph.")
    if extra:
        extra(doc)
    doc.save(path)
    return path


def find(pkg, text):
    for p in wm.paragraphs_of(pkg.document):
        at = wm.accepted_text(p).find(text)
        if at >= 0:
            return p, at, at + len(text)
    raise AssertionError(text)


def ours(writer, action):
    ids = set(writer.revision_ids)
    return lambda el: action if el.get(w("id")) in ids else "keep"


def test_word_level_redline_round_trip(tmp_path):
    pkg = dm.Package(make_doc(tmp_path / "a.docx"))
    before = copy.deepcopy(pkg.document)
    writer = dm.Writer(pkg, ME, "GBH", "2026-10-20T09:00:00Z")
    p, a, b = find(pkg, "within 5 minutes of ignition")
    made, (first, last) = dm.tracked_replace(p, a, b, "within 3 minutes after ignition", writer)
    assert [el.tag for el in made] == [w("del"), w("ins"), w("del"), w("ins")]  # 'of' -> 'after' first, then 5 -> 3
    # W6.2 and W6.4: rejecting the harness's revisions gives the original back, text AND formatting, paragraph by
    # paragraph; W6.3: accepting them gives the intended text
    assert wm.lines(pkg.document, ours(writer, "reject")) == wm.lines(before, wm.KEEP)
    assert wm.accepted_text(p) == "The sprinkler must activate within 3 minutes after ignition, or the test stops."
    # the new 3 replaces the bold 5, so it is bold too; every revision has the author and the one timestamp
    ins = [el for el in made if el.tag == w("ins")]
    assert any(r.find(f"{w('rPr')}/{w('b')}") is not None for r in ins[1].iter(w("r")))
    assert {el.get(w("author")) for el in made} == {ME} and {el.get(w("date")) for el in made} == {writer.date}
    assert first is not None and last is not None and p.index(first) <= p.index(last)
    pkg.touch(pkg.names["document"])
    assert dm.structure_problems(pkg) == []


def test_deleted_text_is_deltext_with_preserved_spaces(tmp_path):
    pkg = dm.Package(make_doc(tmp_path / "a.docx"))
    writer = dm.Writer(pkg, ME)
    p, a, b = find(pkg, "logged and recorded every")
    dm.tracked_replace(p, a, b, "logged every", writer)
    dele = next(p.iter(w("del")))
    texts = [t.text for t in dele.iter(w("delText"))]
    assert "".join(texts) == "and recorded " and not list(dele.iter(w("t")))
    assert all(t.get(wm.XML_SPACE) == "preserve" for t in dele.iter(w("delText")) if t.text != t.text.strip())


def test_inserted_words_take_the_formatting_before_them(tmp_path):
    pkg = dm.Package(make_doc(tmp_path / "a.docx"))
    writer = dm.Writer(pkg, ME)
    p, a, b = find(pkg, "5 minutes")
    made, _ = dm.tracked_replace(p, a, b, "5 whole minutes", writer)
    ins = made[0]
    assert wm.run_text(ins[0]) == "whole " and ins[0].find(f"{w('rPr')}/{w('b')}") is None  # after the space, plain


def test_the_writer_refuses_what_it_should_not_touch(tmp_path):
    pkg = dm.Package(hand_written(tmp_path / "hand.docx"))
    writer = dm.Writer(pkg, ME)
    p, a, b = find(pkg, "grew quickly")
    with pytest.raises(dm.Refused) as refused:  # inside A. Reviewer's pending insertion
        dm.tracked_replace(p, a, b, "grew", writer)
    assert refused.value.needs_person and "A. Reviewer" in str(refused.value)
    p, a, b = find(pkg, "fire grew")  # spans plain text and the pending insertion
    with pytest.raises(dm.Refused):
        dm.tracked_replace(p, a, b, "fire rose", writer)


def test_hyperlinks_and_fields_are_refused(tmp_path):
    body = """<w:p><w:hyperlink r:id="rId9"
      xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <w:r><w:t>the link text</w:t></w:r></w:hyperlink>
      <w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> DATE </w:instrText></w:r>
      <w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>20 October</w:t></w:r>
      <w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>"""
    pkg = dm.Package(hand_written(tmp_path / "h.docx", body=body))
    writer = dm.Writer(pkg, ME)
    p, a, b = find(pkg, "link")
    with pytest.raises(dm.Refused, match="hyperlink"):
        dm.tracked_replace(p, a, b, "url", writer)
    p, a, b = find(pkg, "20 October")
    with pytest.raises(dm.Refused, match="field"):
        dm.tracked_replace(p, a, b, "21 October", writer)


@pytest.mark.parametrize("decide", [wm.ACCEPT, wm.REJECT,
                                    wm.deciding(authors=["A. Reviewer"], action="accept", others="keep"),
                                    wm.deciding(authors=["B. Checker"], action="accept", others="reject"),
                                    wm.deciding(ids=["2", "6"], action="reject", others="keep")])
def test_resolve_does_what_the_reader_predicts(tmp_path, decide):
    """The writer's accept/reject and the reader's lines() are two separate implementations: they must agree on every
    decision, paragraph marks and rows included."""
    root = wm.read_parts(hand_written(tmp_path / "hand.docx"))["document"]
    predicted = wm.lines(root, decide, formatting=False)
    done = copy.deepcopy(root)
    dm.resolve(done, decide)
    assert wm.lines(done, wm.KEEP, formatting=False) == predicted


def test_reject_of_formatting_restores_it(tmp_path):
    root = wm.read_parts(hand_written(tmp_path / "hand.docx"))["document"]
    done = copy.deepcopy(root)
    dm.resolve(done, wm.REJECT)
    hot = next(r for r in done.iter(w("r")) if wm.run_text(r) == "Hot")
    assert hot.find(f"{w('rPr')}/{w('b')}") is None and hot.find(f"{w('rPr')}/{w('rPrChange')}") is None
    assert not list(done.iter(*dm.REVISION_TAGS))


def test_paragraphs_inserted_and_deleted(tmp_path):
    pkg = dm.Package(make_doc(tmp_path / "a.docx"))
    before = copy.deepcopy(pkg.document)
    writer = dm.Writer(pkg, ME)
    p, _, _ = find(pkg, "Data are logged")
    dm.insert_paragraphs_after(p, ["First new.", "Second new."], writer)
    last, _, _ = find(pkg, "Last paragraph.")
    dm.insert_paragraphs_after(last, ["After the end."], writer)  # the body's last paragraph: its own mark is tracked
    q, _, _ = find(pkg, "The sprinkler must")
    dm.delete_paragraph(q, writer)
    accepted = wm.texts(pkg.document, wm.ACCEPT)
    assert accepted == ["Plan", "Data are logged and recorded every second.", "First new.", "Second new.", "Test",
                        "HRR", "1", "300 kW", "Last paragraph.", "After the end."]
    assert wm.lines(pkg.document, ours(writer, "reject")) == wm.lines(before, wm.KEEP)
    pkg.touch(pkg.names["document"])
    assert dm.structure_problems(pkg) == []


def test_deleting_the_last_paragraph_tracks_the_mark_before_it(tmp_path):
    pkg = dm.Package(make_doc(tmp_path / "a.docx", extra=lambda doc: doc.add_paragraph("Really last.")))
    before = copy.deepcopy(pkg.document)
    writer = dm.Writer(pkg, ME)
    last, _, _ = find(pkg, "Really last.")
    made = dm.delete_paragraph(last, writer)
    assert [el.tag for el in made] == [w("del"), w("del")]  # its text, and the mark of the paragraph before it
    assert wm.texts(pkg.document, wm.ACCEPT)[-1] == "Last paragraph."  # no stray empty paragraph after accepting
    assert wm.lines(pkg.document, ours(writer, "reject")) == wm.lines(before, wm.KEEP)


def test_a_paragraph_before_a_table_goes_and_the_one_after_stays(tmp_path):
    """As Word does it (checked through COM): an emptied paragraph whose mark is deleted before a table goes; the
    body's last paragraph, after a table, keeps its mark, so an empty paragraph is left."""
    pkg = dm.Package(make_doc(tmp_path / "a.docx"))
    writer = dm.Writer(pkg, ME)
    p, _, _ = find(pkg, "Data are logged")
    dm.delete_paragraph(p, writer)
    last, _, _ = find(pkg, "Last paragraph.")
    dm.delete_paragraph(last, writer)
    accepted = wm.texts(pkg.document, wm.ACCEPT)
    assert accepted[2:] == ["Test", "HRR", "1", "300 kW", ""]
    done = copy.deepcopy(pkg.document)
    dm.resolve(done, wm.ACCEPT)
    assert wm.texts(done, wm.KEEP) == accepted


def test_table_rows_inserted_and_deleted(tmp_path):
    pkg = dm.Package(make_doc(tmp_path / "a.docx"))
    before = copy.deepcopy(pkg.document)
    writer = dm.Writer(pkg, ME)
    rows = list(pkg.document.iter(w("tr")))
    _, new = dm.insert_row_after(rows[1], ["2", "500 kW"], writer)
    assert new.find(f"{w('trPr')}/{w('ins')}") is not None
    trpr = [etree.QName(c).localname for c in new.find(w("trPr"))]
    assert trpr[-1] == "ins"  # after the other row properties (W5)
    dm.delete_row(rows[0], writer)
    assert wm.texts(pkg.document, wm.ACCEPT)[3:8] == ["1", "300 kW", "2", "500 kW", "Last paragraph."]
    assert wm.lines(pkg.document, ours(writer, "reject")) == wm.lines(before, wm.KEEP)
    view = wm.docx_view(pkg)
    assert view["census"]["by_kind"] == {"row inserted": 1, "row deleted": 1}


def test_a_copied_row_brings_no_old_revisions(tmp_path):
    pkg = dm.Package(hand_written(tmp_path / "hand.docx"))
    writer = dm.Writer(pkg, ME)
    tracked_row = list(pkg.document.iter(w("tr")))[1]  # B. Checker's inserted row
    copy_ = dm.clean_copy(tracked_row, writer)
    assert not list(copy_.iter(*dm.REVISION_TAGS))
    assert all(p.get(wm.qn("w14:paraId")) for p in copy_.iter(w("p")))


def test_duplicate_ids_are_found_and_renumbered(tmp_path):
    body = """<w:p><w:ins w:id="7" w:author="X"><w:r><w:t>a</w:t></w:r></w:ins>
      <w:ins w:id="7" w:author="X"><w:r><w:t>b</w:t></w:r></w:ins></w:p>"""
    pkg = dm.Package(hand_written(tmp_path / "dup.docx", body=body))
    pkg.touch(pkg.names["document"])
    assert any("used 2 times" in x for x in dm.structure_problems(pkg))
    writer = dm.Writer(pkg, ME)
    assert dm.renumber_duplicates(pkg, writer) == 1
    assert not any("used 2 times" in x for x in dm.structure_problems(pkg))


def test_new_ids_are_unique_and_paragraph_ids_valid(tmp_path):
    pkg = dm.Package(make_doc(tmp_path / "a.docx"))
    one, two = dm.Writer(pkg, ME), dm.Writer(pkg, "B. Checker")  # two writers on one package share the counter
    assert one.ids.new() != two.ids.new()
    values = {one.ids.hex8() for _ in range(200)}
    assert len(values) == 200 and all(int(v, 16) < 0x80000000 and len(v) == 8 for v in values)


def test_order_checker_catches_misplaced_properties(tmp_path):
    pkg = dm.Package(make_doc(tmp_path / "a.docx"))
    r = next(pkg.document.iter(w("r")))
    rpr = etree.SubElement(r, w("rPr"))  # after the text: wrong
    etree.SubElement(rpr, w("b"))
    pkg.touch(pkg.names["document"])
    assert any("first child" in x for x in dm.structure_problems(pkg))


def test_comments_threads_and_done_flag(tmp_path):
    pkg = dm.Package(make_doc(tmp_path / "a.docx"))
    a = dm.Writer(pkg, "A. Reviewer", "AR", "2026-10-02T10:00:00Z")
    me = dm.Writer(pkg, ME, "GBH")
    p, s, e = find(pkg, "every second")
    cid = dm.comment_on_text(pkg, p, s, e, "How often exactly?", a)
    first = dm.add_reply(pkg, cid, "Once a second.", me)
    second = dm.add_reply(pkg, cid, "Thanks.", a)
    dm.resolve_thread(pkg, cid, me)
    assert dm.structure_problems(pkg) == [] and dm.comment_problems(pkg) == []
    out = pkg.save(tmp_path / "out.docx")
    view = wm.docx_view(out)
    comments = {c["id"]: c for c in view["comments"]}
    assert comments[cid]["anchor"] == "every second" and comments[cid]["done"]
    assert comments[cid]["replies"] == [first, second] and comments[first]["author"] == ME
    assert view["census"]["threads"] == 1 and view["census"]["replies"] == 2 and view["census"]["resolved"] == 1
    assert ME in view["people"] and "A. Reviewer" in view["people"]
    marked = next(x["marked"] for x in view["paragraphs"] if "every second" in x["accepted"])
    assert marked.index(f"[c{first}]") < marked.index(f"[c{second}]")  # replies in their order
    with zipfile.ZipFile(out) as z:
        types = z.read("[Content_Types].xml").decode()
        rels = z.read("word/_rels/document.xml.rels").decode()
    for part in ("comments", "commentsExtended", "commentsIds", "people"):
        assert f"/word/{part}.xml" in types and f"{part}.xml" in rels


def test_save_keeps_parts_byte_for_byte_and_content_types_first(tmp_path):
    source = make_doc(tmp_path / "a.docx")
    labelled = tmp_path / "labelled.docx"
    with zipfile.ZipFile(source) as z, zipfile.ZipFile(labelled, "w", zipfile.ZIP_DEFLATED) as out:
        for info in z.infolist():
            out.writestr(info, z.read(info.filename))
        out.writestr("docMetadata/LabelInfo.xml", "<clbl:labelList xmlns:clbl='x'/>")  # a sensitivity label
    pkg = dm.Package(labelled)
    writer = dm.Writer(pkg, ME)
    p, a, b = find(pkg, "Last paragraph")
    dm.tracked_replace(p, a, b, "Final paragraph", writer)
    pkg.touch(pkg.names["document"])
    dm.set_track_changes(pkg, True)
    out = pkg.save(tmp_path / "out.docx")
    with zipfile.ZipFile(labelled) as old, zipfile.ZipFile(out) as new:
        assert new.namelist()[0] == "[Content_Types].xml"
        assert "docMetadata/LabelInfo.xml" in new.namelist()
        for name in old.namelist():
            if name not in ("word/document.xml", "word/settings.xml"):
                assert old.read(name) == new.read(name), name
        settings = etree.fromstring(new.read("word/settings.xml"))
    names = [etree.QName(c).localname for c in settings]
    assert "trackRevisions" in names and names.index("trackRevisions") < names.index("defaultTabStop")


def test_python_compare_passes_its_own_check(tmp_path):
    def edit(doc):
        doc.paragraphs[1].text = "The sprinkler must activate within 3 minutes of ignition, or the test stops."
        doc.paragraphs[2]._p.getparent().remove(doc.paragraphs[2]._p)
        row = doc.tables[0].add_row()
        row.cells[0].text, row.cells[1].text = "2", "500 kW"
        doc.add_paragraph("A new last paragraph.")
    old, new = make_doc(tmp_path / "v1.docx"), make_doc(tmp_path / "v2.docx", extra=edit)
    pkg, _ = dm.python_compare(old, new, ME)
    out = pkg.save(tmp_path / "redline.docx")
    root = dm.Package(out).document
    assert wm.texts(root, wm.ACCEPT) == wm.texts(dm.Package(new).document, wm.ACCEPT)
    assert wm.texts(root, wm.REJECT) == wm.texts(dm.Package(old).document, wm.ACCEPT)
    census = wm.docx_view(out)["census"]
    assert census["by_author"] == {ME: census["changes"]} and census["by_kind"]["row inserted"] == 1
    assert dm.structure_problems(dm.Package(out)) == []


def test_pair_up_keeps_order_and_pairs_similar_items():
    old = ["alpha beta gamma", "delta epsilon", "zeta eta theta"]
    new = ["alpha beta gamma", "zeta eta theta iota", "new one"]
    pairs = dm.pair_up(old, new, key=lambda x: x, score=dm.similarity)
    assert pairs == [("alpha beta gamma", "alpha beta gamma"), ("delta epsilon", None),
                     ("zeta eta theta", "zeta eta theta iota"), (None, "new one")]
